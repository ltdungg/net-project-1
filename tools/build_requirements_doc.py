from __future__ import annotations

from pathlib import Path
from datetime import date
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
WORK_DIR = ROOT / ".docx-work" / "requirements"
OUTPUT = OUT_DIR / "Tai-lieu-yeu-cau-chuc-nang-va-phi-chuc-nang-NET.docx"
DIAGRAM = WORK_DIR / "nfr-taxonomy.png"

INK = "1F2A28"
GREEN = "173B35"
GREEN_2 = "276255"
BEIGE = "E9DED0"
BEIGE_2 = "F6F1EA"
MUTED = "66706D"
LINE = "C9D2CF"
WHITE = "FFFFFF"
RED = "9C2F2F"
AMBER = "8A6116"
BLUE = "275C78"


FUNCTIONAL_GROUPS = [
    ("FR-01", "Khám phá và danh mục sản phẩm", [
        ("FR-CAT-001", "Hiển thị trang chủ với bộ sưu tập, sản phẩm nổi bật và điều hướng theo nhóm Nữ, Nam, Phụ kiện.", "Khách", "M", "Chỉ dữ liệu active; liên kết mở đúng trang."),
        ("FR-CAT-002", "Liệt kê sản phẩm theo danh mục, có tổng số kết quả và trạng thái không có dữ liệu.", "Khách", "M", "Kết quả đúng danh mục; empty state rõ ràng."),
        ("FR-CAT-003", "Tìm kiếm theo tên, mô tả ngắn, chất liệu và từ khóa liên quan.", "Khách", "M", "Không phân biệt hoa thường; phản hồi đúng truy vấn."),
        ("FR-CAT-004", "Lọc kết hợp theo danh mục, kích cỡ, màu, khoảng giá và tình trạng còn hàng.", "Khách", "S", "Kết hợp nhiều bộ lọc; có thể xóa từng/all."),
        ("FR-CAT-005", "Sắp xếp theo mới nhất, giá tăng/giảm và mức độ nổi bật.", "Khách", "S", "Thứ tự ổn định; giữ nguyên bộ lọc hiện tại."),
        ("FR-CAT-006", "Phân trang hoặc tải thêm sản phẩm mà không tạo bản ghi trùng.", "Khách", "S", "Không thiếu/trùng khi chuyển trang."),
        ("FR-CAT-007", "Chỉ hiển thị sản phẩm và biến thể đang được phép bán.", "Hệ thống", "M", "Draft/archived không xuất hiện ngoài storefront."),
        ("FR-CAT-008", "Cung cấp URL có thể chia sẻ cho trạng thái tìm kiếm, lọc và sắp xếp.", "Khách", "C", "Mở URL tái lập đúng trạng thái."),
        ("FR-CAT-009", "Hiển thị giá khởi điểm, giá so sánh và trạng thái tồn tổng hợp trên thẻ sản phẩm.", "Khách", "M", "Giá/tồn lấy từ backend; định dạng VND."),
        ("FR-CAT-010", "Ghi nhận sự kiện xem danh sách, tìm kiếm, lọc và chọn sản phẩm phục vụ phân tích funnel.", "Hệ thống", "S", "Sự kiện có schema chuẩn, không chứa dữ liệu nhạy cảm."),
    ]),
    ("FR-02", "Chi tiết sản phẩm và biến thể", [
        ("FR-PDP-001", "Hiển thị tên, mô tả, chất liệu, thư viện ảnh, giá, đánh giá và thông tin đổi trả.", "Khách", "M", "Nội dung khớp sản phẩm; ảnh có alt text."),
        ("FR-PDP-002", "Cho phép chọn biến thể theo màu và kích cỡ; cập nhật SKU, giá và tồn tương ứng.", "Khách", "M", "Mỗi lựa chọn ánh xạ đúng một variant."),
        ("FR-PDP-003", "Vô hiệu hóa biến thể không active hoặc hết tồn khả dụng.", "Hệ thống", "M", "Không thể thêm biến thể không bán được."),
        ("FR-PDP-004", "Hiển thị hướng dẫn chọn kích cỡ và thông tin số đo.", "Khách", "S", "Mở được trên mobile/desktop; đóng bằng bàn phím."),
        ("FR-PDP-005", "Cho phép chọn số lượng trong giới hạn mua và tồn khả dụng.", "Khách", "M", "Quantity nguyên dương, tối đa min(10, ATP)."),
        ("FR-PDP-006", "Thêm biến thể đã chọn vào giỏ hàng và phản hồi thành công/lỗi rõ ràng.", "Khách", "M", "Giỏ tăng đúng; header cập nhật số lượng."),
        ("FR-PDP-007", "Hiển thị sản phẩm liên quan theo danh mục mà không lộ sản phẩm draft.", "Khách", "C", "Tối đa cấu hình; không lặp sản phẩm hiện tại."),
        ("FR-PDP-008", "Thiết lập metadata chia sẻ theo từng sản phẩm.", "Hệ thống", "S", "Title/description/image đúng bản ghi."),
    ]),
    ("FR-03", "Giỏ hàng và khuyến mãi", [
        ("FR-CAR-001", "Tạo giỏ khách bằng cookie ẩn danh và duy trì giỏ trong thời hạn cấu hình.", "Hệ thống", "M", "Cookie HttpOnly; giỏ tồn tại qua lần tải lại."),
        ("FR-CAR-002", "Gộp dòng khi thêm lại cùng một biến thể; không tạo dòng trùng.", "Hệ thống", "M", "Một cart chỉ có một dòng/variant."),
        ("FR-CAR-003", "Cho phép tăng, giảm và xóa sản phẩm trong giỏ.", "Khách", "M", "Tổng tiền cập nhật sau mỗi thao tác."),
        ("FR-CAR-004", "Tính lại giá và tồn từ backend khi đọc hoặc thay đổi giỏ.", "Hệ thống", "M", "Không tin giá/tổng do client gửi."),
        ("FR-CAR-005", "Hiển thị tạm tính, giảm giá, phí vận chuyển và tổng thanh toán.", "Khách", "M", "Các thành phần cộng khớp grand total."),
        ("FR-CAR-006", "Áp dụng một mã voucher cho giỏ và giải thích nguyên nhân nếu không hợp lệ.", "Khách", "M", "Kiểm tra thời gian, min spend, giới hạn sử dụng."),
        ("FR-CAR-007", "Cho phép gỡ voucher và hoàn nguyên tổng tiền.", "Khách", "M", "Discount về 0; shipping được tính lại."),
        ("FR-CAR-008", "Hỗ trợ voucher phần trăm, số tiền cố định và miễn phí giao hàng.", "Hệ thống", "S", "Discount không vượt subtotal/max discount."),
        ("FR-CAR-009", "Thông báo ngưỡng còn thiếu để được miễn phí giao hàng.", "Khách", "S", "Giá trị còn thiếu không âm."),
        ("FR-CAR-010", "Xử lý giỏ rỗng, sản phẩm bị ngừng bán và biến thể thiếu tồn.", "Hệ thống", "M", "Có hướng dẫn khắc phục, không lỗi trang."),
    ]),
    ("FR-04", "Đăng ký, đăng nhập và tài khoản", [
        ("FR-ACC-001", "Đăng ký tài khoản bằng email, mật khẩu mạnh, họ tên và số điện thoại.", "Khách", "S", "Email unique; mật khẩu hash; validation đầy đủ."),
        ("FR-ACC-002", "Đăng nhập và cấp phiên truy cập an toàn.", "Khách", "S", "Sai thông tin trả lỗi chung; không lộ tài khoản."),
        ("FR-ACC-003", "Làm mới phiên bằng refresh token xoay vòng và thu hồi token cũ.", "Hệ thống", "S", "Token cũ không dùng lại sau refresh."),
        ("FR-ACC-004", "Đăng xuất và thu hồi phiên hiện tại.", "Khách", "S", "Cookie hết hạn; refresh token bị revoke."),
        ("FR-ACC-005", "Xem và cập nhật hồ sơ cá nhân, số điện thoại, tùy chọn nhận tin.", "Khách", "S", "Chỉ chủ tài khoản cập nhật được."),
        ("FR-ACC-006", "Quản lý nhiều địa chỉ giao hàng và chọn một địa chỉ mặc định.", "Khách", "S", "Tối đa một địa chỉ mặc định/user."),
        ("FR-ACC-007", "Xem danh sách và chi tiết các đơn thuộc tài khoản.", "Khách", "S", "Không truy cập được đơn của người khác."),
        ("FR-ACC-008", "Phân quyền customer, staff, manager và admin theo nguyên tắc ít quyền nhất.", "Hệ thống", "M", "API kiểm tra quyền ở server."),
    ]),
    ("FR-05", "Checkout, đơn hàng và thanh toán", [
        ("FR-CHK-001", "Cho phép guest checkout không bắt buộc tạo tài khoản.", "Khách", "M", "Có thể đặt đơn bằng email/phone hợp lệ."),
        ("FR-CHK-002", "Thu thập người nhận, điện thoại, địa chỉ, phường/xã, quận/huyện, tỉnh/thành và ghi chú.", "Khách", "M", "Trường bắt buộc được kiểm tra server-side."),
        ("FR-CHK-003", "Cho phép chọn COD hoặc thanh toán trực tuyến.", "Khách", "M", "Phương thức hợp lệ được snapshot vào payment."),
        ("FR-CHK-004", "Hiển thị lại toàn bộ sản phẩm và tổng tiền trước khi xác nhận đặt hàng.", "Khách", "M", "Grand total cuối hiển thị trước CTA."),
        ("FR-ORD-001", "Tạo duy nhất một đơn cho mỗi idempotency key hợp lệ.", "Hệ thống", "M", "Gửi lặp trả cùng kết quả, không nhân đôi side effect."),
        ("FR-ORD-002", "Đọc lại giá, voucher và tồn kho trong giao dịch tạo đơn.", "Hệ thống", "M", "Dữ liệu client không được dùng làm nguồn sự thật."),
        ("FR-ORD-003", "Chọn một kho có đủ tồn cho từng dòng và khóa bản ghi khi giữ tồn.", "Hệ thống", "M", "Hai checkout SKU cuối chỉ một đơn thành công."),
        ("FR-ORD-004", "Tạo reservation tồn kho có thời hạn cho đơn online chờ thanh toán.", "Hệ thống", "M", "Reserved tăng đúng; có expiresAt."),
        ("FR-ORD-005", "Tạo mã đơn duy nhất và lưu snapshot sản phẩm, SKU, giá, địa chỉ, voucher.", "Hệ thống", "M", "Thay đổi catalog sau đó không đổi đơn cũ."),
        ("FR-ORD-006", "Cho phép chuyển trạng thái đơn theo state machine đã phê duyệt.", "Nhân viên", "M", "Transition sai bị từ chối; có lịch sử."),
        ("FR-ORD-007", "Hủy đơn trước khi xuất kho và giải phóng reservation tương ứng.", "Nhân viên", "M", "Reserved giảm đúng; ghi lý do/actor."),
        ("FR-ORD-008", "Khi chuyển sang shipped, tiêu thụ reservation và trừ on-hand một lần.", "Hệ thống", "M", "Không trừ trùng khi request lặp."),
        ("FR-PAY-001", "Tạo payment record gắn với đơn, số tiền, tiền tệ và idempotency key.", "Hệ thống", "M", "Một attempt có một payment duy nhất."),
        ("FR-PAY-002", "Xác minh chữ ký, mã sự kiện, số tiền và tiền tệ của webhook thanh toán.", "Hệ thống", "M", "Sai chữ ký trả 401, không đổi order."),
        ("FR-PAY-003", "Chống xử lý trùng hoặc lùi trạng thái khi webhook đến lặp/out-of-order.", "Hệ thống", "M", "Provider event ID unique; terminal không bị lùi."),
        ("FR-PAY-004", "Cho phép mô phỏng thanh toán chỉ trong môi trường local.", "Dev/QA", "C", "Production trả LOCAL_ONLY."),
    ]),
    ("FR-06", "Quản trị catalog, khách hàng, đơn và khuyến mãi", [
        ("FR-ADM-001", "Dashboard hiển thị đơn 30 ngày, doanh thu, khách mới, backlog và SKU tồn thấp.", "Quản lý", "M", "Chỉ số lấy từ dữ liệu giao dịch hiện tại."),
        ("FR-ADM-002", "Liệt kê sản phẩm gồm trạng thái, biến thể, giá, tồn và ảnh chính.", "Nhân viên", "M", "Có draft/active/archived; dữ liệu đúng DB."),
        ("FR-ADM-003", "Tạo sản phẩm, danh mục chính và biến thể đầu tiên trong một giao dịch.", "Quản lý", "M", "Slug/SKU unique; rollback nếu một bước lỗi."),
        ("FR-ADM-004", "Tạo tồn đầu kỳ cho biến thể mới và ghi movement tương ứng.", "Quản lý", "M", "On-hand và sổ cái khớp nhau."),
        ("FR-ADM-005", "Tải ảnh sản phẩm, kiểm tra loại/kích thước và liên kết media với sản phẩm.", "Quản lý", "S", "Chỉ ảnh cho phép; URL truy cập được."),
        ("FR-ADM-006", "Liệt kê, lọc và mở chi tiết đơn hàng.", "Nhân viên", "M", "Tìm được theo mã/email/trạng thái."),
        ("FR-ADM-007", "Cập nhật trạng thái đơn với lý do và audit trail.", "Nhân viên", "M", "Lưu from/to/actor/timestamp."),
        ("FR-ADM-008", "Xem danh sách khách hàng, trạng thái, tổng đơn và tổng chi tiêu.", "Nhân viên", "S", "PII hiển thị theo quyền."),
        ("FR-ADM-009", "Tạo chương trình khuyến mãi và coupon với thời gian, điều kiện, giới hạn.", "Quản lý", "S", "Start < end; code unique; value hợp lệ."),
        ("FR-ADM-010", "Xem báo cáo doanh thu, sản phẩm bán chạy và chỉ số vận hành.", "Quản lý", "S", "Có khoảng thời gian và định nghĩa chỉ số."),
    ]),
    ("FR-07", "Quản lý tồn kho đa kho", [
        ("FR-INV-001", "Hiển thị on-hand, reserved, available, safety stock, ATP và reorder point theo SKU/kho.", "Nhân viên", "M", "Số tổng hợp khớp từng dòng."),
        ("FR-INV-002", "Lọc tồn theo kho, trạng thái sắp hết/hết hàng/an toàn và tìm theo SKU/sản phẩm.", "Nhân viên", "M", "Kết quả đúng điều kiện kết hợp."),
        ("FR-INV-003", "Tạo phiếu nhập kho và tăng on-hand trong một giao dịch nguyên tử.", "Nhân viên", "M", "Operation, item, movement, audit cùng commit."),
        ("FR-INV-004", "Tạo phiếu chuyển kho, trừ kho nguồn và cộng kho đích nguyên tử.", "Nhân viên", "M", "Không mất tồn; kho nguồn đủ available."),
        ("FR-INV-005", "Tạo dòng tồn ở kho đích nếu variant chưa tồn tại tại kho đó.", "Hệ thống", "M", "Unique warehouse/variant; không tạo trùng."),
        ("FR-INV-006", "Chốt kiểm kê theo số đếm thực tế và ghi chênh lệch.", "Nhân viên", "M", "Actual không thấp hơn reserved; có variance."),
        ("FR-INV-007", "Điều chỉnh tăng/giảm tồn có lý do; chỉ manager/admin được thực hiện.", "Quản lý", "M", "RBAC; balance không âm/thấp hơn reserved."),
        ("FR-INV-008", "Mọi phiếu kho hỗ trợ idempotency key và trả kết quả cũ khi gửi lặp.", "Hệ thống", "M", "Không nhân đôi movement."),
        ("FR-INV-009", "Movement tồn kho là append-only; sửa sai bằng movement bù.", "Hệ thống", "M", "Không API sửa/xóa movement đã ghi."),
        ("FR-INV-010", "Tự động giải phóng reservation hết hạn và hủy đơn pending tương ứng.", "Nhân viên", "M", "Reserved giảm; đơn cập nhật một lần."),
        ("FR-INV-011", "Hiển thị lịch sử phiếu kho và sổ cái biến động gần nhất.", "Nhân viên", "M", "Có mã phiếu, lý do, kho, balance after."),
        ("FR-INV-012", "Xuất snapshot tồn kho CSV có BOM UTF-8.", "Nhân viên", "S", "Mở đúng tiếng Việt trong Excel."),
    ]),
    ("FR-08", "AI, báo cáo, audit và tích hợp", [
        ("FR-AI-001", "Trợ lý quản trị trả lời câu hỏi về doanh thu, backlog, tồn thấp và khuyến mãi dựa trên dữ liệu nội bộ.", "Nhân viên", "S", "Câu trả lời có thời điểm, confidence, grounded=true."),
        ("FR-AI-002", "AI chỉ cung cấp đề xuất và đường dẫn thao tác; không tự thay đổi đơn, giá hoặc tồn.", "Hệ thống", "M", "Không có mutation từ prompt."),
        ("FR-AI-003", "Khi dịch vụ AI ngoài không khả dụng, hệ thống dùng phân tích quy tắc hoặc thông báo fallback.", "Hệ thống", "S", "Admin vẫn nhận tóm tắt vận hành."),
        ("FR-RPT-001", "Cung cấp API metric cho đơn, doanh thu, tồn thấp và khách mới 30 ngày.", "Quản lý", "M", "Công thức chỉ số được cố định và kiểm thử."),
        ("FR-AUD-001", "Ghi audit cho thay đổi nhạy cảm với actor, action, entity, before/after và correlation ID.", "Hệ thống", "M", "Tra cứu được chuỗi thao tác."),
        ("FR-INT-001", "Ghi sự kiện order.created vào transactional outbox.", "Hệ thống", "M", "Order và event cùng commit."),
        ("FR-INT-002", "Cung cấp health endpoint kiểm tra ứng dụng và kết nối cơ sở dữ liệu.", "DevOps", "M", "Trả 200 chỉ khi dependency bắt buộc sẵn sàng."),
        ("FR-INT-003", "Mọi lỗi API dùng envelope thống nhất, stable code và correlation ID.", "Hệ thống", "M", "Frontend xử lý được không dựa vào message."),
    ]),
]


NFR_GROUPS = [
    ("NFR-01", "Khả dụng và trải nghiệm sử dụng", [
        ("NFR-USA-001", "Sử dụng", "Luồng mua hàng cốt lõi hoàn thành trên mobile không cần hướng dẫn ngoài giao diện.", "≥ 85% người thử hoàn tất; median ≤ 5 phút", "Usability test", "M"),
        ("NFR-USA-002", "Học dùng", "Nhân viên mới thực hiện nhập kho và xử lý đơn sau hướng dẫn ngắn.", "≥ 90% hoàn tất sau ≤ 30 phút đào tạo", "UAT", "S"),
        ("NFR-USA-003", "Khả năng tiếp cận", "Storefront và checkout đáp ứng WCAG 2.2 mức AA.", "0 lỗi critical trong axe; keyboard flow pass", "A11y audit", "M"),
        ("NFR-USA-004", "Responsive", "Hỗ trợ viewport 360–1920 px, không tràn ngang ở luồng chính.", "Pass ma trận thiết bị mục tiêu", "Visual QA", "M"),
        ("NFR-USA-005", "Phản hồi", "Mọi thao tác bất đồng bộ có loading, success, empty và error state phù hợp.", "100% flow Must có đủ trạng thái", "UX checklist", "M"),
        ("NFR-USA-006", "Ngôn ngữ", "Nội dung người dùng là tiếng Việt rõ nghĩa; tiền tệ và ngày giờ nhất quán.", "VND/vi-VN; 0 chuỗi kỹ thuật lộ ra UI", "Content QA", "M"),
    ]),
    ("NFR-02", "Hiệu năng và khả năng mở rộng", [
        ("NFR-PER-001", "Web performance", "Trang chủ/PLP/PDP đạt Core Web Vitals trên mobile thực tế.", "LCP p75 ≤ 2,5s; INP ≤ 200ms; CLS ≤ 0,1", "RUM/Lighthouse", "M"),
        ("NFR-PER-002", "API đọc", "API catalog, cart và dashboard phổ biến đáp ứng nhanh ở tải mục tiêu.", "p95 ≤ 500ms; p99 ≤ 1s", "Load test", "M"),
        ("NFR-PER-003", "API ghi", "Tạo đơn và nghiệp vụ kho không tính thời gian provider ngoài.", "p95 ≤ 1s; p99 ≤ 2s", "Load test", "M"),
        ("NFR-PER-004", "Dung lượng", "Baseline phục vụ đồng thời người dùng và nhân viên vận hành của MVP.", "300 concurrent; 20 req/s sustained; 50 req/s burst", "Load test", "S"),
        ("NFR-PER-005", "Database", "Truy vấn danh sách không quét toàn bảng khi dữ liệu tăng.", "Query plan dùng index; không N+1 trên API chính", "DB review", "M"),
        ("NFR-PER-006", "Media", "Ảnh sản phẩm được tối ưu kích thước và cache phù hợp.", "Ảnh hiển thị ≤ 300KB mục tiêu; cache ≥ 7 ngày", "Network audit", "S"),
        ("NFR-PER-007", "Mở rộng ngang", "API không phụ thuộc session memory cục bộ.", "Thêm instance không mất phiên/giao dịch", "Architecture test", "S"),
    ]),
    ("NFR-03", "Độ tin cậy, toàn vẹn và phục hồi", [
        ("NFR-REL-001", "Sẵn sàng", "Dịch vụ giao dịch đạt mức sẵn sàng sau giai đoạn ổn định.", "≥ 99,9% mỗi tháng, trừ bảo trì báo trước", "Uptime report", "M"),
        ("NFR-REL-002", "Toàn vẹn đơn", "Một idempotency key không tạo nhiều order/payment/reservation.", "0 side effect trùng trong concurrency test", "Integration test", "M"),
        ("NFR-REL-003", "Toàn vẹn tồn", "Không để reserved > on-hand hoặc on-hand < 0.", "Constraint + 100% test cạnh tranh pass", "DB/test", "M"),
        ("NFR-REL-004", "Giao dịch", "Order, reservation, payment và outbox commit/rollback cùng nhau khi cần.", "Fault injection không tạo trạng thái mồ côi", "Integration test", "M"),
        ("NFR-REL-005", "Retry", "Tích hợp ngoài có timeout, retry giới hạn và idempotency.", "≤ 3 retry exponential; không retry lỗi 4xx", "Resilience test", "S"),
        ("NFR-REL-006", "Khôi phục", "Có backup định kỳ và quy trình phục hồi dữ liệu.", "RPO ≤ 24h; RTO ≤ 4h", "Restore drill", "M"),
        ("NFR-REL-007", "Triển khai", "Có rollback khi release thất bại.", "Khôi phục version trước ≤ 15 phút", "Release drill", "M"),
        ("NFR-REL-008", "Graceful shutdown", "Instance ngừng nhận request mới và hoàn tất giao dịch đang chạy.", "0 request ghi bị cắt trong rolling restart test", "Ops test", "S"),
    ]),
    ("NFR-04", "Bảo mật và quyền riêng tư", [
        ("NFR-SEC-001", "Truyền dữ liệu", "Mọi traffic production dùng TLS hiện hành và HSTS.", "TLS 1.2+; HSTS ≥ 180 ngày", "Security scan", "M"),
        ("NFR-SEC-002", "Xác thực", "Mật khẩu được hash bằng thuật toán thích hợp; token có hạn và xoay vòng.", "BCrypt cost ≥ 12; access ≤ 15 phút", "Code/security review", "M"),
        ("NFR-SEC-003", "Phân quyền", "Authorization server-side theo customer/staff/manager/admin.", "100% API nhạy cảm có negative test", "Security test", "M"),
        ("NFR-SEC-004", "CSRF/CORS", "Cookie xác thực có SameSite/HttpOnly/Secure và kiểm tra origin.", "Cross-origin mutation bị từ chối", "Security test", "M"),
        ("NFR-SEC-005", "Injection", "Query parameter hóa; input/output validation theo allow-list.", "0 Critical/High từ SAST/DAST", "Security scan", "M"),
        ("NFR-SEC-006", "Bí mật", "Secret không nằm trong source, image, log hoặc response.", "Secret scan pass; rotate được không rebuild", "CI audit", "M"),
        ("NFR-SEC-007", "Dữ liệu nhạy cảm", "Không lưu PAN/CVV; PII được giảm thiểu và masking theo quyền.", "0 dữ liệu thẻ; PII list view được kiểm soát", "Privacy review", "M"),
        ("NFR-SEC-008", "Rate limit", "Login, tra cứu đơn và webhook được bảo vệ chống lạm dụng.", "429 theo policy; không khóa nhầm hàng loạt", "Abuse test", "S"),
        ("NFR-SEC-009", "Audit", "Audit log nhạy cảm append-only và không chứa secret.", "Truy vết actor/action/time/correlation", "Audit review", "M"),
        ("NFR-SEC-010", "Vulnerability", "Dependency và container được quét trước phát hành.", "0 Critical/High mở hoặc có risk acceptance", "CI report", "M"),
    ]),
    ("NFR-05", "Linh động, bảo trì và chuyển giao", [
        ("NFR-MAI-001", "Kiến trúc", "Backend tổ chức modular monolith theo domain, tránh phụ thuộc vòng.", "Module boundary review pass", "Architecture review", "S"),
        ("NFR-MAI-002", "API", "API version/contract tương thích ngược trong phạm vi release.", "Breaking change có version và migration plan", "Contract test", "M"),
        ("NFR-MAI-003", "Kiểm thử", "Business rule quan trọng có unit/integration test tự động.", "≥ 80% service/domain; 100% rule critical", "Coverage report", "S"),
        ("NFR-MAI-004", "Mã nguồn", "Code format, static analysis và review bắt buộc trước merge.", "CI pass; tối thiểu 1 reviewer", "CI/repository", "M"),
        ("NFR-MAI-005", "Migration", "Schema thay đổi bằng migration có thứ tự, immutable sau release.", "Fresh install và upgrade path đều pass", "Migration test", "M"),
        ("NFR-MAI-006", "Cấu hình", "Ngưỡng vận hành và credentials cấu hình ngoài code.", "Thay đổi env không sửa source", "Config audit", "S"),
        ("NFR-MAI-007", "Tài liệu", "Có README, API guide, runbook, ERD và quyết định kiến trúc.", "Onboarding thành viên mới ≤ 1 ngày", "Handover UAT", "S"),
    ]),
    ("NFR-06", "Tương thích, cài đặt và chuẩn tổ chức", [
        ("NFR-COM-001", "Trình duyệt", "Hỗ trợ 2 phiên bản mới nhất của Chrome, Edge, Firefox và Safari.", "Pass browser matrix cho flow Must", "Cross-browser test", "M"),
        ("NFR-COM-002", "Triển khai", "Môi trường local chạy bằng một lệnh Docker Compose.", "Fresh clone → healthy ≤ 10 phút", "Onboarding test", "M"),
        ("NFR-COM-003", "Nhất quán", "Dev/test/prod dùng cùng artifact và cấu hình qua environment.", "Không rebuild riêng cho từng môi trường", "Pipeline audit", "S"),
        ("NFR-COM-004", "Chuẩn dữ liệu", "Thời gian lưu UTC, hiển thị Asia/Ho_Chi_Minh; tiền dùng integer VND.", "0 lỗi timezone/rounding trong test", "Data test", "M"),
        ("NFR-COM-005", "CSV", "File xuất dùng UTF-8 BOM và escape đúng chuẩn CSV.", "Mở đúng trong Excel/LibreOffice", "Compatibility test", "S"),
        ("NFR-COM-006", "Tích hợp", "Webhook/API dùng JSON UTF-8, stable error code và correlation ID.", "Contract test pass", "Contract test", "M"),
    ]),
    ("NFR-07", "Pháp lý, quy tắc bên ngoài và đạo đức", [
        ("NFR-EXT-001", "Quyền riêng tư", "Thu thập, sử dụng và lưu PII theo mục đích đã công bố và consent phù hợp.", "Privacy checklist được phê duyệt", "Legal review", "M"),
        ("NFR-EXT-002", "Lưu trữ", "Có thời hạn giữ và quy trình xóa/ẩn danh dữ liệu tài khoản, log, audit.", "Retention matrix + job kiểm chứng", "Compliance test", "S"),
        ("NFR-EXT-003", "Thương mại", "Hiển thị chính sách giá, giao nhận, đổi trả, điều khoản trước xác nhận mua.", "Legal sign-off trước go-live", "UAT/legal", "M"),
        ("NFR-EXT-004", "Thanh toán", "Không đưa dữ liệu thẻ vào hệ thống; tuân theo phạm vi provider/PCI được duyệt.", "SAQ/scope được xác nhận", "Security/legal", "M"),
        ("NFR-EXT-005", "AI có trách nhiệm", "AI không tự động thay đổi dữ liệu và phải nêu rõ đây là đề xuất dựa trên snapshot.", "100% response có grounded/time; mutation=0", "AI evaluation", "M"),
        ("NFR-EXT-006", "Không phân biệt", "Gợi ý AI không sử dụng thuộc tính nhạy cảm hoặc suy diễn cá nhân không cần thiết.", "Prompt/data review pass", "AI/privacy review", "S"),
    ]),
    ("NFR-08", "Quan sát và vận hành", [
        ("NFR-OPS-001", "Logging", "Log có cấu trúc, correlation ID, severity và không chứa secret/PII dư thừa.", "100% request lỗi tra được end-to-end", "Log audit", "M"),
        ("NFR-OPS-002", "Metrics", "Theo dõi latency, error rate, traffic, saturation và business event chính.", "Dashboard RED + order/payment/inventory", "Ops review", "M"),
        ("NFR-OPS-003", "Alert", "Cảnh báo theo triệu chứng và có runbook xử lý.", "Alert test gửi đúng kênh ≤ 5 phút", "Game day", "M"),
        ("NFR-OPS-004", "Health", "Liveness/readiness phân biệt và phản ánh dependency bắt buộc.", "Orchestrator không route vào instance chưa ready", "Deployment test", "M"),
        ("NFR-OPS-005", "SLO", "Có error budget và báo cáo SLO theo tháng sau go-live.", "Báo cáo availability/latency/error", "Ops report", "C"),
    ]),
]


QUALITY_SCENARIOS = [
    ("QS-01", "Cạnh tranh SKU cuối", "Hai khách xác nhận mua cùng một SKU còn 1 ATP", "Giờ cao điểm", "Inventory/Order", "Hệ thống khóa/kiểm tra tồn nguyên tử; một request thành công, một request nhận lỗi tồn", "Không oversell; reserved ≤ on-hand"),
    ("QS-02", "Webhook gửi lặp", "Provider gửi cùng event thanh toán 5 lần", "Production", "Payment", "Chỉ event đầu tạo side effect; các event sau trả received/duplicate", "Một payment transition; một email/outbox"),
    ("QS-03", "DB lỗi giữa tạo đơn", "Kết nối lỗi sau khi insert order nhưng trước reservation", "Fault injection", "Checkout", "Toàn bộ giao dịch rollback hoặc tiếp tục an toàn theo idempotency", "Không order mồ côi; retry cho cùng kết quả"),
    ("QS-04", "Tải cao catalog", "300 người dùng đồng thời duyệt PLP/PDP", "Campaign", "Catalog", "Cache và DB phục vụ trong ngân sách latency", "API p95 ≤ 500ms; lỗi < 1%"),
    ("QS-05", "Khôi phục backup", "Mất database primary", "Disaster drill", "Toàn hệ thống", "Khôi phục bản backup gần nhất và xác minh dữ liệu cốt lõi", "RPO ≤ 24h; RTO ≤ 4h"),
    ("QS-06", "Tấn công quyền", "Customer gọi API điều chỉnh tồn", "Mọi thời điểm", "Authorization", "Server từ chối trước business service và ghi security event phù hợp", "HTTP 403; inventory không đổi"),
    ("QS-07", "AI provider timeout", "Dịch vụ AI không phản hồi trong 30 giây", "Admin đang xem dashboard", "AI Copilot", "Chuyển sang tóm tắt quy tắc hoặc thông báo fallback; không khóa UI", "Phản hồi hữu ích ≤ 3s sau timeout/fallback"),
    ("QS-08", "Rolling restart", "Một instance bị dừng khi đang xử lý request ghi", "Deploy", "API", "Instance graceful shutdown; traffic chuyển sang instance ready", "0 request ghi bị cắt; rollback ≤ 15 phút"),
    ("QS-09", "Nội dung không truy cập", "Khách dùng bàn phím và screen reader để checkout", "Mobile/desktop", "Storefront", "Focus order, label, error và CTA được đọc/điều khiển đúng", "WCAG 2.2 AA; flow hoàn tất"),
    ("QS-10", "CSV tiếng Việt", "Nhân viên tải báo cáo tồn và mở bằng Excel", "Windows", "Inventory Export", "File có BOM, escape dấu phẩy/nháy kép và dữ liệu đúng snapshot", "Không lỗi font; số dòng khớp DB"),
]


TRACE_ROWS = [
    ("E01 Catalog", "FR-CAT-001…010", "UC-01/02", "products, categories, variants", "GET /api/products", "TC-CAT-*"),
    ("E02 Product", "FR-PDP-001…008", "UC-03", "products, variants, media, inventory_items", "GET /api/products/{slug}", "TC-PDP-*"),
    ("E03 Cart", "FR-CAR-001…010", "UC-04/05", "carts, cart_items, coupons", "/api/cart; /api/vouchers", "TC-CAR-*"),
    ("E04 Account", "FR-ACC-001…008", "UC-06/07", "users, roles, addresses, refresh_tokens", "/api/auth; /api/account", "TC-ACC-*"),
    ("E05 Checkout", "FR-CHK-001…004", "UC-08", "carts, customer_orders", "POST /api/orders", "TC-CHK-*"),
    ("E06 Order", "FR-ORD-001…008", "UC-09/10", "customer_orders, order_items, history", "/api/orders; /api/admin/orders", "TC-ORD-*"),
    ("E07 Payment", "FR-PAY-001…004", "UC-11", "payments, payment_events", "/api/payments", "TC-PAY-*"),
    ("E08 Admin", "FR-ADM-001…010", "UC-12…15", "catalog/order/user/promotion tables", "/api/admin/*", "TC-ADM-*"),
    ("E09 Inventory", "FR-INV-001…012", "UC-16…20", "inventory_*; warehouses", "/api/admin/inventory/*", "TC-INV-*"),
    ("E10 AI/Report", "FR-AI/RPT-*", "UC-21/22", "orders, inventory, promotions", "/api/admin/ai; metrics", "TC-AI/RPT-*"),
    ("E11 Audit/Integration", "FR-AUD/INT-*", "UC-23", "audit_logs, outbox_events", "health/error/outbox", "TC-INT-*"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120) -> None:
    total = int(sum(widths_in) * 1440)
    tbl_pr = table._tbl.tblPr
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [int(width * 1440) for width in widths_in]
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(widths_in[min(idx, len(widths_in) - 1)])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_cell_text(cell, value, *, bold=False, color=INK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(value))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, headers: list[str], widths: list[float], body_size=8.7) -> None:
    set_repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=WHITE, size=9)
        set_cell_shading(table.rows[0].cells[idx], GREEN)
    for row_idx, row in enumerate(table.rows[1:], 1):
        prevent_row_split(row)
        if row_idx % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, BEIGE_2)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(body_size)
    set_table_geometry(table, widths)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_numbering(doc: Document, fmt: str, text: str, font: str = "Calibri") -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_bullet(doc, text: str, bullet_id: int, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="Requirement Bullet")
    marker = p.add_run("•  ")
    marker.bold = True
    marker.font.color.rgb = RGBColor.from_string(GREEN)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_callout(doc, label: str, text: str, fill=BEIGE_2) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), GREEN_2)
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(label.upper() + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor.from_string(INK)


def make_diagram() -> None:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (2400, 1260), "#FAF8F4")
    draw = ImageDraw.Draw(canvas)
    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    font_root = ImageFont.truetype(bold_path, 34)
    font_group = ImageFont.truetype(bold_path, 29)
    font_child = ImageFont.truetype(regular_path, 25)
    font_note = ImageFont.truetype(regular_path, 22)

    def centered_text(box, label, font, fill):
        lines = label.split("\n")
        line_height = font.getbbox("Ag")[3] + 5
        total_height = line_height * len(lines)
        x1, y1, x2, y2 = box
        y = y1 + ((y2 - y1) - total_height) / 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font)
            width = bbox[2] - bbox[0]
            draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=fill)
            y += line_height

    def box(rect, label, fill, outline, font, color, radius=18):
        draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=3)
        centered_text(rect, label, font, color)

    root = (780, 80, 1620, 190)
    groups = {
        "product": (90, 350, 690, 460),
        "org": (900, 350, 1500, 460),
        "ext": (1710, 350, 2310, 460),
    }
    box(root, "YÊU CẦU PHI CHỨC NĂNG", "#173B35", "#173B35", font_root, "white")
    box(groups["product"], "YÊU CẦU SẢN PHẨM", "#276255", "#276255", font_group, "white")
    box(groups["org"], "YÊU CẦU TỔ CHỨC", "#276255", "#276255", font_group, "white")
    box(groups["ext"], "YÊU CẦU BÊN NGOÀI", "#276255", "#276255", font_group, "white")

    line = "#71807B"
    root_cx = (root[0] + root[2]) // 2
    bus_y = 270
    draw.line((root_cx, root[3], root_cx, bus_y), fill=line, width=4)
    group_centers = [(rect[0] + rect[2]) // 2 for rect in groups.values()]
    draw.line((group_centers[0], bus_y, group_centers[-1], bus_y), fill=line, width=4)
    for rect in groups.values():
        cx = (rect[0] + rect[2]) // 2
        draw.line((cx, bus_y, cx, rect[1]), fill=line, width=4)

    children = {
        "product": ["Khả dụng", "Hiệu năng", "Độ tin cậy", "Linh động"],
        "org": ["Chuyển giao", "Cài đặt", "Tiêu chuẩn", "Vận hành"],
        "ext": ["Tương thích", "Hợp quy", "Pháp lý", "Đạo đức &\nan toàn"],
    }
    starts = {"product": 42, "org": 852, "ext": 1662}
    for group, labels in children.items():
        parent = groups[group]
        parent_cx = (parent[0] + parent[2]) // 2
        child_bus = 555
        draw.line((parent_cx, parent[3], parent_cx, child_bus), fill="#9AA8A4", width=3)
        rects = []
        for i, label in enumerate(labels):
            x1 = starts[group] + i * 165
            y1 = 650 if i % 2 == 0 else 835
            rect = (x1, y1, x1 + 145, y1 + 100)
            rects.append(rect)
        centers = [(r[0] + r[2]) // 2 for r in rects]
        draw.line((centers[0], child_bus, centers[-1], child_bus), fill="#9AA8A4", width=3)
        for rect, label in zip(rects, labels):
            cx = (rect[0] + rect[2]) // 2
            draw.line((cx, child_bus, cx, rect[1]), fill="#9AA8A4", width=3)
            box(rect, label, "white", "#99AAA5", font_child, "#1F2A28", radius=15)

    note = "Khung phân loại được điều chỉnh cho đồ án website thương mại điện tử NÉT"
    bbox = draw.textbbox((0, 0), note, font=font_note)
    draw.text(((2400 - (bbox[2] - bbox[0])) / 2, 1165), note, font=font_note, fill="#66706D")
    canvas.save(DIAGRAM, format="PNG", optimize=True)


def configure_styles(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(GREEN)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(16)

    for style_name, size, before, after, color in [
        ("Heading 1", 16, 18, 10, GREEN),
        ("Heading 2", 13, 14, 7, GREEN_2),
        ("Heading 3", 11.5, 10, 5, GREEN),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Requirement Bullet" not in styles:
        bullet_style = styles.add_style("Requirement Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet_style = styles["Requirement Bullet"]
    bullet_style.base_style = normal
    bullet_style.font.name = "Calibri"
    bullet_style.font.size = Pt(10.5)
    bullet_style.paragraph_format.space_after = Pt(4)
    bullet_style.paragraph_format.line_spacing = 1.18

    if "Small Note" not in styles:
        note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Small Note"]
    note.base_style = normal
    note.font.name = "Calibri"
    note.font.size = Pt(8.5)
    note.font.color.rgb = RGBColor.from_string(MUTED)
    note.paragraph_format.space_after = Pt(4)

    list_bullet = styles["List Bullet"]
    list_bullet.font.name = "Calibri"
    list_bullet.font.size = Pt(10.5)
    list_bullet.font.color.rgb = RGBColor.from_string(INK)
    list_bullet.paragraph_format.left_indent = Inches(0.375)
    list_bullet.paragraph_format.first_line_indent = Inches(-0.188)
    list_bullet.paragraph_format.space_after = Pt(4)
    list_bullet.paragraph_format.line_spacing = 1.18

    bullet_id = add_numbering(doc, "bullet", "•", font="Arial")
    decimal_id = add_numbering(doc, "decimal", "%1.")
    return bullet_id, decimal_id


def add_running_furniture(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "NÉT FASHION  ·  ĐẶC TẢ YÊU CẦU HỆ THỐNG"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("ĐỒ ÁN WEBSITE BÁN QUẦN ÁO TRỰC TUYẾN")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(GREEN_2)

    p = doc.add_paragraph(style="Title")
    p.add_run("TÀI LIỆU YÊU CẦU\nCHỨC NĂNG VÀ PHI CHỨC NĂNG")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Functional Requirements & Non-functional Requirements Specification")

    table = doc.add_table(rows=5, cols=2)
    data = [
        ("Hệ thống", "NÉT – Nền tảng thương mại điện tử thời trang"),
        ("Loại tài liệu", "Software Requirements Specification – Chuyên đề FR/NFR"),
        ("Phiên bản", "1.0 – Baseline phục vụ thiết kế, phát triển và kiểm thử"),
        ("Ngày phát hành", "01/09/2026"),
    ]
    set_cell_text(table.rows[0].cells[0], "Thông tin", bold=True, color=WHITE, size=9.5)
    set_cell_text(table.rows[0].cells[1], "Giá trị", bold=True, color=WHITE, size=9.5)
    set_cell_shading(table.rows[0].cells[0], GREEN)
    set_cell_shading(table.rows[0].cells[1], GREEN)
    set_repeat_header(table.rows[0])
    for idx, (key, value) in enumerate(data, 1):
        set_cell_text(table.rows[idx].cells[0], key, bold=True, color=GREEN, size=9.5)
        set_cell_text(table.rows[idx].cells[1], value, size=9.5)
        set_cell_shading(table.rows[idx].cells[0], BEIGE)
        if idx % 2:
            set_cell_shading(table.rows[idx].cells[1], BEIGE_2)
    set_table_geometry(table, [1.55, 4.95])

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_callout(doc, "Mục đích", "Định nghĩa đầy đủ những gì hệ thống phải thực hiện và mức chất lượng phải đạt; làm baseline cho thiết kế, coding, kiểm thử, UAT và nghiệm thu đồ án.", BEIGE_2)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(75)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NÉT / PRODUCT · BUSINESS ANALYSIS · ENGINEERING")
    r.font.name = "Calibri"
    r.font.size = Pt(9)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def add_control(doc: Document, bullet_id: int) -> None:
    doc.add_heading("Kiểm soát tài liệu", level=1)
    table = doc.add_table(rows=4, cols=4)
    headers = ["Phiên bản", "Ngày", "Nội dung", "Trạng thái"]
    rows = [
        ("0.1", "30/08/2026", "Khung yêu cầu ban đầu", "Draft"),
        ("0.5", "31/08/2026", "Bổ sung nghiệp vụ inventory và admin", "Review"),
        ("1.0", "01/09/2026", "Baseline FR/NFR đầy đủ cho đồ án", "Proposed baseline"),
    ]
    for idx, header in enumerate(headers): set_cell_text(table.rows[0].cells[idx], header)
    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row): set_cell_text(table.rows[r_idx].cells[c_idx], value, size=9)
    style_table(table, headers, [0.85, 0.9, 3.75, 1.0], 9)

    doc.add_heading("Vai trò phê duyệt", level=2)
    for text in [
        "Product Owner/PM: xác nhận phạm vi, ưu tiên và tiêu chí chấp nhận nghiệp vụ.",
        "Business Analyst: đảm bảo yêu cầu rõ ràng, nhất quán, có thể kiểm thử và truy vết.",
        "Frontend/Backend/DevOps: xác nhận tính khả thi kỹ thuật và kế hoạch kiểm chứng NFR.",
        "Giảng viên/người hướng dẫn: nhận xét mức độ đầy đủ và phù hợp với mục tiêu đồ án.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("Quy ước", level=2)
    p = doc.add_paragraph()
    p.add_run("M / Must").bold = True
    p.add_run(": bắt buộc trong baseline; ")
    p.add_run("S / Should").bold = True
    p.add_run(": quan trọng nhưng có thể xếp sau Must; ")
    p.add_run("C / Could").bold = True
    p.add_run(": thực hiện khi còn nguồn lực. “Hệ thống phải” biểu thị yêu cầu bắt buộc; tiêu chí nghiệm thu là điều kiện có thể quan sát/đo lường.")

    doc.add_heading("Mục lục nội dung", level=2)
    toc = [
        "1. Tổng quan và phạm vi",
        "2. Tác nhân, giả định và nguyên tắc đặc tả",
        "3. Yêu cầu chức năng",
        "4. Luồng nghiệp vụ trọng yếu",
        "5. Khung yêu cầu phi chức năng",
        "6. Danh mục yêu cầu phi chức năng định lượng",
        "7. Quality Attribute Scenarios",
        "8. Ma trận truy vết và kế hoạch nghiệm thu",
        "9. Phụ lục checklist",
    ]
    for item in toc:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(3)
    doc.add_page_break()


def add_overview(doc: Document, bullet_id: int) -> None:
    doc.add_heading("1. Tổng quan và phạm vi", level=1)
    doc.add_heading("1.1 Mục tiêu tài liệu", level=2)
    doc.add_paragraph(
        "Tài liệu này là baseline yêu cầu cho website bán quần áo trực tuyến NÉT. Nội dung tập trung vào hai lớp: yêu cầu chức năng mô tả hành vi hệ thống và yêu cầu phi chức năng mô tả mức chất lượng, ràng buộc, khả năng vận hành. Mỗi yêu cầu có mã duy nhất để liên kết với use case, dữ liệu, API và test case."
    )
    add_callout(doc, "Phạm vi baseline", f"{sum(len(x[2]) for x in FUNCTIONAL_GROUPS)} yêu cầu chức năng và {sum(len(x[2]) for x in NFR_GROUPS)} yêu cầu phi chức năng, bao phủ toàn bộ storefront, back-office, tồn kho, thanh toán, AI và vận hành.")

    doc.add_heading("1.2 Phạm vi trong đồ án", level=2)
    for text in [
        "Storefront responsive: trang chủ, danh mục, tìm kiếm, chi tiết sản phẩm, giỏ hàng và voucher.",
        "Checkout: guest/member, COD/thanh toán online, tạo đơn, giữ tồn và theo dõi trạng thái.",
        "Tài khoản: đăng ký, đăng nhập, hồ sơ, sổ địa chỉ và lịch sử đơn.",
        "Back-office: dashboard, catalog, khách hàng, đơn hàng, khuyến mãi, báo cáo và AI Copilot.",
        "Inventory đa kho: nhập, chuyển, kiểm kê, điều chỉnh, reservation, sổ cái và CSV.",
        "Nền tảng kỹ thuật: API, cơ sở dữ liệu, bảo mật, audit, outbox, Docker, CI/CD và quan sát vận hành.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("1.3 Ngoài phạm vi baseline", level=2)
    for text in [
        "Ứng dụng mobile native, marketplace đa nhà bán, đa tiền tệ và đa ngôn ngữ.",
        "POS/ERP toàn diện, tối ưu phân bổ đơn đa kho tự động, split shipment nâng cao.",
        "Loyalty/điểm thưởng, gift card, subscription và hoàn trả self-service đầy đủ.",
        "Lưu hoặc xử lý dữ liệu thẻ; phần nhập thẻ thuộc cổng thanh toán được chứng nhận.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("2. Tác nhân, giả định và nguyên tắc đặc tả", level=1)
    actors = [
        ("Khách vãng lai", "Duyệt sản phẩm, quản lý giỏ, áp voucher, guest checkout."),
        ("Khách thành viên", "Thực hiện quyền khách vãng lai và quản lý hồ sơ, địa chỉ, lịch sử đơn."),
        ("Nhân viên vận hành", "Xử lý đơn, xem khách hàng, nhập/chuyển/kiểm kê tồn."),
        ("Quản lý", "Quản lý catalog, điều chỉnh tồn, khuyến mãi, báo cáo và AI."),
        ("Quản trị viên", "Toàn quyền cấu hình, phân quyền và kiểm soát vận hành."),
        ("Cổng thanh toán", "Gửi kết quả giao dịch qua webhook có chữ ký."),
        ("DevOps/Support", "Triển khai, health check, giám sát, backup, khôi phục và xử lý sự cố."),
    ]
    table = doc.add_table(rows=len(actors)+1, cols=2)
    for i, h in enumerate(["Tác nhân", "Trách nhiệm/phạm vi tương tác"]): set_cell_text(table.rows[0].cells[i], h)
    for r_idx, row in enumerate(actors, 1):
        set_cell_text(table.rows[r_idx].cells[0], row[0], bold=True, color=GREEN, size=9)
        set_cell_text(table.rows[r_idx].cells[1], row[1], size=9)
    style_table(table, ["Tác nhân", "Trách nhiệm/phạm vi tương tác"], [1.55, 4.95], 9)

    doc.add_heading("2.1 Nguyên tắc chất lượng của yêu cầu", level=2)
    for text in [
        "Đơn nghĩa: một yêu cầu chỉ diễn đạt một hành vi hoặc một ràng buộc có thể kiểm chứng.",
        "Có nguồn sự thật: giá, khuyến mãi, quyền và tồn kho được xác nhận ở backend.",
        "Có thể truy vết: mã FR/NFR liên kết được với use case, bảng dữ liệu, API và test evidence.",
        "Có thể nghiệm thu: ưu tiên dùng điều kiện Given/When/Then hoặc ngưỡng định lượng.",
        "Nhất quán trạng thái: order, payment, fulfillment và inventory tuân theo state machine/quy tắc giao dịch.",
    ]:
        add_bullet(doc, text, bullet_id)


def add_functional_requirements(doc: Document) -> None:
    doc.add_heading("3. Yêu cầu chức năng", level=1)
    doc.add_paragraph(
        "Các yêu cầu dưới đây mô tả hành vi quan sát được của hệ thống. Cột “Nghiệm thu tóm tắt” là điều kiện tối thiểu; test case chi tiết phải mở rộng cả happy path, validation, authorization, concurrency và idempotency khi phù hợp."
    )
    headers = ["Mã", "Yêu cầu", "Actor", "Ưu tiên", "Nghiệm thu tóm tắt"]
    widths = [0.78, 3.0, 0.72, 0.5, 1.5]
    section_number = 1
    for _, title, requirements in FUNCTIONAL_GROUPS:
        doc.add_heading(f"3.{section_number} {title}", level=2)
        table = doc.add_table(rows=len(requirements)+1, cols=5)
        for i, h in enumerate(headers): set_cell_text(table.rows[0].cells[i], h)
        for row_idx, row in enumerate(requirements, 1):
            for col_idx, value in enumerate(row):
                align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(table.rows[row_idx].cells[col_idx], value, bold=(col_idx == 0), color=GREEN if col_idx == 0 else INK, size=8.5, align=align)
        style_table(table, headers, widths, 8.5)
        section_number += 1


def add_critical_flows(doc: Document, bullet_id: int, decimal_id: int) -> None:
    doc.add_heading("4. Luồng nghiệp vụ trọng yếu", level=1)
    doc.add_paragraph("Phần này làm rõ những yêu cầu có nhiều bước, side effect và rủi ro nhất. Đây là đầu vào trực tiếp cho sequence diagram, transaction boundary và integration test.")

    flows = [
        ("4.1 Đặt hàng và giữ tồn", "FR-CHK-001…004; FR-ORD-001…005", [
            "Khách nhập thông tin nhận hàng, chọn phương thức thanh toán và xác nhận đơn.",
            "Backend xác thực payload, idempotency key và đọc lại variant/giá/tồn.",
            "Backend khóa dòng inventory phù hợp; kiểm tra available và tính lại voucher/tổng tiền.",
            "Trong một giao dịch, tạo order, item snapshot, payment, reservation, history và outbox.",
            "Cập nhật reserved bằng điều kiện chống race; xóa cart items sau khi commit thành công.",
            "Trả mã đơn, thời hạn reservation và hành động tiếp theo cho frontend.",
        ], "Nếu bất kỳ dòng nào thiếu tồn hoặc giao dịch lỗi, không được để lại order/payment/reservation mồ côi. Gửi lặp cùng idempotency key phải trả kết quả trước đó."),
        ("4.2 Thanh toán trực tuyến", "FR-PAY-001…004", [
            "Khách được chuyển tới bước thanh toán cho payment đã tạo.",
            "Provider gửi webhook; hệ thống đọc raw body và xác minh HMAC/signature.",
            "Hệ thống kiểm tra event ID, payment ID, amount, currency và trạng thái hiện tại.",
            "Event hợp lệ được lưu một lần; payment chuyển trạng thái và order được confirm nếu phù hợp.",
            "Outbox/notification được xử lý tách biệt, lỗi gửi không rollback payment.",
        ], "Webhook sai chữ ký không thay đổi dữ liệu. Event trùng/out-of-order không tạo side effect trùng hoặc làm lùi trạng thái terminal."),
        ("4.3 Chuyển kho", "FR-INV-004…009", [
            "Nhân viên chọn dòng tồn kho nguồn, kho đích, số lượng, lý do và reference.",
            "Backend kiểm tra role, idempotency và khóa dòng nguồn.",
            "Nếu cần, tạo dòng inventory đích trước khi tạo operation item.",
            "Trong một giao dịch: trừ nguồn, cộng đích, tạo operation, operation item, hai movement và audit.",
            "Kết quả trả mã phiếu; gửi lặp trả phiếu cũ.",
        ], "Tổng on-hand toàn hệ thống không đổi qua một transfer. Kho nguồn không được thấp hơn reserved; source và destination phải khác nhau."),
        ("4.4 Chuyển trạng thái đơn", "FR-ORD-006…008; FR-ADM-007", [
            "Nhân viên mở chi tiết đơn và chọn trạng thái tiếp theo trong danh sách hợp lệ.",
            "Backend khóa order, kiểm tra transition và yêu cầu reason cho hủy/hoàn trả.",
            "Nếu hủy trước ship: release reservation. Nếu shipped: consume reservation và trừ on-hand.",
            "Cập nhật order/payment/fulfillment liên quan; ghi history và audit.",
            "Phát outbox event để gửi thông báo hoặc tích hợp tiếp theo.",
        ], "Trạng thái terminal không được chuyển lùi tùy ý. Mọi side effect inventory phải idempotent và khớp lịch sử đơn."),
    ]
    for heading, trace, steps, exception in flows:
        doc.add_heading(heading, level=2)
        p = doc.add_paragraph()
        p.add_run("Truy vết: ").bold = True
        p.add_run(trace)
        flow_decimal_id = add_numbering(doc, "decimal", "%1.")
        for step in steps:
            p = doc.add_paragraph(style="Requirement Bullet")
            apply_numbering(p, flow_decimal_id)
            p.add_run(step)
        add_callout(doc, "Điều kiện bất biến", exception, "EEF4F1")


def add_nfr_framework(doc: Document) -> None:
    doc.add_heading("5. Khung yêu cầu phi chức năng", level=1)
    doc.add_paragraph(
        "Yêu cầu phi chức năng không mô tả một màn hình riêng lẻ mà quy định chất lượng và ràng buộc xuyên suốt. Dựa trên sơ đồ tham chiếu, tài liệu chia NFR thành ba nhánh chính: yêu cầu sản phẩm, yêu cầu của tổ chức và yêu cầu bên ngoài; đồng thời bổ sung bảo mật, khả năng quan sát và AI có trách nhiệm phù hợp với hệ thống hiện đại."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    picture = run.add_picture(str(DIAGRAM), width=Inches(6.45))
    picture._inline.docPr.set("descr", "Cây phân loại yêu cầu phi chức năng gồm yêu cầu sản phẩm, yêu cầu tổ chức và yêu cầu bên ngoài, mỗi nhóm có bốn thuộc tính con.")
    picture._inline.docPr.set("title", "Cây phân loại yêu cầu phi chức năng")
    doc.add_paragraph("Hình 1. Cây phân loại yêu cầu phi chức năng áp dụng cho đồ án NÉT", style="Small Note").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_callout(doc, "Cách đọc", "Mỗi NFR phải có mục tiêu đo lường và phương pháp kiểm chứng. Những câu như “hệ thống nhanh”, “bảo mật tốt” hoặc “dễ dùng” không được xem là yêu cầu hoàn chỉnh nếu chưa có ngưỡng/evidence.")


def add_nfr_requirements(doc: Document) -> None:
    doc.add_heading("6. Danh mục yêu cầu phi chức năng định lượng", level=1)
    headers = ["Mã", "Thuộc tính", "Yêu cầu và mục tiêu", "Kiểm chứng", "Ưu tiên"]
    widths = [0.86, 0.92, 3.35, 0.92, 0.45]
    section_number = 1
    for _, title, requirements in NFR_GROUPS:
        doc.add_heading(f"6.{section_number} {title}", level=2)
        table = doc.add_table(rows=len(requirements)+1, cols=5)
        for idx, header in enumerate(headers): set_cell_text(table.rows[0].cells[idx], header)
        for row_idx, row in enumerate(requirements, 1):
            req_id, attribute, requirement, target, verification, priority = row
            combined = requirement + "\nMục tiêu: " + target
            values = (req_id, attribute, combined, verification, priority)
            for col_idx, value in enumerate(values):
                align = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 4 else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(table.rows[row_idx].cells[col_idx], value, bold=(col_idx == 0), color=GREEN if col_idx == 0 else INK, size=8.6, align=align)
        style_table(table, headers, widths, 8.6)
        section_number += 1


def add_quality_scenarios(doc: Document) -> None:
    doc.add_heading("7. Quality Attribute Scenarios", level=1)
    doc.add_paragraph("Quality scenario chuyển NFR thành tình huống kiểm thử có ngữ cảnh, kích thích, thành phần chịu tác động, phản ứng và thước đo. Mười scenario dưới đây là bộ tối thiểu cho kiểm thử kiến trúc và release readiness.")
    headers = ["ID / Tình huống", "Kích thích & môi trường", "Phản ứng mong đợi", "Thước đo"]
    widths = [1.2, 1.85, 2.35, 1.1]
    table = doc.add_table(rows=len(QUALITY_SCENARIOS)+1, cols=4)
    for idx, header in enumerate(headers): set_cell_text(table.rows[0].cells[idx], header)
    for row_idx, row in enumerate(QUALITY_SCENARIOS, 1):
        qid, name, stimulus, env, artifact, response, measure = row
        values = (f"{qid}\n{name}", f"{stimulus}\nMôi trường: {env}\nThành phần: {artifact}", response, measure)
        for col_idx, value in enumerate(values):
            set_cell_text(table.rows[row_idx].cells[col_idx], value, bold=(col_idx == 0), color=GREEN if col_idx == 0 else INK, size=8.5)
    style_table(table, headers, widths, 8.5)


def add_traceability(doc: Document, bullet_id: int) -> None:
    doc.add_heading("8. Ma trận truy vết và kế hoạch nghiệm thu", level=1)
    doc.add_heading("8.1 Ma trận truy vết cấp epic", level=2)
    headers = ["Epic", "FR", "Use case", "Dữ liệu chính", "API", "Test"]
    widths = [0.75, 0.95, 0.72, 1.55, 1.65, 0.88]
    table = doc.add_table(rows=len(TRACE_ROWS)+1, cols=6)
    for idx, header in enumerate(headers): set_cell_text(table.rows[0].cells[idx], header)
    for row_idx, row in enumerate(TRACE_ROWS, 1):
        for col_idx, value in enumerate(row):
            set_cell_text(table.rows[row_idx].cells[col_idx], value, bold=(col_idx == 0), color=GREEN if col_idx == 0 else INK, size=8.2)
    style_table(table, headers, widths, 8.2)

    doc.add_heading("8.2 Definition of Ready cho yêu cầu", level=2)
    for text in [
        "Có mã duy nhất, mô tả đơn nghĩa, actor và giá trị/nguyên nhân nghiệp vụ.",
        "Có priority, acceptance criteria, business rule liên quan và dependency đã nhận diện.",
        "UI/diagram/API/data contract đủ để team ước lượng mà không phải đoán hành vi cốt lõi.",
        "NFR ảnh hưởng trực tiếp đã được gắn và có cách đo/evidence.",
        "Không còn câu hỏi mở làm thay đổi đáng kể scope hoặc transaction boundary.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("8.3 Definition of Done và điều kiện nghiệm thu", level=2)
    for text in [
        "Code review, static analysis, unit/integration test và migration test đều pass.",
        "Acceptance criteria Must có evidence; negative authorization và validation được kiểm thử.",
        "Không còn lỗi Sev-1/Sev-2; Critical/High security issue đã xử lý hoặc có risk acceptance.",
        "NFR performance, backup/restore, monitoring, health check và rollback được kiểm chứng theo ngưỡng.",
        "Traceability được cập nhật từ FR/NFR tới test; PO/PM ký UAT cho phạm vi release.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("8.4 Evidence bắt buộc theo nhóm NFR", level=2)
    evidence = [
        ("Hiệu năng", "Kịch bản tải, dataset, biểu đồ p50/p95/p99, throughput, error rate."),
        ("Bảo mật", "SAST/SCA/container scan, authorization negative test, secret scan, DAST."),
        ("Độ tin cậy", "Concurrency/idempotency test, fault injection, backup restore, rollback drill."),
        ("Khả dụng", "A11y report, browser/device matrix, usability findings và trạng thái xử lý."),
        ("Vận hành", "Dashboard, alert test, runbook, health/readiness và log correlation evidence."),
    ]
    table = doc.add_table(rows=len(evidence)+1, cols=2)
    for i, h in enumerate(["Nhóm", "Evidence nghiệm thu"]): set_cell_text(table.rows[0].cells[i], h)
    for idx, row in enumerate(evidence, 1):
        set_cell_text(table.rows[idx].cells[0], row[0], bold=True, color=GREEN, size=9)
        set_cell_text(table.rows[idx].cells[1], row[1], size=9)
    style_table(table, ["Nhóm", "Evidence nghiệm thu"], [1.35, 5.15], 9)


def add_appendix(doc: Document, bullet_id: int) -> None:
    doc.add_heading("9. Phụ lục checklist", level=1)
    doc.add_heading("9.1 Checklist review yêu cầu chức năng", level=2)
    for text in [
        "[ ] Tất cả màn hình/flow có loading, empty, error, success và authorization state.",
        "[ ] Business rule về giá, voucher, tồn và trạng thái được thực thi ở backend.",
        "[ ] Mọi mutation quan trọng có idempotency hoặc chiến lược chống gửi lặp.",
        "[ ] Order item và promotion lưu snapshot đủ để lịch sử không bị thay đổi theo catalog.",
        "[ ] Inventory side effect khớp operation, movement, reservation và audit.",
        "[ ] Mọi API trả stable error code, validation detail phù hợp và correlation ID.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("9.2 Checklist review yêu cầu phi chức năng", level=2)
    for text in [
        "[ ] Mỗi NFR có ngưỡng đo, môi trường đo và owner cung cấp evidence.",
        "[ ] Ngưỡng performance gắn với tải/dataset cụ thể, không chỉ đo một request cục bộ.",
        "[ ] Security/privacy được review từ design, không chờ đến cuối dự án.",
        "[ ] Backup đã restore thử; alert đã kích hoạt thử; rollback đã diễn tập.",
        "[ ] Các ngoại lệ NFR có risk acceptance, thời hạn xử lý và người chịu trách nhiệm.",
        "[ ] AI chỉ đưa đề xuất, không có quyền mutation; dữ liệu prompt không chứa PII không cần thiết.",
    ]:
        add_bullet(doc, text, bullet_id)

    doc.add_heading("9.3 Thuật ngữ", level=2)
    terms = [
        ("ATP", "Available to Promise = on-hand – reserved – safety stock."),
        ("Idempotency", "Cùng một yêu cầu gửi lặp không tạo thêm side effect."),
        ("p95", "95% mẫu có thời gian phản hồi nhỏ hơn hoặc bằng giá trị công bố."),
        ("RPO/RTO", "Mức mất dữ liệu tối đa / thời gian khôi phục dịch vụ mục tiêu."),
        ("SLO", "Mục tiêu chất lượng dịch vụ có thể đo trong một khoảng thời gian."),
        ("Outbox", "Bảng sự kiện giao dịch giúp phát event đáng tin cậy sau commit."),
        ("Grounded AI", "Câu trả lời AI dựa trên dữ liệu/snapshot được cung cấp, không tự bịa nguồn."),
    ]
    table = doc.add_table(rows=len(terms)+1, cols=2)
    for i, h in enumerate(["Thuật ngữ", "Định nghĩa"]): set_cell_text(table.rows[0].cells[i], h)
    for idx, row in enumerate(terms, 1):
        set_cell_text(table.rows[idx].cells[0], row[0], bold=True, color=GREEN, size=9)
        set_cell_text(table.rows[idx].cells[1], row[1], size=9)
    style_table(table, ["Thuật ngữ", "Định nghĩa"], [1.35, 5.15], 9)

    doc.add_heading("9.4 Nguồn nội bộ", level=2)
    p = doc.add_paragraph(style="Small Note")
    p.add_run("Tài liệu được tổng hợp từ: ").bold = True
    p.add_run("Project Overview, BRD/SRS, PRD, Business Rules, Use Cases, ERD/Data Dictionary, mã nguồn hiện hành và sơ đồ phân loại NFR do người dùng cung cấp. Các ngưỡng NFR là baseline đề xuất cho đồ án và phải được hiệu chỉnh bằng kết quả kiểm thử thực tế trước go-live.")


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Tài liệu yêu cầu chức năng và phi chức năng – NÉT Fashion"
    props.subject = "Software Requirements Specification"
    props.author = "Nhóm đồ án NÉT"
    props.keywords = "Functional requirements, Non-functional requirements, SRS, E-commerce, Inventory"
    props.comments = "Baseline 1.0"


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    make_diagram()
    doc = Document()
    bullet_id, decimal_id = configure_styles(doc)
    add_running_furniture(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_control(doc, bullet_id)
    add_overview(doc, bullet_id)
    add_functional_requirements(doc)
    add_critical_flows(doc, bullet_id, decimal_id)
    add_nfr_framework(doc)
    add_nfr_requirements(doc)
    add_quality_scenarios(doc)
    add_traceability(doc, bullet_id)
    add_appendix(doc, bullet_id)
    doc.save(OUTPUT)
    print(OUTPUT.name)


if __name__ == "__main__":
    build()
