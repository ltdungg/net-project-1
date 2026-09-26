from __future__ import annotations

"""Tạo lại bộ tài liệu đồ án theo đúng phiên bản NET Commerce đang chạy.

Nguồn sự thật của nội dung là mã nguồn Java/React và các migration Flyway V1–V7.
Script tạo ba tài liệu Word cùng một bộ sơ đồ kỹ thuật, tránh để các bản báo cáo
độc lập bị lệch thuật ngữ hoặc trạng thái triển khai.
"""

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
IMG_DIR = ROOT / "artifacts" / "current-documentation-diagrams"
MAIN_DOCX = OUT_DIR / "Bao-cao-do-an-website-ban-quan-ao-truc-tuyen.docx"
REQUIREMENTS_DOCX = OUT_DIR / "Tai-lieu-yeu-cau-chuc-nang-va-phi-chuc-nang-NET.docx"
SUMMARY_DOCX = OUT_DIR / "Bao-cao-tong-hop-du-an-cua-hang-thoi-trang-truc-tuyen.docx"

OUT_DIR.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)

# narrative_proposal preset, với override có tên: A4 Academic Report.
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MID_BLUE = "5B9BD5"
INK = "243447"
MUTED = "667085"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F9"
GRID = "D0D5DD"
GREEN = "2E7D32"
PALE_GREEN = "E8F5E9"
GOLD = "9A6700"
PALE_GOLD = "FFF5D6"
RED = "9B1C1C"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"

FONT_REG = Path(r"C:\Windows\Fonts\calibri.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\calibrib.ttf")
FONT_ITALIC = Path(r"C:\Windows\Fonts\calibrii.ttf")


PROJECT_FACTS = [
    ("Tên đề tài", "Xây dựng website bán quần áo trực tuyến NÉT Commerce"),
    ("Loại sản phẩm", "Đồ án phần mềm có storefront, trang quản trị, API và hạ tầng Docker"),
    ("Thời gian", "08 tuần (02 tháng)"),
    ("Nhân sự", "05 người: PM, Frontend, Backend 1, Backend 2 và DevOps"),
    ("Phiên bản tài liệu", "5.0 — đồng bộ mã nguồn và Flyway V7 ngày 02/09/2026"),
    ("Quy mô dữ liệu demo", "22 sản phẩm, 58 biến thể, 03 mã voucher"),
    ("Cơ sở dữ liệu", "33 bảng nghiệp vụ và 01 view phân khúc khách hàng"),
]


FUNCTIONAL_REQUIREMENTS = [
    ("FR-AUT-01", "Đăng ký tài khoản bằng email, mật khẩu và họ tên; email phải duy nhất.", "Must", "Đã cài đặt"),
    ("FR-AUT-02", "Gửi và xác minh email bằng token một lần có thời hạn.", "Must", "Đã cài đặt"),
    ("FR-AUT-03", "Đăng nhập bằng email/mật khẩu và trả access token 15 phút.", "Must", "Đã cài đặt"),
    ("FR-AUT-04", "Xoay vòng refresh token opaque 30 ngày; token cũ bị vô hiệu hóa.", "Must", "Đã cài đặt"),
    ("FR-AUT-05", "Đăng xuất và thu hồi refresh token hiện hành.", "Must", "Đã cài đặt"),
    ("FR-AUT-06", "Quên/đặt lại mật khẩu bằng token một lần; không lộ email có tồn tại hay không.", "Must", "Đã cài đặt"),
    ("FR-ACC-01", "Xem và cập nhật hồ sơ cá nhân.", "Must", "Đã cài đặt"),
    ("FR-ACC-02", "Thêm, sửa, xóa và chọn địa chỉ giao hàng mặc định.", "Must", "Đã cài đặt"),
    ("FR-ACC-03", "Quản lý danh sách sản phẩm yêu thích.", "Should", "Đã cài đặt"),
    ("FR-CAT-01", "Hiển thị danh sách sản phẩm đang bán, phân trang và sắp xếp.", "Must", "Đã cài đặt"),
    ("FR-CAT-02", "Tìm kiếm, lọc theo danh mục, khoảng giá và thuộc tính biến thể.", "Must", "Đã cài đặt"),
    ("FR-CAT-03", "Hiển thị chi tiết sản phẩm, media, chất liệu, biến thể size/màu và giá.", "Must", "Đã cài đặt"),
    ("FR-CAT-04", "Chỉ cho chọn biến thể đang hoạt động và còn ATP.", "Must", "Đã cài đặt"),
    ("FR-CAR-01", "Tạo giỏ cho khách vãng lai hoặc người đã đăng nhập.", "Must", "Đã cài đặt"),
    ("FR-CAR-02", "Thêm SKU, đổi số lượng, xóa dòng và tính lại tổng ở server.", "Must", "Đã cài đặt"),
    ("FR-PRO-01", "Áp hoặc bỏ mã voucher; kiểm tra thời gian, ngưỡng và giới hạn sử dụng.", "Must", "Đã cài đặt"),
    ("FR-PRO-02", "Hỗ trợ giảm cố định, phần trăm có trần và miễn phí vận chuyển.", "Must", "Đã cài đặt"),
    ("FR-CHK-01", "Checkout revalidate SKU, giá, tồn, địa chỉ và voucher trong transaction.", "Must", "Đã cài đặt"),
    ("FR-CHK-02", "Tạo đơn idempotent, dòng hàng snapshot và lịch sử trạng thái.", "Must", "Đã cài đặt"),
    ("FR-CHK-03", "Tạo reservation theo kho và từ chối oversell khi ATP không đủ.", "Must", "Đã cài đặt"),
    ("FR-PAY-01", "Hỗ trợ COD và thanh toán online mô phỏng.", "Must", "Đã cài đặt"),
    ("FR-PAY-02", "Xử lý payment event có chữ ký và chống lặp theo provider_event_id.", "Must", "Một phần"),
    ("FR-ORD-01", "Khách thành viên xem danh sách và chi tiết đơn của chính mình.", "Must", "Đã cài đặt"),
    ("FR-ORD-02", "Khách vãng lai theo dõi đơn bằng token ký, TTL 07 ngày, gửi trong header.", "Must", "Một phần"),
    ("FR-ORD-03", "Nhân viên tra cứu đơn và chuyển trạng thái theo state machine.", "Must", "Đã cài đặt"),
    ("FR-INV-01", "Tra cứu tồn theo kho/SKU, hiển thị on-hand, reserved, ATP và ngưỡng đặt lại.", "Must", "Đã cài đặt"),
    ("FR-INV-02", "Nhập kho nhiều dòng với idempotency key và movement append-only.", "Must", "Đã cài đặt"),
    ("FR-INV-03", "Chuyển kho nguyên tử: trừ kho nguồn, cộng kho đích và ghi hai movement.", "Must", "Đã cài đặt"),
    ("FR-INV-04", "Kiểm kê và điều chỉnh tồn có lý do, version check và audit.", "Must", "Đã cài đặt"),
    ("FR-INV-05", "Giải phóng reservation hết hạn bằng tác vụ quản trị an toàn lặp lại.", "Should", "Một phần"),
    ("FR-INV-06", "Xuất tồn kho CSV UTF-8 để đối soát.", "Should", "Đã cài đặt"),
    ("FR-FUL-01", "Tự tạo một phiếu giao cho mỗi kho tham gia cấp hàng khi checkout.", "Must", "Đã cài đặt"),
    ("FR-FUL-02", "Hiển thị danh sách/KPI và chi tiết phiếu giao, dòng hàng, địa chỉ, lịch sử.", "Must", "Đã cài đặt"),
    ("FR-FUL-03", "Cập nhật hãng, dịch vụ, tracking, ETA, chi phí và ghi chú thủ công.", "Must", "Đã cài đặt"),
    ("FR-FUL-04", "Điều khiển pending → processing → ready_to_ship → shipped → delivered.", "Must", "Đã cài đặt"),
    ("FR-FUL-05", "Hỗ trợ failed, returned, cancelled; bắt buộc lý do tại các nhánh ngoại lệ.", "Must", "Đã cài đặt"),
    ("FR-FUL-06", "Khi shipped, chỉ consume reservation và xuất tồn thuộc đúng phiếu giao.", "Must", "Đã cài đặt"),
    ("FR-FUL-07", "COD chỉ ghi paid khi mọi phiếu giao đã delivered.", "Must", "Đã cài đặt"),
    ("FR-CUS-01", "Quản trị tìm kiếm, lọc trạng thái và xem hồ sơ 360 độ của khách hàng.", "Must", "Đã cài đặt"),
    ("FR-CUS-02", "Tự tính sáu phân khúc at_risk, vip, loyal, new, potential, prospect.", "Must", "Đã cài đặt"),
    ("FR-CUS-03", "Quản trị khóa/mở trạng thái tài khoản theo quyền.", "Must", "Đã cài đặt"),
    ("FR-ADM-01", "Quản lý sản phẩm, biến thể, media và trạng thái xuất bản.", "Must", "Đã cài đặt"),
    ("FR-ADM-02", "Quản lý chương trình khuyến mãi và coupon.", "Must", "Đã cài đặt"),
    ("FR-ADM-03", "Xem dashboard, báo cáo tổng quan và tài chính.", "Must", "Đã cài đặt"),
    ("FR-AI-01", "AI Copilot trả lời trên snapshot KPI và đưa hành động gợi ý, không tự ghi dữ liệu.", "Should", "Đã cài đặt dạng rule-based"),
    ("FR-AUD-01", "Ghi audit log và outbox event cho các thay đổi nghiệp vụ trọng yếu.", "Must", "Đã cài đặt phần lõi"),
]


NON_FUNCTIONAL_REQUIREMENTS = [
    ("NFR-PER-01", "Hiệu năng", "API đọc catalog/cart p95 ≤ 500 ms ở tải mục tiêu 100 request/giây.", "K6/JMeter + log latency"),
    ("NFR-PER-02", "Hiệu năng", "Trang quan trọng đạt LCP ≤ 2,5 giây trên thiết bị di động 4G mô phỏng.", "Lighthouse"),
    ("NFR-REL-01", "Tin cậy", "Checkout và inventory mutation phải nguyên tử; không sinh order mồ côi.", "Integration test PostgreSQL"),
    ("NFR-REL-02", "Tin cậy", "Idempotency áp dụng cho order, payment, movement, operation và fulfillment.", "Retry/duplicate test"),
    ("NFR-REL-03", "Tin cậy", "Không oversell: reserved ≥ 0 và reserved ≤ on_hand.", "Constraint + concurrency test"),
    ("NFR-AVA-01", "Khả dụng", "Mục tiêu MVP đạt 99,5%/tháng; health phản ánh API, PostgreSQL và Redis.", "Health/SLO report"),
    ("NFR-SEC-01", "Bảo mật", "Mật khẩu băm bằng bộ mã hóa Spring Security; không ghi mật khẩu/token thô vào log.", "Security review"),
    ("NFR-SEC-02", "Bảo mật", "RBAC được kiểm tra ở server cho customer, staff, manager và admin.", "Authorization test"),
    ("NFR-SEC-03", "Bảo mật", "Access token 15 phút; refresh token 30 ngày được hash, rotate và revoke.", "Auth lifecycle test"),
    ("NFR-SEC-04", "Bảo mật", "Guest order token ký, TTL 07 ngày, giữ ở sessionStorage và truyền X-Guest-Order-Token.", "Contract/security test"),
    ("NFR-SEC-05", "Bảo mật", "Request ghi phải qua kiểm tra Origin và validation; correlation ID đi xuyên request.", "Filter/integration test"),
    ("NFR-USA-01", "Khả dụng", "Giao diện responsive từ 360 px; trạng thái rỗng/lỗi/loading có hướng dẫn rõ.", "UAT đa thiết bị"),
    ("NFR-ACC-01", "Truy cập", "Luồng mua hàng đáp ứng WCAG 2.2 AA ở mức thiết kế: bàn phím, label, focus, contrast.", "Axe + manual test"),
    ("NFR-MAI-01", "Bảo trì", "Backend modular monolith theo domain; migration Flyway đã chạy là bất biến.", "Architecture review"),
    ("NFR-MAI-02", "Bảo trì", "Business rule trọng yếu có unit/integration test và stable error code.", "CI/test report"),
    ("NFR-COM-01", "Tương thích", "Hỗ trợ hai phiên bản gần nhất của Chrome, Edge, Firefox và Safari.", "Browser matrix"),
    ("NFR-DEP-01", "Triển khai", "Fresh clone chạy local bằng Docker Compose trong ≤ 10 phút sau khi cấu hình env.", "Onboarding test"),
    ("NFR-DAT-01", "Dữ liệu", "Thời gian lưu UTC; hiển thị Asia/Ho_Chi_Minh; tiền lưu số nguyên VND ở cột *_cents.", "Data contract test"),
    ("NFR-OBS-01", "Quan sát", "Log có cấu trúc, severity, correlation ID và không chứa PII dư thừa.", "Log audit"),
    ("NFR-OBS-02", "Quan sát", "Theo dõi error rate, latency, traffic, saturation và sự kiện order/payment/inventory.", "Dashboard vận hành"),
    ("NFR-BCK-01", "Khôi phục", "Mục tiêu RPO ≤ 24 giờ và RTO ≤ 4 giờ cho PostgreSQL ở phạm vi đồ án.", "Restore drill"),
    ("NFR-PRI-01", "Riêng tư", "PII chỉ hiển thị cho đúng chủ thể/nhân sự có quyền; có ma trận lưu trữ và xóa dữ liệu.", "Privacy checklist"),
    ("NFR-AI-01", "AI có trách nhiệm", "Copilot phải nêu thời điểm snapshot, không dùng thuộc tính nhạy cảm và không tự mutation.", "Prompt/evaluation test"),
]


USE_CASES = [
    ("UC-01", "Đăng ký và xác minh email", "Khách", "Tạo tài khoản hợp lệ và kích hoạt email", "Email chưa tồn tại", ["Nhập thông tin đăng ký", "Server validate và băm mật khẩu", "Tạo user + customer role", "Sinh auth token", "Xác minh token", "Đánh dấu email_verified_at"]),
    ("UC-02", "Đăng nhập và duy trì phiên", "Thành viên", "Nhận phiên truy cập an toàn", "Tài khoản active và email đã xác minh", ["Gửi email/mật khẩu", "Rate limit và kiểm tra Origin", "Xác thực mật khẩu", "Cấp access + refresh token", "Rotate refresh khi hết access", "Logout thu hồi refresh"]),
    ("UC-03", "Tìm và chọn sản phẩm", "Khách/Thành viên", "Chọn đúng SKU đang bán", "Catalog có dữ liệu", ["Tìm/lọc/sắp xếp", "Mở trang chi tiết", "Xem media và size guide", "Chọn size/màu", "Kiểm tra ATP", "Thêm vào giỏ hoặc yêu thích"]),
    ("UC-04", "Quản lý giỏ và voucher", "Khách/Thành viên", "Tạo giỏ có tổng tiền do server tính", "SKU active", ["Thêm SKU", "Điều chỉnh số lượng", "Áp mã voucher", "Server kiểm tra điều kiện", "Trả subtotal/discount/shipping/grand total", "Có thể bỏ voucher"]),
    ("UC-05", "Checkout đa kho", "Khách/Thành viên", "Tạo đơn, giữ tồn và phiếu giao nhất quán", "Giỏ hợp lệ và có địa chỉ", ["Gửi cart/address/payment/idempotency key", "Khóa và kiểm tra ATP", "Tạo order + item snapshot", "Tạo reservation theo kho", "Nhóm reservation theo warehouse", "Tạo fulfillment + fulfillment items", "Tạo payment và outbox", "Commit rồi trả mã đơn"]),
    ("UC-06", "Thanh toán và theo dõi đơn", "Khách/Thành viên", "Xác nhận thanh toán/trạng thái mà không lộ dữ liệu", "Đơn đã tạo", ["COD hoặc online mock", "Xử lý payment event idempotent", "Member đọc bằng JWT", "Guest gửi signed token trong header", "Hiển thị timeline đơn và từng kiện"]),
    ("UC-07", "Quản lý tồn kho", "Staff/Manager", "Biến động tồn có chứng từ và sổ cái", "Có quyền kho", ["Tra cứu ATP", "Nhập kho/chuyển kho/kiểm kê/điều chỉnh", "Validate kho và SKU", "Kiểm tra version/idempotency", "Post operation", "Cập nhật balance và append movement", "Ghi audit/outbox"]),
    ("UC-08", "Điều phối vận chuyển", "Order Operator", "Đưa từng kiện từ chờ xử lý đến hoàn tất", "Đơn paid hoặc cod_pending", ["Mở danh sách phiếu giao", "Chuyển processing", "Pick/pack và ready_to_ship", "Nhập carrier/tracking/ETA", "Chuyển shipped để consume reservation", "Cập nhật delivered/failed/returned", "Tổng hợp trạng thái đơn và COD"]),
    ("UC-09", "Quản lý khách hàng", "Staff/Manager/Admin", "Tìm kiếm, xem 360 độ và kiểm soát trạng thái", "Có quyền backoffice", ["Tìm/lọc customer", "Xem KPI và lịch sử đơn", "Đọc segment động", "Khóa/mở tài khoản theo policy", "Ghi audit"]),
    ("UC-10", "Quản lý catalog", "Staff/Manager/Admin", "Tạo sản phẩm và media đúng chuẩn", "Có quyền catalog", ["Nhập thông tin sản phẩm", "Tạo biến thể SKU", "Upload media", "Kiểm tra slug/SKU duy nhất", "Lưu draft hoặc publish"]),
    ("UC-11", "Quản lý khuyến mãi", "Manager/Admin", "Tạo promotion/coupon trong thời gian hợp lệ", "Có quyền promotion", ["Chọn loại giảm", "Đặt giá trị/ngưỡng/trần", "Thiết lập thời gian và giới hạn", "Tạo coupon", "Kích hoạt", "Theo dõi sử dụng"]),
    ("UC-12", "Sử dụng AI Copilot", "Manager/Admin", "Nhận tóm tắt và gợi ý có căn cứ", "Có snapshot metrics", ["Chọn câu hỏi quản trị", "Server lấy snapshot order/inventory/promotion", "Rule engine tạo insight", "Trả nguồn và thời điểm", "Người quản trị tự quyết định hành động"]),
]


TABLE_CATALOG = [
    ("Identity", "roles", "Vai trò RBAC"), ("Identity", "users", "Tài khoản và hồ sơ"),
    ("Identity", "user_roles", "Gán vai trò"), ("Identity", "refresh_tokens", "Refresh token đã băm"),
    ("Identity", "auth_tokens", "Token xác minh/đặt lại"), ("Identity", "addresses", "Sổ địa chỉ"),
    ("Catalog", "categories", "Cây danh mục"), ("Catalog", "products", "Sản phẩm"),
    ("Catalog", "product_categories", "Gắn sản phẩm–danh mục"), ("Catalog", "product_variants", "SKU size/màu"),
    ("Catalog", "product_media", "Ảnh/media"), ("Catalog", "wishlist_items", "Sản phẩm yêu thích"),
    ("Cart/Promotion", "carts", "Giỏ guest/member"), ("Cart/Promotion", "cart_items", "Dòng giỏ"),
    ("Cart/Promotion", "promotions", "Chương trình giảm giá"), ("Cart/Promotion", "coupons", "Mã voucher"),
    ("Order", "customer_orders", "Đơn hàng"), ("Order", "order_items", "Snapshot dòng hàng"),
    ("Order", "order_status_history", "Lịch sử đơn"), ("Order", "order_promotions", "Khuyến mãi trên đơn"),
    ("Payment", "payments", "Giao dịch thanh toán"), ("Payment", "payment_events", "Sự kiện provider"),
    ("Inventory", "warehouses", "Kho"), ("Inventory", "inventory_items", "Balance kho/SKU"),
    ("Inventory", "inventory_movements", "Sổ cái biến động"), ("Inventory", "inventory_operations", "Chứng từ kho"),
    ("Inventory", "inventory_operation_items", "Dòng chứng từ kho"), ("Inventory", "inventory_reservations", "Giữ tồn cho đơn"),
    ("Fulfillment", "fulfillments", "Phiếu giao theo kho"), ("Fulfillment", "fulfillment_items", "Dòng hàng trong kiện"),
    ("Fulfillment", "fulfillment_events", "Lịch sử kiện"), ("Governance", "audit_logs", "Nhật ký thao tác"),
    ("Governance", "outbox_events", "Sự kiện tích hợp"),
]


API_CATALOG = [
    ("Auth", "POST", "/api/auth/register|verify-email|login|refresh|logout", "Public/Authenticated", "Vòng đời tài khoản và phiên"),
    ("Account", "GET/POST/PATCH/DELETE", "/api/account/profile|addresses|wishlist", "Customer", "Hồ sơ, địa chỉ, yêu thích"),
    ("Catalog", "GET", "/api/products; /api/products/{slug}", "Public", "Danh sách và chi tiết"),
    ("Cart", "GET/POST/PATCH/DELETE", "/api/cart; /api/cart/items/{id}; /api/vouchers", "Guest/Customer", "Giỏ và voucher"),
    ("Order", "GET/POST", "/api/orders; /api/orders/{id}; /api/orders/{id}/guest", "Guest/Customer", "Đặt và theo dõi đơn"),
    ("Payment", "POST", "/api/payments/mock/confirm; /api/payments/webhook", "Customer/Provider", "Xác nhận thanh toán"),
    ("Admin Order", "GET/PATCH", "/api/admin/orders; /api/admin/orders/{id}/status", "Staff+", "Điều hành đơn"),
    ("Admin Customer", "GET/PATCH", "/api/admin/customers; /segments; /{id}/status", "Staff+", "Customer 360 và segment"),
    ("Admin Catalog", "GET/POST", "/api/admin/products; /options; /{id}/media; /uploads", "Staff+", "Quản lý catalog/media"),
    ("Inventory", "GET/POST", "/api/admin/inventory; /receipts; /transfers; /counts; /adjust; /export", "Staff+", "Tác nghiệp đa kho"),
    ("Fulfillment", "GET/PATCH", "/api/admin/fulfillments; /{id}/tracking; /{id}/status", "Staff+", "Điều phối vận chuyển"),
    ("Insights", "GET/POST", "/api/admin/metrics; /reports/overview; /finance; /ai; /promotions", "Manager+", "KPI, tài chính, Copilot"),
    ("Health", "GET", "/api/health", "Public", "Readiness ứng dụng"),
]


TIMELINE = [
    ("Tuần 1", "Khởi động, khảo sát mô phỏng, charter, vision, kiến trúc nền", "PM/BE1/DevOps"),
    ("Tuần 2", "Auth, user, catalog schema/API; shell storefront và design system", "FE/BE1/DevOps"),
    ("Tuần 3", "Catalog UI, media, wishlist, profile/address", "FE/BE1"),
    ("Tuần 4", "Cart, voucher, pricing và contract checkout", "FE/BE1/BE2"),
    ("Tuần 5", "Order/payment, reservation, inventory operations", "BE1/BE2/FE"),
    ("Tuần 6", "Customer segments, fulfillment đa kho, admin screens", "BE2/FE/PM"),
    ("Tuần 7", "Báo cáo, AI Copilot, hardening bảo mật, integration test", "Cả nhóm"),
    ("Tuần 8", "UAT, Docker smoke, sửa lỗi, hoàn thiện báo cáo và bàn giao", "Cả nhóm"),
]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def img_font(size: int, bold: bool = False, italic: bool = False):
    path = FONT_BOLD if bold else FONT_ITALIC if italic else FONT_REG
    return ImageFont.truetype(str(path), size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=font)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def draw_center(draw, box, text, font, fill=INK, gap=5):
    x1, y1, x2, y2 = box
    lines = wrap_text(draw, text, font, int(x2 - x1 - 24))
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + gap * (len(lines) - 1)
    y = (y1 + y2 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        x = (x1 + x2 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=f"#{fill}" if len(fill) == 6 else fill)
        y += height + gap


def diagram_canvas(title: str, subtitle: str = "", size=(2600, 1500)):
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.text((70, 40), title, font=img_font(48, True), fill=f"#{NAVY}")
    if subtitle:
        draw.text((72, 103), subtitle, font=img_font(25), fill=f"#{MUTED}")
    draw.line((70, 150, size[0] - 70, 150), fill=f"#{GRID}", width=3)
    return image, draw


def box(draw, rect, title, subtitle="", fill=LIGHT_GRAY, outline=BLUE, title_size=29, subtitle_size=22, radius=18):
    draw.rounded_rectangle(rect, radius=radius, fill=f"#{fill}", outline=f"#{outline}", width=4)
    x1, y1, x2, y2 = rect
    if subtitle:
        draw_center(draw, (x1 + 10, y1 + 14, x2 - 10, y1 + (y2 - y1) * 0.58), title, img_font(title_size, True), NAVY)
        draw_center(draw, (x1 + 10, y1 + (y2 - y1) * 0.52, x2 - 10, y2 - 8), subtitle, img_font(subtitle_size), MUTED)
    else:
        draw_center(draw, rect, title, img_font(title_size, True), NAVY)


def arrow(draw, start, end, color=MUTED, width=4, head=14):
    draw.line((start, end), fill=f"#{color}", width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    p1 = (end[0] - head * math.cos(angle - math.pi / 6), end[1] - head * math.sin(angle - math.pi / 6))
    p2 = (end[0] - head * math.cos(angle + math.pi / 6), end[1] - head * math.sin(angle + math.pi / 6))
    draw.polygon((end, p1, p2), fill=f"#{color}")


def ortho(draw, points, color=MUTED, width=4, head=14):
    draw.line(points, fill=f"#{color}", width=width, joint="curve")
    if len(points) >= 2:
        arrow(draw, points[-2], points[-1], color=color, width=width, head=head)


def save_diagram(image: Image.Image, name: str) -> Path:
    path = IMG_DIR / name
    image.save(path, dpi=(220, 220), optimize=True)
    return path


def make_context_diagram() -> Path:
    image, draw = diagram_canvas("Bối cảnh hệ thống NÉT Commerce", "Phạm vi đang triển khai và các chủ thể tương tác")
    box(draw, (950, 505, 1650, 930), "NÉT Commerce", "Storefront + Admin + BFF + Java API", PALE_BLUE, DARK_BLUE, 43, 27)
    actors = [
        ((100, 260, 680, 460), "Khách vãng lai", "Duyệt hàng, giỏ, checkout, tra cứu đơn"),
        ((100, 560, 680, 760), "Khách thành viên", "Hồ sơ, địa chỉ, yêu thích, lịch sử đơn"),
        ((100, 860, 680, 1060), "Nhân viên quản trị", "Catalog, đơn, kho, khách hàng, vận chuyển"),
        ((100, 1160, 680, 1360), "Manager / Admin", "Báo cáo, tài chính, khuyến mãi, AI Copilot"),
    ]
    for rect, title, subtitle in actors:
        box(draw, rect, title, subtitle, LIGHT_GRAY, BLUE, 30, 21)
        arrow(draw, (rect[2], (rect[1] + rect[3]) // 2), (950, 610 + (rect[1] - 260) * 0.2), BLUE)
    externals = [
        ((1920, 300, 2500, 500), "PostgreSQL 17", "33 bảng + 1 view; Flyway V1–V7"),
        ((1920, 610, 2500, 810), "Redis 7.4", "Rate limit và dữ liệu tạm"),
        ((1920, 920, 2500, 1120), "Upload volume", "Media sản phẩm lưu cục bộ"),
        ((1920, 1230, 2500, 1430), "Tích hợp ngoài", "Payment mock; vận chuyển nhập thủ công"),
    ]
    for rect, title, subtitle in externals:
        box(draw, rect, title, subtitle, PALE_GREEN if title != "Tích hợp ngoài" else PALE_GOLD, GREEN if title != "Tích hợp ngoài" else GOLD, 30, 21)
        arrow(draw, (1650, 620 + (rect[1] - 300) * 0.2), (rect[0], (rect[1] + rect[3]) // 2), GREEN if title != "Tích hợp ngoài" else GOLD)
    return save_diagram(image, "01-system-context.png")


def make_architecture_diagram() -> Path:
    image, draw = diagram_canvas("Kiến trúc triển khai hiện tại", "Modular monolith Java phía sau BFF; dữ liệu quan hệ trên PostgreSQL")
    layers = [
        (220, "TRẢI NGHIỆM", ["React 19", "Vinext routes", "Tailwind / Base UI", "Storefront + Admin"]),
        (480, "BFF / BIÊN", ["/api/commerce/**", "Cookie / token forwarding", "Origin normalization", "Error envelope"]),
        (740, "JAVA API", ["Spring Boot 4.1.1", "Spring Security", "Domain services", "JDBC + transaction"]),
        (1000, "DỮ LIỆU", ["PostgreSQL 17", "Redis 7.4", "Flyway V1–V7", "Upload volume"]),
        (1260, "VẬN HÀNH", ["Docker Compose", "Health checks", "Structured logs", "Smoke tests"]),
    ]
    for y, label, items in layers:
        draw.rounded_rectangle((90, y, 2510, y + 180), radius=18, fill=f"#{LIGHT_GRAY}", outline=f"#{GRID}", width=3)
        draw.rounded_rectangle((90, y, 430, y + 180), radius=18, fill=f"#{NAVY}")
        draw_center(draw, (100, y + 10, 420, y + 170), label, img_font(27, True), WHITE)
        x = 500
        for item in items:
            box(draw, (x, y + 30, x + 450, y + 150), item, "", PALE_BLUE, MID_BLUE, 25)
            x += 500
    for y in [400, 660, 920, 1180]:
        arrow(draw, (1300, y), (1300, y + 80), DARK_BLUE, 6, 18)
    return save_diagram(image, "02-deployment-architecture.png")


def make_use_case_diagram(title: str, filename: str, actors: list[tuple[str, str]], groups: list[tuple[str, list[str]]]) -> Path:
    image, draw = diagram_canvas(title, "Sơ đồ phân miền để các đường kết nối không chồng chéo", size=(2800, 1650))
    draw.rounded_rectangle((620, 205, 2680, 1540), radius=20, fill="#FBFCFE", outline=f"#{DARK_BLUE}", width=4)
    draw.text((660, 225), "Biên hệ thống NÉT Commerce", font=img_font(29, True), fill=f"#{DARK_BLUE}")
    actor_y = 330
    actor_ports = []
    for actor, detail in actors:
        box(draw, (70, actor_y, 500, actor_y + 180), actor, detail, PALE_GOLD, GOLD, 28, 20)
        actor_ports.append(actor_y + 90)
        actor_y += 310
    group_height = 1180 // max(1, len(groups))
    for gi, (group, cases) in enumerate(groups):
        y1 = 300 + gi * group_height
        y2 = y1 + group_height - 25
        draw.rounded_rectangle((720, y1, 2560, y2), radius=16, fill=f"#{LIGHT_GRAY}", outline=f"#{GRID}", width=3)
        draw.text((755, y1 + 18), group, font=img_font(27, True), fill=f"#{NAVY}")
        cols = 2 if len(cases) > 3 else 1
        for ci, case in enumerate(cases):
            col = ci % cols
            row = ci // cols
            width = 780 if cols == 2 else 1620
            x1 = 820 + col * 860
            y = y1 + 78 + row * 112
            box(draw, (x1, y, x1 + width, y + 82), case, "", PALE_BLUE, BLUE, 22, radius=38)
        # Mỗi nhóm chỉ có một đường actor→group; chi tiết liên kết được thể hiện trong nhóm.
        actor_index = min(gi, len(actor_ports) - 1)
        ortho(draw, [(500, actor_ports[actor_index]), (590, actor_ports[actor_index]), (590, (y1 + y2) // 2), (720, (y1 + y2) // 2)], BLUE, 4, 14)
    return save_diagram(image, filename)


def make_sequence_diagram(title: str, filename: str, participants: list[str], messages: list[tuple[int, int, str, str]]) -> Path:
    width = 3000
    height = 360 + 115 * len(messages)
    image, draw = diagram_canvas(title, "Mỗi thông điệp ở một hàng riêng; đường trả về dùng nét đứt", size=(width, height))
    left = 180
    spacing = (width - 2 * left) / (len(participants) - 1)
    xs = [int(left + i * spacing) for i in range(len(participants))]
    for x, participant in zip(xs, participants):
        box(draw, (x - 170, 190, x + 170, 300), participant, "", PALE_BLUE, BLUE, 23)
        draw.line((x, 300, x, height - 100), fill=f"#{GRID}", width=3)
    y = 370
    for number, (source, target, text_value, kind) in enumerate(messages, start=1):
        x1, x2 = xs[source], xs[target]
        color = GREEN if kind == "return" else BLUE
        if kind == "return":
            dash = 18
            step = 30 if x2 > x1 else -30
            pos = x1
            while (pos < x2 if step > 0 else pos > x2):
                end = min(pos + dash, x2) if step > 0 else max(pos - dash, x2)
                draw.line((pos, y, end, y), fill=f"#{color}", width=3)
                pos += step
            arrow(draw, (x1, y), (x2, y), color, 1, 13)
        else:
            arrow(draw, (x1, y), (x2, y), color, 4, 14)
        label_box = (min(x1, x2) + 18, y - 48, max(x1, x2) - 18, y - 5)
        draw_center(draw, label_box, f"{number}. {text_value}", img_font(20), INK)
        y += 115
    return save_diagram(image, filename)


def draw_entity(draw, rect, name: str, fields: list[str], fill=LIGHT_GRAY):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=14, fill=f"#{fill}", outline=f"#{BLUE}", width=3)
    draw.rectangle((x1, y1, x2, y1 + 60), fill=f"#{NAVY}")
    draw_center(draw, (x1 + 5, y1 + 3, x2 - 5, y1 + 57), name, img_font(23, True), WHITE)
    y = y1 + 75
    for field in fields:
        draw.text((x1 + 18, y), field, font=img_font(18), fill=f"#{INK}")
        y += 32


def relation(draw, points, label: str, label_xy: tuple[int, int], color=MUTED):
    ortho(draw, points, color=color, width=3, head=12)
    draw.rounded_rectangle((label_xy[0] - 50, label_xy[1] - 20, label_xy[0] + 50, label_xy[1] + 20), radius=8, fill="white", outline=f"#{GRID}")
    draw_center(draw, (label_xy[0] - 48, label_xy[1] - 18, label_xy[0] + 48, label_xy[1] + 18), label, img_font(17, True), MUTED)


def make_erd_identity() -> Path:
    image, draw = diagram_canvas("ERD 1 — Định danh và tài khoản", "PK/FK rút gọn; chi tiết đầy đủ nằm trong từ điển dữ liệu")
    entities = {
        "users": ((1010, 500, 1590, 930), ["PK id varchar(64)", "email, password_hash", "status, email_verified_at", "locked_until, profile", "created_at, updated_at"]),
        "roles": ((100, 270, 580, 520), ["PK id", "UQ code", "name, description"]),
        "user_roles": ((670, 270, 1160, 520), ["PK/FK user_id", "PK/FK role_id", "assigned_at"]),
        "refresh_tokens": ((1700, 240, 2250, 520), ["PK id; FK user_id", "UQ token_hash", "expires_at, revoked_at"]),
        "auth_tokens": ((1700, 610, 2250, 920), ["PK id; FK user_id", "type, UQ token_hash", "expires_at, consumed_at"]),
        "addresses": ((200, 700, 760, 1080), ["PK id; FK user_id", "recipient, phone", "line1, district, city", "is_default"]),
        "wishlist_items": ((820, 1100, 1390, 1370), ["PK/FK user_id", "PK/FK product_id", "created_at"]),
        "products(ref)": ((1700, 1110, 2250, 1360), ["PK id", "slug, name", "status"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(1010, 600), (900, 600), (900, 520)], "1:N", (900, 560))
    relation(draw, [(670, 390), (580, 390)], "N:1", (625, 350))
    relation(draw, [(1590, 590), (1640, 590), (1640, 380), (1700, 380)], "1:N", (1640, 450))
    relation(draw, [(1590, 730), (1700, 730)], "1:N", (1645, 690))
    relation(draw, [(1010, 800), (760, 800)], "1:N", (880, 760))
    relation(draw, [(1250, 930), (1250, 1100)], "1:N", (1295, 1015))
    relation(draw, [(1390, 1235), (1700, 1235)], "N:1", (1545, 1195))
    return save_diagram(image, "10-erd-identity.png")


def make_erd_catalog() -> Path:
    image, draw = diagram_canvas("ERD 2 — Catalog và media", "Quan hệ được bố trí theo luồng danh mục → sản phẩm → SKU → tồn")
    entities = {
        "categories": ((100, 300, 650, 650), ["PK id; FK parent_id", "UQ slug", "name, sort_order", "status"]),
        "product_categories": ((760, 300, 1310, 590), ["PK/FK product_id", "PK/FK category_id", "is_primary, sort_order"]),
        "products": ((1430, 300, 2000, 680), ["PK id; UQ slug", "name, description", "material, featured", "status, published_at"]),
        "product_variants": ((1430, 840, 2000, 1260), ["PK id; FK product_id", "UQ sku", "size, color", "price_cents, cost_cents", "status"]),
        "product_media": ((760, 910, 1310, 1260), ["PK id; FK product_id", "FK variant_id", "storage_key, public_url", "alt_text, sort_order"]),
        "inventory_items(ref)": ((2130, 870, 2660, 1240), ["PK id", "FK warehouse_id", "FK variant_id", "on_hand, reserved"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(650, 470), (760, 470)], "1:N", (705, 430))
    relation(draw, [(1310, 445), (1430, 445)], "N:1", (1370, 405))
    relation(draw, [(1715, 680), (1715, 840)], "1:N", (1760, 760))
    relation(draw, [(1430, 1010), (1310, 1010)], "1:N", (1370, 970))
    relation(draw, [(2000, 1050), (2130, 1050)], "1:N", (2065, 1010))
    relation(draw, [(375, 300), (375, 225), (520, 225), (520, 300)], "0:N", (450, 205))
    return save_diagram(image, "11-erd-catalog.png")


def make_erd_cart_promotion() -> Path:
    image, draw = diagram_canvas("ERD 3 — Giỏ hàng và khuyến mãi", "Không nối chéo: mỗi bảng phụ thuộc được đặt cạnh bảng cha")
    entities = {
        "users(ref)": ((100, 300, 600, 590), ["PK id", "email, status"]),
        "carts": ((780, 270, 1370, 650), ["PK id; FK user_id", "UQ guest_token", "coupon_code", "expires_at"]),
        "cart_items": ((1570, 270, 2170, 650), ["PK id; FK cart_id", "FK variant_id", "quantity", "unit_price_cents"]),
        "product_variants(ref)": ((2250, 300, 2750, 610), ["PK id; UQ sku", "price_cents", "status"]),
        "promotions": ((280, 930, 900, 1330), ["PK id", "type, value", "min/max amount", "starts_at, ends_at", "status, rules_json"]),
        "coupons": ((1080, 930, 1680, 1300), ["PK id; FK promotion_id", "UQ code", "usage_limit/count", "status"]),
        "order_promotions(ref)": ((1880, 930, 2520, 1300), ["PK id; FK order_id", "FK promotion_id/coupon_id", "name_snapshot", "discount_cents"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(600, 445), (780, 445)], "1:N", (690, 405))
    relation(draw, [(1370, 445), (1570, 445)], "1:N", (1470, 405))
    relation(draw, [(2170, 445), (2250, 445)], "N:1", (2210, 405))
    relation(draw, [(900, 1120), (1080, 1120)], "1:N", (990, 1080))
    relation(draw, [(1680, 1120), (1880, 1120)], "1:N", (1780, 1080))
    return save_diagram(image, "12-erd-cart-promotion.png")


def make_erd_order_payment() -> Path:
    image, draw = diagram_canvas("ERD 4 — Đơn hàng và thanh toán", "Order là aggregate trung tâm; history và payment event là append-only")
    entities = {
        "customer_orders": ((960, 430, 1660, 960), ["PK id; UQ order_number", "FK user_id", "status/payment/fulfillment", "*_total_cents", "shipping_address_json", "UQ idempotency_key"]),
        "order_items": ((100, 300, 700, 690), ["PK id; FK order_id", "FK variant_id", "SKU/name snapshot", "quantity, price, total"]),
        "order_status_history": ((100, 900, 700, 1240), ["PK id; FK order_id", "from_status, to_status", "reason, actor_user_id", "metadata_json"]),
        "payments": ((2050, 300, 2650, 730), ["PK id; FK order_id", "provider, method", "status, amount_cents", "UQ idempotency_key", "captured_at"]),
        "payment_events": ((2050, 930, 2650, 1280), ["PK id; FK payment_id", "provider/event_id UQ", "signature_valid", "payload_json"]),
        "inventory_reservations(ref)": ((760, 1100, 1380, 1420), ["PK id; FK order_id", "FK inventory_item_id", "quantity, status", "expires_at"]),
        "order_promotions(ref)": ((1450, 1100, 1990, 1410), ["PK id; FK order_id", "FK promotion/coupon", "discount_cents"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(960, 610), (700, 610)], "1:N", (830, 570))
    relation(draw, [(960, 800), (820, 800), (820, 1070), (700, 1070)], "1:N", (820, 940))
    relation(draw, [(1660, 600), (2050, 600)], "1:N", (1855, 560))
    relation(draw, [(2350, 730), (2350, 930)], "1:N", (2395, 830))
    relation(draw, [(1120, 960), (1120, 1100)], "1:N", (1165, 1030))
    relation(draw, [(1510, 960), (1510, 1100)], "1:N", (1555, 1030))
    return save_diagram(image, "13-erd-order-payment.png")


def make_erd_inventory() -> Path:
    image, draw = diagram_canvas("ERD 5 — Inventory đa kho", "Balance, chứng từ và sổ cái được tách để bảo toàn khả năng đối soát")
    entities = {
        "warehouses": ((100, 500, 620, 810), ["PK id; UQ code", "name, address", "status"]),
        "inventory_items": ((890, 480, 1510, 860), ["PK id; FK warehouse/variant", "UQ warehouse+variant", "on_hand, reserved", "safety_stock, version"]),
        "inventory_movements": ((1870, 260, 2490, 640), ["PK id; FK inventory_item", "type, quantity_delta", "balance_after", "UQ idempotency_key"]),
        "inventory_reservations": ((1870, 850, 2490, 1220), ["PK id; FK inventory_item", "FK order_id", "quantity, status", "expires_at"]),
        "inventory_operations": ((100, 1040, 740, 1420), ["PK id; UQ operation_number", "type, status", "FK source/destination warehouse", "UQ idempotency_key"]),
        "inventory_operation_items": ((880, 1060, 1570, 1450), ["PK id; FK operation_id", "FK variant_id", "FK source/destination item", "quantity/expected/actual"]),
        "product_variants(ref)": ((920, 190, 1490, 390), ["PK id; UQ sku", "size, color"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(620, 650), (890, 650)], "1:N", (755, 610))
    relation(draw, [(1205, 390), (1205, 480)], "1:N", (1250, 435))
    relation(draw, [(1510, 600), (1700, 600), (1700, 450), (1870, 450)], "1:N", (1700, 520))
    relation(draw, [(1510, 760), (1700, 760), (1700, 1035), (1870, 1035)], "1:N", (1700, 900))
    relation(draw, [(740, 1230), (880, 1230)], "1:N", (810, 1190))
    relation(draw, [(450, 1040), (450, 900), (300, 900), (300, 810)], "N:1", (450, 940))
    relation(draw, [(1570, 1260), (1680, 1260), (1680, 790), (1510, 790)], "N:1", (1680, 1020))
    return save_diagram(image, "14-erd-inventory.png")


def make_erd_fulfillment() -> Path:
    image, draw = diagram_canvas("ERD 6 — Fulfillment và truy vết", "Một phiếu giao cho mỗi cặp order–warehouse trong phạm vi MVP")
    entities = {
        "customer_orders(ref)": ((100, 370, 680, 760), ["PK id", "status/payment status", "fulfillment_status"]),
        "fulfillments": ((930, 330, 1620, 900), ["PK id; UQ shipment_number", "FK order_id/warehouse_id", "status, carrier, tracking", "address_snapshot_json", "version, timestamps", "UQ idempotency_key"]),
        "warehouses(ref)": ((990, 1050, 1550, 1310), ["PK id; UQ code", "name, status"]),
        "fulfillment_items": ((1900, 240, 2540, 620), ["PK id; FK fulfillment_id", "FK order_item_id", "UQ reservation_id", "quantity"]),
        "fulfillment_events": ((1900, 810, 2540, 1200), ["PK id; FK fulfillment_id", "from_status, to_status", "reason, actor", "metadata_json"]),
        "inventory_reservations(ref)": ((1880, 1290, 2550, 1510), ["PK id", "status, quantity", "inventory_item_id"]),
        "order_items(ref)": ((110, 940, 680, 1230), ["PK id; FK order_id", "SKU snapshot", "quantity"]),
    }
    for name, (rect, fields) in entities.items(): draw_entity(draw, rect, name, fields, PALE_BLUE if "ref" not in name else PALE_GOLD)
    relation(draw, [(680, 560), (930, 560)], "1:N", (805, 520))
    relation(draw, [(1275, 900), (1275, 1050)], "N:1", (1320, 975))
    relation(draw, [(1620, 480), (1900, 480)], "1:N", (1760, 440))
    relation(draw, [(1620, 730), (1760, 730), (1760, 980), (1900, 980)], "1:N", (1760, 850))
    relation(draw, [(1100, 900), (800, 900), (800, 1085), (680, 1085)], "1:N", (800, 990))
    relation(draw, [(2540, 430), (2690, 430), (2690, 1400), (2550, 1400)], "0:1", (2640, 690))
    return save_diagram(image, "15-erd-fulfillment.png")


def make_segment_diagram() -> Path:
    image, draw = diagram_canvas("Quy tắc phân khúc khách hàng RFM-2026.1", "View động admin_customer_segments; xét theo thứ tự từ trên xuống")
    rules = [
        ("1", "AT RISK", "Đã có đơn paid và lần paid gần nhất > 90 ngày", PALE_RED, RED),
        ("2", "VIP", "Doanh thu ≥ 5.000.000đ hoặc ≥ 5 đơn paid", PALE_GOLD, GOLD),
        ("3", "LOYAL", "≥ 2 đơn paid và có giao dịch trong 90 ngày", PALE_GREEN, GREEN),
        ("4", "NEW", "Tài khoản ≤ 30 ngày và ≤ 1 đơn paid", PALE_BLUE, BLUE),
        ("5", "POTENTIAL", "Đúng 1 đơn paid", LIGHT_GRAY, DARK_BLUE),
        ("6", "PROSPECT", "Chưa có đơn paid", LIGHT_GRAY, MUTED),
    ]
    y = 220
    for number, name, condition, fill, outline in rules:
        draw.ellipse((120, y + 28, 200, y + 108), fill=f"#{outline}")
        draw_center(draw, (120, y + 28, 200, y + 108), number, img_font(28, True), WHITE)
        box(draw, (270, y, 2380, y + 135), name, condition, fill, outline, 28, 22)
        if y < 220 + 5 * 190:
            arrow(draw, (1325, y + 135), (1325, y + 190), MUTED, 3, 12)
            draw.text((1360, y + 148), "không thỏa", font=img_font(18, italic=True), fill=f"#{MUTED}")
        y += 190
    draw.text((280, 1400), "Chỉ tính doanh thu từ đơn payment_status='paid'; loại đơn cancelled/returned/refunded. Không gán nhãn thủ công.", font=img_font(22, italic=True), fill=f"#{MUTED}")
    return save_diagram(image, "16-customer-segmentation.png")


def make_timeline_diagram() -> Path:
    image, draw = diagram_canvas("Kế hoạch thực hiện 08 tuần — nhóm 05 người", "Mỗi tuần có deliverable kiểm chứng được")
    x0, y0 = 130, 330
    width = 285
    colors = [PALE_BLUE, PALE_GREEN, PALE_GOLD, LIGHT_GRAY]
    for index, (week, scope, owners) in enumerate(TIMELINE):
        x = x0 + index * 300
        draw.ellipse((x + 100, 205, x + 180, 285), fill=f"#{BLUE}")
        draw_center(draw, (x + 100, 205, x + 180, 285), str(index + 1), img_font(28, True), WHITE)
        box(draw, (x, y0, x + width, y0 + 620), week, scope, colors[index % len(colors)], BLUE, 28, 21)
        draw_center(draw, (x + 10, y0 + 490, x + width - 10, y0 + 590), f"Owner: {owners}", img_font(20, True), DARK_BLUE)
        if index < 7:
            arrow(draw, (x + width, y0 + 310), (x + 300, y0 + 310), BLUE, 4, 14)
    box(draw, (400, 1090, 2200, 1320), "Definition of Done", "Code review + test pass + migration/contract cập nhật + demo được + tài liệu đồng bộ", PALE_GREEN, GREEN, 31, 23)
    return save_diagram(image, "17-project-timeline.png")


def generate_diagrams() -> dict[str, Path]:
    paths = {
        "context": make_context_diagram(),
        "architecture": make_architecture_diagram(),
        "uc_storefront": make_use_case_diagram(
            "Use Case — Khách hàng và storefront", "03-use-case-storefront.png",
            [("Khách vãng lai", "Không cần tài khoản"), ("Thành viên", "Đã xác thực")],
            [("Khám phá", ["Tìm/lọc sản phẩm", "Xem chi tiết và biến thể", "Xem hướng dẫn size", "Thêm yêu thích"]),
             ("Mua hàng", ["Quản lý giỏ", "Áp voucher", "Checkout", "Thanh toán"]),
             ("Sau mua", ["Theo dõi đơn", "Xem lịch sử", "Quản lý hồ sơ", "Quản lý địa chỉ"])],
        ),
        "uc_admin": make_use_case_diagram(
            "Use Case — Quản trị cửa hàng", "04-use-case-admin.png",
            [("Staff", "Vận hành hằng ngày"), ("Manager", "Quản lý và phân tích"), ("Admin", "Quyền cao nhất")],
            [("Bán hàng", ["Quản lý đơn", "Xem customer 360", "Lọc phân khúc", "Quản lý catalog"]),
             ("Điều hành", ["Quản lý promotion", "Xem báo cáo", "Xem tài chính", "Dùng AI Copilot"]),
             ("Quản trị", ["Khóa/mở tài khoản", "Xem audit", "Cấu hình", "Theo dõi health"])],
        ),
        "uc_inventory": make_use_case_diagram(
            "Use Case — Inventory đa kho", "05-use-case-inventory.png",
            [("Nhân viên kho", "Tác nghiệp"), ("Manager", "Giám sát/phê duyệt"), ("Hệ thống", "Tự động giữ/xuất tồn")],
            [("Tra cứu", ["Xem ATP theo SKU/kho", "Lọc tồn thấp", "Xuất CSV"]),
             ("Chứng từ", ["Nhập kho", "Chuyển kho", "Kiểm kê", "Điều chỉnh"]),
             ("Tự động", ["Giữ tồn khi checkout", "Release hết hạn", "Consume khi shipped", "Append movement"])],
        ),
        "uc_fulfillment": make_use_case_diagram(
            "Use Case — Điều phối vận chuyển", "06-use-case-fulfillment.png",
            [("Order Operator", "Điều phối kiện"), ("Nhân viên kho", "Pick/pack/bàn giao"), ("Khách hàng", "Theo dõi kết quả")],
            [("Chuẩn bị", ["Xem backlog", "Chuyển processing", "Pick/pack", "Ready to ship"]),
             ("Bàn giao", ["Nhập carrier/tracking", "Chuyển shipped", "Ghi lần giao", "Cập nhật ETA"]),
             ("Kết thúc", ["Delivered", "Failed", "Returned", "Cancelled"])],
        ),
        "seq_auth": make_sequence_diagram(
            "Sequence — Đăng nhập và xoay vòng token", "07-sequence-auth.png",
            ["Browser", "React/BFF", "AuthController", "AuthService", "PostgreSQL"],
            [(0, 1, "POST /api/commerce/auth/login", "call"), (1, 2, "Forward + Origin hợp lệ", "call"),
             (2, 3, "login(email, password)", "call"), (3, 4, "Đọc user/role; verify hash", "call"),
             (4, 3, "User active + roles", "return"), (3, 4, "Lưu refresh token hash", "call"),
             (3, 2, "Access 15m + refresh 30d", "return"), (2, 1, "ApiEnvelope", "return"),
             (1, 0, "Lưu phiên; điều hướng", "return"), (0, 1, "POST /refresh khi access hết hạn", "call"),
             (1, 3, "Rotate refresh", "call"), (3, 4, "Revoke cũ + insert mới", "call")],
        ),
        "seq_checkout": make_sequence_diagram(
            "Sequence — Checkout, reservation và fulfillment đa kho", "08-sequence-checkout.png",
            ["Customer", "React/BFF", "Order API", "PostgreSQL", "Payment", "Fulfillment"],
            [(0, 1, "Xác nhận địa chỉ/payment/voucher", "call"), (1, 2, "POST /orders + idempotency key", "call"),
             (2, 3, "Khóa cart/SKU; revalidate giá và ATP", "call"), (3, 2, "Snapshot hợp lệ", "return"),
             (2, 3, "Insert order/items/history", "call"), (2, 3, "Reserve inventory theo kho", "call"),
             (2, 5, "Nhóm reservation theo warehouse", "call"), (5, 3, "Insert fulfillment/items/events", "call"),
             (2, 4, "Tạo payment COD/online", "call"), (4, 3, "Insert payment + event/outbox", "call"),
             (3, 2, "Commit transaction", "return"), (2, 1, "Order + guest token nếu cần", "return"),
             (1, 0, "Hiển thị mã đơn và từng kiện", "return")],
        ),
        "seq_fulfillment": make_sequence_diagram(
            "Sequence — Chuyển phiếu giao sang shipped/delivered", "09-sequence-fulfillment.png",
            ["Admin UI", "Fulfillment API", "State Machine", "PostgreSQL", "Inventory", "Payment/Audit"],
            [(0, 1, "PATCH tracking + expectedVersion", "call"), (1, 3, "Update tracking optimistic", "call"),
             (0, 1, "PATCH status=shipped", "call"), (1, 2, "Validate transition + payment ready", "call"),
             (2, 1, "Cho phép", "return"), (1, 4, "Consume reservation của kiện", "call"),
             (4, 3, "Giảm on_hand/reserved + movement", "call"), (1, 3, "Event + aggregate order status", "call"),
             (1, 5, "Audit + outbox", "call"), (0, 1, "PATCH status=delivered", "call"),
             (1, 3, "Ghi delivered event", "call"), (1, 5, "Nếu mọi kiện delivered: capture COD", "call")],
        ),
        "erd_identity": make_erd_identity(),
        "erd_catalog": make_erd_catalog(),
        "erd_cart": make_erd_cart_promotion(),
        "erd_order": make_erd_order_payment(),
        "erd_inventory": make_erd_inventory(),
        "erd_fulfillment": make_erd_fulfillment(),
        "segments": make_segment_diagram(),
        "timeline": make_timeline_diagram(),
    }
    return paths


# ---------------------------------------------------------------------------
# Word authoring helpers — narrative_proposal preset + A4 Academic override
# ---------------------------------------------------------------------------


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths: list[float], indent_dxa=120):
    total_dxa = sum(round(w * 1440) for w in widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(round(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=8.5, color=MUTED)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Mở tài liệu trong Word và chọn Update Field để cập nhật mục lục."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, fallback, end):
        run._r.append(node)
    set_run_font(run, size=10, color=MUTED, italic=True)


def configure_document(doc: Document, running_title: str):
    section = doc.sections[0]
    # Named override: A4 Academic Report; content width remains exactly 6.5 in.
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.77)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = running_title
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], size=8.5, color=MUTED, bold=True)
    header.paragraph_format.space_after = Pt(2)
    p_pr = header._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), GRID)
    borders.append(bottom)
    p_pr.append(borders)
    footer = section.footer.paragraphs[0]
    footer.add_run("NÉT Commerce  |  ")
    set_run_font(footer.runs[0], size=8.5, color=MUTED)
    add_page_field(footer)


def set_properties(doc: Document, title: str, subject: str):
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "Nhóm đồ án NÉT Commerce"
    props.keywords = "đồ án, thương mại điện tử, Spring Boot, React, PostgreSQL, inventory, fulfillment"
    props.comments = "Tài liệu được đồng bộ từ mã nguồn và migration Flyway V1–V7."


def add_cover(doc: Document, title: str, subtitle: str, document_type: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(document_type.upper()), size=12, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(title), size=28, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(38)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(subtitle), size=14, color=DARK_BLUE, italic=True)

    table = doc.add_table(rows=len(PROJECT_FACTS), cols=2)
    set_table_geometry(table, [1.75, 4.75])
    set_repeat_header(table.rows[0])
    for idx, (label, value) in enumerate(PROJECT_FACTS):
        set_cell_shading(table.cell(idx, 0), LIGHT_GRAY)
        set_run_font(table.cell(idx, 0).paragraphs[0].add_run(label), size=9.5, color=NAVY, bold=True)
        set_run_font(table.cell(idx, 1).paragraphs[0].add_run(value), size=9.5, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(64)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("TP. HỒ CHÍ MINH — THÁNG 09/2026"), size=11, color=MUTED, bold=True)
    doc.add_page_break()


def add_para(doc: Document, text: str, *, bold_lead: str | None = None, italic=False, align=None, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=11, color=INK, bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), size=11, color=INK, italic=italic)
    else:
        set_run_font(p.add_run(text), size=11, color=INK, italic=italic)
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        set_run_font(p.add_run(item), size=11, color=INK)


def create_numbering_instance(doc: Document) -> int:
    """Tạo một numbering instance riêng để mỗi quy trình bắt đầu lại từ 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract_num.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.extend((start, number_format, level_text, level_justification))

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "540")
    indentation.set(qn("w:hanging"), "280")
    paragraph_properties.extend((tabs, indentation))
    level.append(paragraph_properties)

    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    run_properties.extend((fonts, size))
    level.append(run_properties)
    abstract_num.append(level)

    # ECMA-376 yêu cầu abstractNum đứng trước các phần tử num.
    abstract_position = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(abstract_position, abstract_num)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    cleanup_position = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:numIdMacAtCleanup")),
        len(numbering),
    )
    numbering.insert(cleanup_position, number)
    return num_id


def add_numbered(doc: Document, items: list[str]):
    num_id = create_numbering_instance(doc)
    for item in items:
        p = doc.add_paragraph()
        paragraph_properties = p._p.get_or_add_pPr()
        numbering_properties = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        numbering_properties.extend((level, number))
        paragraph_properties.insert(0, numbering_properties)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        set_run_font(p.add_run(item), size=11, color=INK)


def add_callout(doc: Document, label: str, text: str, kind="info"):
    palette = {
        "info": (PALE_BLUE, BLUE),
        "success": (PALE_GREEN, GREEN),
        "warning": (PALE_GOLD, GOLD),
        "risk": (PALE_RED, RED),
    }
    fill, accent = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_repeat_header(table.rows[0])
    set_cell_shading(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(f"{label}: "), size=10, color=accent, bold=True)
    set_run_font(p.add_run(text), size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[tuple], widths: list[float], font_size=8.8, caption: str | None = None):
    if caption:
        p = doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(caption), size=9, color=MUTED, italic=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_shading(table.cell(0, index), NAVY)
        p = table.cell(0, index).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(header), size=font_size, color=WHITE, bold=True)
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "FAFBFC")
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(str(value)), size=font_size, color=INK)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc: Document, path: Path, caption: str, number: int, width=6.35, alt_text: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("name", f"Hình {number}")
    doc_pr.set("descr", alt_text or caption)
    cp = doc.add_paragraph(style="Caption")
    set_run_font(cp.add_run(f"Hình {number}. {caption}"), size=9, color=MUTED, italic=True)


def add_chapter(doc: Document, number: str, title: str):
    p = doc.add_heading(f"CHƯƠNG {number}. {title}", level=1)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True


def add_landscape_figures(doc: Document, figures: list[tuple[Path, str, int]]):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.35)
    for index, (path, caption, number) in enumerate(figures):
        if index:
            doc.add_page_break()
        add_figure(doc, path, caption, number, width=9.9, alt_text=caption)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.77)
    section.header_distance = section.footer_distance = Inches(0.45)


def add_detailed_use_case(doc: Document, use_case: tuple):
    code, title, actor, objective, precondition, flow = use_case
    doc.add_heading(f"{code} — {title}", level=2)
    add_table(
        doc,
        ["Thuộc tính", "Nội dung"],
        [
            ("Tác nhân chính", actor),
            ("Mục tiêu", objective),
            ("Tiền điều kiện", precondition),
            ("Kích hoạt", f"Tác nhân chọn chức năng {title.lower()}."),
            ("Hậu điều kiện", "Dữ liệu được commit nguyên tử; response có mã lỗi ổn định khi thất bại."),
        ],
        [1.55, 4.95],
        8.8,
    )
    doc.add_heading("Luồng chính", level=3)
    add_numbered(doc, flow)
    doc.add_heading("Luồng thay thế/ngoại lệ", level=3)
    add_bullets(doc, [
        "Dữ liệu đầu vào không hợp lệ: trả 400 với error code; không ghi dữ liệu dở dang.",
        "Không đủ quyền hoặc không đúng chủ thể: trả 401/403 và không vào business service.",
        "Version/idempotency xung đột: trả 409; client tải lại snapshot trước khi thử lại.",
        "Dependency tạm thời lỗi: rollback transaction, giữ correlation ID để truy vết.",
    ])


def requirement_domain(code: str) -> str:
    return {
        "AUT": "Xác thực", "ACC": "Tài khoản", "CAT": "Catalog", "CAR": "Giỏ hàng",
        "PRO": "Khuyến mãi", "CHK": "Checkout", "PAY": "Thanh toán", "ORD": "Đơn hàng",
        "INV": "Tồn kho", "FUL": "Vận chuyển", "CUS": "Khách hàng", "ADM": "Quản trị",
        "AI": "AI", "AUD": "Kiểm soát",
    }.get(code.split("-")[1], "Khác")


def build_main_document(diagrams: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc, "ĐỒ ÁN NÉT COMMERCE — BÁO CÁO TỔNG HỢP")
    set_properties(doc, "Báo cáo đồ án website bán quần áo trực tuyến", "Phân tích, thiết kế, cài đặt, kiểm thử và triển khai NÉT Commerce")
    add_cover(doc, "XÂY DỰNG WEBSITE BÁN QUẦN ÁO TRỰC TUYẾN", "Phân tích nghiệp vụ · Thiết kế kỹ thuật · Cài đặt · Kiểm thử · Vận hành", "Báo cáo đồ án")

    doc.add_heading("LỜI CAM ĐOAN VÀ PHẠM VI SỐ LIỆU", level=1)
    add_para(doc, "Báo cáo được biên soạn từ mã nguồn đang có trong workspace, API Java Spring Boot, giao diện React/Vinext, Docker Compose và migration Flyway V1–V7. Các con số 22 sản phẩm, 58 biến thể, 33 bảng và kết quả kiểm thử được đối chiếu với phiên bản hiện tại.")
    add_para(doc, "Số liệu khảo sát n=200 trong chương 1 là bộ dữ liệu mô phỏng để minh họa phương pháp phân tích đồ án; không được diễn giải như khảo sát thực địa. Những chức năng chưa tích hợp production như hãng vận chuyển, cổng thanh toán thật, hoàn tiền và outbox dispatcher được nêu rõ là giới hạn hoặc hướng phát triển.")
    add_callout(doc, "NGUYÊN TẮC ĐỒNG BỘ", "Khi tài liệu và mã nguồn khác nhau, migration Flyway cùng contract API đang chạy là nguồn sự thật. Mọi thay đổi schema mới phải dùng migration tiếp theo, không sửa migration đã áp dụng.", "warning")

    doc.add_heading("TÓM TẮT ĐỒ ÁN", level=1)
    add_para(doc, "NÉT Commerce là hệ thống thương mại điện tử thời trang B2C gồm storefront cho khách hàng và backoffice cho nhân viên. Người mua có thể duyệt sản phẩm, chọn biến thể, quản lý giỏ, áp voucher, đăng ký/đăng nhập, checkout, thanh toán COD hoặc online mô phỏng và theo dõi từng kiện hàng. Nhân viên quản lý catalog, tồn kho đa kho, đơn hàng, vận chuyển, khách hàng, phân khúc, khuyến mãi, báo cáo và AI Copilot.")
    add_para(doc, "Kiến trúc hiện tại dùng React 19, Vinext, TypeScript và Tailwind ở frontend/BFF; Java 21 và Spring Boot 4.1.1 ở backend; PostgreSQL 17, Redis 7.4 và volume media ở lớp dữ liệu. Fulfillment được tự tạo theo từng kho khi checkout, liên kết reservation đến từng dòng kiện, bảo đảm chỉ xuất đúng tồn khi chuyển sang shipped. View admin_customer_segments tính động sáu nhóm khách hàng theo ưu tiên RFM-2026.1.")
    add_para(doc, "Phạm vi xác minh hiện có 69 bài kiểm thử backend, trong đó 10 bài PostgreSQL/Testcontainers, và 12 bài kiểm thử frontend; lint, typecheck, build và smoke test local đều đạt tại lần nghiệm thu gần nhất. Hệ thống chạy bằng Docker Compose trên cổng 8787 cho frontend và 8080 cho API.")
    add_callout(doc, "TỪ KHÓA", "thương mại điện tử; thời trang; Java Spring Boot; React; PostgreSQL; inventory đa kho; customer segmentation; fulfillment; Docker", "info")

    doc.add_heading("MỤC LỤC", level=1)
    add_toc_field(doc.add_paragraph())
    doc.add_heading("CÁCH ĐỌC BÁO CÁO", level=1)
    add_table(doc, ["Mục tiêu đọc", "Chương", "Bằng chứng"], [
        ("Hiểu lý do và phạm vi", "1–2", "Bối cảnh, khảo sát mô phỏng, actor, FR/NFR"),
        ("Đánh giá nghiệp vụ", "3 và 6", "Use case, sequence, state machine, business rule"),
        ("Đánh giá kỹ thuật", "4–5", "Kiến trúc, API, ERD 33 bảng + 1 view"),
        ("Đánh giá sản phẩm", "7–8", "Module đã cài đặt, test và giới hạn"),
        ("Đánh giá quản trị dự án", "9", "Timeline 08 tuần, owner, RACI, RAID"),
    ], [2.0, 1.0, 3.5], 9)

    add_chapter(doc, "1", "TỔNG QUAN VÀ LÝ DO CHỌN ĐỀ TÀI")
    doc.add_heading("1.1 Bối cảnh", level=2)
    add_para(doc, "Mua quần áo trực tuyến đòi hỏi người dùng chọn đúng size/màu, hiểu giá cuối, biết sản phẩm còn tại kho nào và theo dõi giao hàng minh bạch. Ở phía vận hành, catalog, tồn kho, đơn, thanh toán và vận chuyển phải cùng một nguồn dữ liệu; nếu tách rời sẽ dễ sai SKU, oversell, chậm xử lý và khó đối soát.")
    doc.add_heading("1.2 Vấn đề nghiên cứu", level=2)
    add_bullets(doc, [
        "Trải nghiệm mua hàng dễ đứt đoạn giữa tìm kiếm, chọn biến thể, voucher, địa chỉ và thanh toán.",
        "Tồn kho phải phản ánh theo từng kho/SKU và phân biệt on-hand, reserved, ATP.",
        "Một đơn có thể lấy hàng từ nhiều kho, do đó cần nhiều phiếu giao nhưng vẫn tổng hợp về một order.",
        "Quản trị cần customer 360 và phân khúc có quy tắc minh bạch thay vì gán nhãn cảm tính.",
        "Mọi mutation quan trọng cần idempotency, audit và lịch sử để điều tra sự cố.",
    ])
    doc.add_heading("1.3 Mục tiêu", level=2)
    add_table(doc, ["Mã", "Mục tiêu", "Chỉ báo hoàn thành"], [
        ("MT-01", "Hoàn thiện luồng mua hàng end-to-end", "Browse → cart → voucher → checkout → tracking chạy trên local Docker"),
        ("MT-02", "Xây dựng backoffice có thể vận hành", "Catalog, order, customer, inventory, fulfillment, reports, AI có màn hình/API"),
        ("MT-03", "Bảo đảm tính đúng dữ liệu", "Transaction, constraint, state machine, idempotency và test"),
        ("MT-04", "Có hồ sơ đồ án truy vết được", "FR → UC → API → bảng → test"),
    ], [0.8, 2.5, 3.2], 8.8)
    doc.add_heading("1.4 Khảo sát mô phỏng và cách sử dụng", level=2)
    add_para(doc, "Nhóm xây dựng bộ dữ liệu mô phỏng 200 phản hồi để thực hành phương pháp khảo sát. Các tỷ lệ sau dùng cho ưu tiên giả thuyết, không phải kết quả thực địa: không chắc size 68%; sản phẩm khác ảnh/mô tả 59%; đổi trả khó hoặc thiếu thông tin 54%; phí xuất hiện muộn 47%; không biết tồn size/màu 46%; tracking không rõ 35%; lo thanh toán online 26%.")
    add_callout(doc, "GIỚI HẠN NGHIÊN CỨU", "Muốn công bố kết luận thị trường phải phát hành bảng hỏi thật, mô tả mẫu, loại phản hồi không hợp lệ và tính sai số. Báo cáo hiện chỉ dùng dữ liệu mô phỏng để thiết kế phạm vi đồ án.", "warning")
    doc.add_heading("1.5 Phạm vi", level=2)
    add_table(doc, ["Trong phạm vi đã cài đặt", "Ngoài phạm vi/tiếp theo"], [
        ("Storefront, auth/account, wishlist, cart/voucher, checkout, order tracking", "Marketplace nhiều nhà bán, loyalty điểm thưởng, recommendation cá nhân hóa"),
        ("Catalog/admin, inventory đa kho, customer segment, fulfillment", "Webhook/API hãng vận chuyển và nhãn giao hàng tự động"),
        ("COD và payment online mock", "Cổng thanh toán production, refund/chargeback đầy đủ"),
        ("AI Copilot dạng rule-based trên snapshot", "LLM production có guardrail/evaluation và quyền thực thi"),
    ], [3.25, 3.25], 8.8)
    add_figure(doc, diagrams["context"], "Bối cảnh hệ thống và các chủ thể", 1)

    add_chapter(doc, "2", "PHÂN TÍCH SẢN PHẨM VÀ YÊU CẦU")
    doc.add_heading("2.1 Tầm nhìn sản phẩm", level=2)
    add_para(doc, "Dành cho người mua thời trang cần một trải nghiệm rõ ràng và đội cửa hàng cần vận hành thống nhất, NÉT Commerce là nền tảng B2C kết nối catalog, SKU, khuyến mãi, tồn kho, đơn hàng, khách hàng và giao nhận trong một hệ thống có thể kiểm chứng. Khác với một giao diện demo tĩnh, mọi tổng tiền, giữ tồn và chuyển trạng thái đều do backend chịu trách nhiệm.")
    doc.add_heading("2.2 Nhóm người dùng", level=2)
    add_table(doc, ["Nhóm", "Nhu cầu chính", "Quyền điển hình"], [
        ("Khách vãng lai", "Tìm hàng nhanh, checkout ít bước, tra cứu đơn an toàn", "Catalog, cart, voucher, guest checkout/tracking"),
        ("Khách thành viên", "Lưu hồ sơ/địa chỉ/yêu thích và xem lịch sử", "Customer-owned resources"),
        ("Staff", "Xử lý catalog, đơn, kho, khách và giao hàng", "Tác nghiệp backoffice"),
        ("Manager", "Theo dõi KPI, tài chính, promotion, ngoại lệ", "Read/approve/analysis"),
        ("Admin", "Quyền cao nhất và kiểm soát tài khoản", "Toàn bộ backoffice theo policy"),
    ], [1.35, 3.1, 2.05], 8.8)
    doc.add_heading("2.3 Yêu cầu chức năng", level=2)
    fr_rows = [(code, requirement_domain(code), text, priority, status) for code, text, priority, status in FUNCTIONAL_REQUIREMENTS]
    add_table(doc, ["Mã", "Miền", "Yêu cầu", "Ưu tiên", "Trạng thái"], fr_rows, [0.85, 0.9, 3.55, 0.55, 0.65], 7.8, "Bảng 2.1. Danh mục yêu cầu chức năng đồng bộ với phiên bản hiện tại")
    doc.add_heading("2.4 Yêu cầu phi chức năng", level=2)
    add_table(doc, ["Mã", "Nhóm", "Yêu cầu/tiêu chí", "Cách kiểm chứng"], NON_FUNCTIONAL_REQUIREMENTS, [0.9, 1.0, 3.75, 0.85], 7.8, "Bảng 2.2. Yêu cầu phi chức năng có tiêu chí đo")

    add_chapter(doc, "3", "USE CASE VÀ ĐẶC TẢ NGHIỆP VỤ")
    add_para(doc, "Use case được chia theo miền để tránh một sơ đồ tổng quá dày. Mỗi đường chỉ nối tác nhân với một nhóm chức năng; chi tiết include/extend và điều kiện nằm trong đặc tả văn bản, giúp sơ đồ không có đường chồng chéo.")
    add_landscape_figures(doc, [
        (diagrams["uc_storefront"], "Use Case phía khách hàng", 2),
        (diagrams["uc_admin"], "Use Case trang quản trị", 3),
        (diagrams["uc_inventory"], "Use Case Inventory đa kho", 4),
        (diagrams["uc_fulfillment"], "Use Case điều phối vận chuyển", 5),
    ])
    doc.add_heading("3.1 Danh mục và đặc tả use case", level=2)
    for use_case in USE_CASES:
        add_detailed_use_case(doc, use_case)

    add_chapter(doc, "4", "KIẾN TRÚC VÀ THIẾT KẾ KỸ THUẬT")
    doc.add_heading("4.1 Lựa chọn kiến trúc", level=2)
    add_para(doc, "Giải pháp dùng modular monolith cho Java API để phù hợp đội 05 người và thời gian 08 tuần: giao dịch xuyên order–inventory–fulfillment vẫn ở một database transaction, nhưng mã được tách theo domain package. Frontend dùng BFF /api/commerce/** để chuẩn hóa cookie/header, giảm việc browser gọi trực tiếp API và giữ contract UI ổn định.")
    add_figure(doc, diagrams["architecture"], "Kiến trúc triển khai hiện tại", 6)
    doc.add_heading("4.2 Tech stack", level=2)
    add_table(doc, ["Lớp", "Công nghệ", "Lý do"], [
        ("Frontend", "React 19, Vinext, TypeScript 5.9, Tailwind CSS 4, Base UI", "SSR/RSC, type safety, component tái sử dụng, responsive"),
        ("BFF", "Vinext route handlers", "Forward auth/guest headers, normalize Origin và lỗi"),
        ("Backend", "Java 21, Spring Boot 4.1.1, Security, JDBC, Validation", "Transaction rõ, RBAC, ecosystem kiểm thử"),
        ("Data", "PostgreSQL 17, Flyway V1–V7", "Constraint quan hệ, JSONB, migration có thứ tự"),
        ("Cache", "Redis 7.4", "Rate limit và dữ liệu tạm"),
        ("DevOps", "Docker Compose, multi-stage Dockerfile, healthcheck", "Môi trường local tái lập được"),
        ("Testing", "JUnit/Spring Test/Testcontainers; Vitest", "Unit + PostgreSQL integration + frontend rule tests"),
    ], [1.1, 2.4, 3.0], 8.7)
    doc.add_heading("4.3 Bảo mật", level=2)
    add_bullets(doc, [
        "JWT access 15 phút; refresh token opaque 30 ngày lưu dạng hash và rotate.",
        "Token xác minh email/đặt lại mật khẩu dùng một lần; có expires_at và consumed_at.",
        "Guest order token ký có hạn 07 ngày và lưu sessionStorage. Contract yêu cầu X-Guest-Order-Token; màn tracking đã tuân thủ, nhưng online-payment client còn ghép token vào query và được ghi nhận là defect.",
        "Spring Security kiểm tra RBAC ở server; OriginValidationFilter bảo vệ request ghi.",
        "Validation, stable error code, correlation ID, audit log và idempotency giảm rủi ro thao tác lặp.",
    ])
    doc.add_heading("4.4 Danh mục API", level=2)
    add_table(doc, ["Miền", "Method", "Endpoint", "Quyền", "Mục đích"], API_CATALOG, [0.8, 0.7, 2.55, 0.9, 1.55], 7.4)

    add_chapter(doc, "5", "THIẾT KẾ CƠ SỞ DỮ LIỆU")
    doc.add_heading("5.1 Nguyên tắc", level=2)
    add_bullets(doc, [
        "Schema hiện tại có 33 bảng nghiệp vụ; không tính flyway_schema_history. View admin_customer_segments là read model, không phải bảng vật lý.",
        "Khóa chính là varchar(64), do application sinh; tài liệu cũ mô tả UUID đã được loại bỏ.",
        "Tiền lưu ở cột *_cents kiểu bigint/numeric số nguyên VND để không phát sinh sai số dấu phẩy động.",
        "Order item, promotion và địa chỉ giao hàng lưu snapshot để lịch sử không thay đổi theo catalog/profile.",
        "Movement, order history, payment event và fulfillment event là dữ liệu append-only về mặt nghiệp vụ.",
        "Refund chưa có bảng vật lý trong V1–V7; đây là hướng phát triển, không tuyên bố đã cài đặt.",
    ])
    doc.add_heading("5.2 Danh mục 33 bảng và 01 view", level=2)
    add_table(doc, ["Miền", "Bảng", "Vai trò"], TABLE_CATALOG, [1.25, 2.2, 3.05], 8.3)
    add_table(doc, ["Đối tượng", "Loại", "Nguồn", "Mục đích"], [
        ("admin_customer_segments", "VIEW", "users + customer_orders", "Tổng hợp KPI khách hàng và segment_code theo RFM-2026.1"),
    ], [2.1, 0.7, 1.8, 1.9], 8.5)
    doc.add_heading("5.3 ERD theo miền", level=2)
    add_para(doc, "ERD được tách thành sáu hình. Mỗi hình chỉ hiển thị các cột định danh và cột nghiệp vụ chính; toàn bộ cột, nullability, default, khóa và bất biến nằm ở docs/ba/data-dictionary, mỗi bảng một file.")
    add_landscape_figures(doc, [
        (diagrams["erd_identity"], "ERD định danh và tài khoản", 7),
        (diagrams["erd_catalog"], "ERD catalog và media", 8),
        (diagrams["erd_cart"], "ERD giỏ hàng và khuyến mãi", 9),
        (diagrams["erd_order"], "ERD đơn hàng và thanh toán", 10),
        (diagrams["erd_inventory"], "ERD Inventory đa kho", 11),
        (diagrams["erd_fulfillment"], "ERD fulfillment và truy vết", 12),
    ])

    add_chapter(doc, "6", "QUY TRÌNH NGHIỆP VỤ TRỌNG YẾU")
    doc.add_heading("6.1 Xác thực", level=2)
    add_figure(doc, diagrams["seq_auth"], "Sequence đăng nhập và refresh token rotation", 13)
    doc.add_heading("6.2 Checkout đa kho", level=2)
    add_figure(doc, diagrams["seq_checkout"], "Sequence checkout tạo reservation và fulfillment", 14)
    doc.add_heading("6.3 Vận chuyển", level=2)
    add_figure(doc, diagrams["seq_fulfillment"], "Sequence shipped/delivered và tác động kho, COD", 15)
    doc.add_heading("6.4 State machine fulfillment", level=2)
    add_table(doc, ["Từ trạng thái", "Có thể chuyển đến", "Điều kiện chính", "Tác động"], [
        ("pending", "processing, cancelled", "Đơn paid hoặc cod_pending; hủy phải có lý do", "Cancel trước dispatch release reservation"),
        ("processing", "ready_to_ship, cancelled", "Đã bắt đầu pick/pack", "Không xuất tồn"),
        ("ready_to_ship", "shipped, cancelled", "Có carrier và tracking", "Shipped consume đúng reservation"),
        ("shipped", "delivered, failed, returned", "Đã bàn giao vật lý", "Không cho cancel trực tiếp"),
        ("failed", "processing, shipped, returned", "Có lý do; tăng delivery_attempts", "Chờ retry hoặc hoàn"),
        ("delivered", "returned", "Có xác nhận giao", "Tổng hợp COD khi tất cả kiện giao"),
        ("returned", "—", "Hàng quay về kho", "Không tự restock; phải kiểm định và nhập kho"),
        ("cancelled", "—", "Chưa dispatch", "Release reservation"),
    ], [1.05, 1.45, 2.5, 1.5], 8.1)
    doc.add_heading("6.5 Phân khúc khách hàng", level=2)
    add_figure(doc, diagrams["segments"], "Quy tắc ưu tiên phân khúc RFM-2026.1", 16)
    doc.add_heading("6.6 Quy tắc nghiệp vụ cốt lõi", level=2)
    rules = [
        ("BR-01", "Giá/tổng tiền", "Server luôn tính lại giá, voucher, phí và tổng trước khi tạo đơn."),
        ("BR-02", "Tồn kho", "ATP = on_hand − reserved − safety_stock; không được âm hoặc oversell."),
        ("BR-03", "Idempotency", "Cùng idempotency key phải trả lại kết quả cũ hoặc conflict, không nhân đôi side effect."),
        ("BR-04", "Đơn đa kho", "Số fulfillment bằng số warehouse khác nhau trong reservation được cấp."),
        ("BR-05", "Xuất kho", "Chỉ status shipped mới consume reservation và giảm on_hand/reserved."),
        ("BR-06", "Hủy", "Đã shipped không được hủy; đơn đã paid/authorized phải hoàn tiền trước khi hủy."),
        ("BR-07", "Hoàn hàng", "Returned sau dispatch không tự cộng tồn; kho kiểm định rồi lập receipt/adjustment."),
        ("BR-08", "COD", "Chỉ paid khi tất cả fulfillment delivered; refused/returned trước delivery hủy payment."),
        ("BR-09", "Guest tracking", "Token 07 ngày trong header; không đặt token trên URL/log."),
        ("BR-10", "Segment", "Segment được tính động theo thứ tự ưu tiên; không gán thủ công."),
        ("BR-11", "Đồng thời", "expectedVersion bắt buộc khi sửa tracking/status fulfillment."),
        ("BR-12", "AI", "Copilot chỉ tư vấn trên snapshot và không có quyền mutation."),
    ]
    add_table(doc, ["Mã", "Miền", "Quy tắc"], rules, [0.75, 1.1, 4.65], 8.5)

    add_chapter(doc, "7", "CÀI ĐẶT VÀ CÁC PHÂN HỆ")
    doc.add_heading("7.1 Frontend", level=2)
    add_table(doc, ["Phân hệ", "Route tiêu biểu", "Chức năng"], [
        ("Storefront", "/, /san-pham, /san-pham/[slug]", "Landing, PLP/PDP, tìm/lọc, biến thể"),
        ("Auth", "/dang-ky, /dang-nhap, /xac-minh-email, /quen-mat-khau", "Vòng đời tài khoản"),
        ("Account", "/tai-khoan/**, /yeu-thich", "Hồ sơ, địa chỉ, lịch sử, wishlist"),
        ("Checkout", "/gio-hang, /thanh-toan, /don-hang/[id]", "Giỏ, voucher, checkout, tracking"),
        ("Admin", "/admin/**", "Dashboard, catalog, order, customer, inventory, fulfillment, report, AI"),
    ], [1.1, 2.4, 3.0], 8.6)
    doc.add_heading("7.2 Backend", level=2)
    add_table(doc, ["Module Java", "Trách nhiệm"], [
        ("auth/account/wishlist", "Xác thực, token, hồ sơ, địa chỉ, yêu thích"),
        ("catalog/cart", "Sản phẩm, media, giỏ, voucher và pricing"),
        ("order/payment", "Checkout, order state, payment event"),
        ("inventory", "ATP, receipt, transfer, count, adjust, reservation expiry"),
        ("fulfillment", "Phiếu giao, tracking, state machine, consume/release tồn"),
        ("admincustomer/adminorder/admin.insights", "Customer 360, segment, order admin, report, finance, Copilot"),
        ("common/config", "ApiEnvelope, error, correlation ID, security, idempotency"),
    ], [2.4, 4.1], 8.7)
    doc.add_heading("7.3 How to run", level=2)
    add_callout(doc, "YÊU CẦU", "Docker Desktop đang chạy; cổng 8787 và 8080 chưa bị chiếm; thực hiện từ thư mục platform.", "info")
    add_table(doc, ["Bước", "Lệnh PowerShell", "Kết quả"], [
        ("1", "Copy-Item .env.example .env", "Tạo cấu hình frontend/BFF"),
        ("2", "Copy-Item backend/.env.example backend/.env", "Tạo cấu hình Java API"),
        ("3", "Set-Location backend; docker --context default compose up -d --build", "PostgreSQL, Redis, API :8080"),
        ("4", "Set-Location ..; docker --context default compose up -d --build", "Frontend/BFF :8787"),
        ("5", "Invoke-WebRequest http://localhost:8080/api/health", "API trả HTTP 200"),
        ("6", "Invoke-WebRequest http://localhost:8787/api/health", "BFF health trả HTTP 200"),
        ("7", "node ops/smoke-test.mjs", "Smoke test storefront/admin"),
    ], [0.45, 3.85, 2.2], 8.2)
    add_para(doc, "Tài khoản demo quản trị local: owner@example.com / OwnerNet2026. Chỉ dùng trong môi trường demo; phải đổi secret và credentials trước mọi môi trường chia sẻ.", italic=True)

    add_chapter(doc, "8", "KIỂM THỬ, NGHIỆM THU VÀ GIỚI HẠN")
    doc.add_heading("8.1 Kết quả kiểm thử", level=2)
    add_table(doc, ["Hạng mục", "Kết quả", "Phạm vi"], [
        ("Backend unit/integration", "69 test pass", "Trong đó 10 PostgreSQL/Testcontainers; auth, order, payment, inventory, segmentation, fulfillment"),
        ("Frontend unit", "12 test pass", "Catalog, promotion, order state, inventory rule"),
        ("Frontend static", "Lint + typecheck + build pass", "TypeScript và bundle production"),
        ("Runtime local", "Frontend/API/PostgreSQL/Redis healthy", "Health trực tiếp và qua BFF"),
        ("Smoke", "Các endpoint/màn hình trọng yếu HTTP 200", "Login, segments, customer, fulfillment, admin routes"),
    ], [1.55, 1.25, 3.7], 8.5)
    doc.add_heading("8.2 Tiêu chí nghiệm thu", level=2)
    add_bullets(doc, [
        "Đăng ký → xác minh → đăng nhập → refresh → logout hoạt động theo contract.",
        "Voucher hợp lệ/sai/hết hạn cho kết quả đúng và checkout revalidate.",
        "Checkout một hoặc nhiều kho tạo đúng order, reservation, fulfillment và payment.",
        "Shipped chỉ xuất tồn thuộc kiện; returned không tự restock.",
        "Customer segments trả đủ sáu nhóm và chi tiết customer 360.",
        "Admin routes tải được sau đăng nhập; request không đủ quyền bị từ chối.",
    ])
    doc.add_heading("8.3 Giới hạn đã biết", level=2)
    add_table(doc, ["Giới hạn", "Ảnh hưởng", "Hướng xử lý"], [
        ("Chưa tích hợp hãng vận chuyển", "Carrier/tracking/ETA nhập thủ công; không có webhook scan", "Adapter provider + webhook idempotent"),
        ("Guest online-payment lệch contract", "Client còn gắn token vào query nên nhánh guest có thể bị 400 và token có thể lọt vào URL log", "Gọi /orders/{id}/guest và truyền X-Guest-Order-Token như order tracking"),
        ("Expiry chưa đồng bộ payment aggregate", "Job có thể hủy order nhưng để payments.status ở pending", "Cập nhật payment và order trong cùng transaction, bổ sung integration test"),
        ("Reservation COD không được job quét", "Đơn COD ở confirmed nên reservation 30 ngày có thể không được giải phóng khi hết hạn", "Chốt chính sách TTL COD và mở rộng điều kiện truy vấn của scheduler"),
        ("Webhook thành công đến sau khi hủy đơn", "Payment có thể thành succeeded/paid trong khi order vẫn cancelled", "Chặn hoặc reconcile event với trạng thái terminal của order"),
        ("Payment mới là mock", "Không dùng để thu tiền thật", "Tích hợp provider và đối soát"),
        ("Chưa có refund aggregate", "Paid order cần quy trình hoàn tiền ngoài hệ thống trước khi hủy", "Thêm migration refunds và service idempotent"),
        ("Returned chưa tự đánh giá chất lượng", "Kho phải kiểm định và lập receipt/adjustment", "Return inspection/disposition"),
        ("Outbox chưa có dispatcher production", "Sự kiện mới dừng ở lưu trữ", "Worker retry/dead-letter/metrics"),
        ("AI rule-based", "Không hội thoại ngôn ngữ tự nhiên sâu", "LLM có grounding, guardrail và evaluation"),
    ], [1.8, 2.2, 2.5], 8.3)

    add_chapter(doc, "9", "QUẢN LÝ DỰ ÁN TRONG 02 THÁNG")
    add_figure(doc, diagrams["timeline"], "Timeline 08 tuần và đầu mối thực hiện", 17)
    doc.add_heading("9.1 Phân công đầu việc", level=2)
    add_table(doc, ["Vai trò", "Trách nhiệm đầu cuối", "Deliverable chính"], [
        ("PM", "Phạm vi, backlog, lịch, RAID, UAT, tài liệu", "Charter, PRD/SRS, sprint plan, acceptance"),
        ("Frontend", "Storefront, auth/account UI, checkout, admin UI, BFF", "Responsive pages, state/error UX, build"),
        ("Backend 1", "Auth, account, catalog, cart, order, security", "API contract, transaction, tests"),
        ("Backend 2", "Inventory, payment, customer segment, fulfillment", "State machines, V7, integration tests"),
        ("DevOps", "Docker, env, CI, health, log, backup/restore", "Compose, runbook, smoke, release evidence"),
    ], [1.05, 3.1, 2.35], 8.5)
    doc.add_heading("9.2 WBS theo tuần", level=2)
    add_table(doc, ["Mốc", "Phạm vi", "Owner"], TIMELINE, [0.75, 4.4, 1.35], 8.5)
    doc.add_heading("9.3 RACI rút gọn", level=2)
    add_table(doc, ["Workstream", "PM", "FE", "BE1", "BE2", "DevOps"], [
        ("Vision/Scope/Acceptance", "A/R", "C", "C", "C", "I"),
        ("Storefront/BFF/Admin UI", "C", "A/R", "C", "C", "C"),
        ("Auth/Catalog/Cart/Order", "C", "C", "A/R", "C", "C"),
        ("Inventory/Payment/Segment/Fulfillment", "C", "C", "C", "A/R", "C"),
        ("Docker/CI/Observability/Release", "A", "C", "C", "C", "R"),
        ("UAT/Tài liệu/Bàn giao", "A/R", "C", "C", "C", "C"),
    ], [2.5, 0.65, 0.65, 0.65, 0.65, 1.4], 8.2)
    doc.add_heading("9.4 Rủi ro chính", level=2)
    add_table(doc, ["Rủi ro", "Xác suất/Tác động", "Ứng phó", "Owner"], [
        ("Phạm vi vượt 08 tuần", "Cao/Cao", "Must-first, freeze scope cuối tuần 6", "PM"),
        ("Oversell hoặc sai reservation", "TB/Rất cao", "Constraint, transaction, integration test", "BE2"),
        ("Auth/token lộ qua log/URL", "TB/Rất cao", "Hash token; TTL; guest header; log scrub", "BE1"),
        ("FE và API lệch contract", "TB/Cao", "BFF, error envelope, contract review mỗi sprint", "FE/BE1"),
        ("Migration làm hỏng dữ liệu", "Thấp/Rất cao", "Flyway immutable, backup, fresh/upgrade test", "DevOps"),
        ("Tài liệu lệch code", "Cao/TB", "Sinh Word từ một script và kiểm tra traceability", "PM"),
    ], [2.0, 1.25, 2.45, 0.8], 8.2)

    add_chapter(doc, "10", "ĐÁNH GIÁ VÀ HƯỚNG PHÁT TRIỂN")
    doc.add_heading("10.1 Kết quả", level=2)
    add_para(doc, "Đồ án đã chuyển từ website trưng bày sang một hệ thống giao dịch có backend Java, xác thực đầy đủ, inventory đa kho, customer segmentation và fulfillment theo từng kho. Điểm mạnh nổi bật là chuỗi dữ liệu checkout–reservation–fulfillment–inventory được gắn kết trong transaction, giúp giải thích rõ vì sao một đơn có nhiều kiện và khi nào tồn thực sự bị xuất.")
    doc.add_heading("10.2 Hướng phát triển ưu tiên", level=2)
    add_numbered(doc, [
        "Tích hợp cổng thanh toán production, refund aggregate và đối soát.",
        "Tích hợp adapter hãng vận chuyển, webhook, nhãn giao hàng và retry/dead-letter.",
        "Bổ sung return inspection/disposition để quyết định restock, cách ly hoặc loại bỏ.",
        "Triển khai outbox dispatcher, metrics/SLO và backup/restore drill định kỳ.",
        "Nâng AI Copilot sang LLM có grounding, quyền tối thiểu, audit và evaluation.",
        "Thực hiện khảo sát thực địa và usability test để thay dữ liệu mô phỏng bằng bằng chứng người dùng.",
    ])
    add_callout(doc, "KẾT LUẬN", "Phiên bản hiện tại phù hợp mục tiêu đồ án 08 tuần và môi trường local Docker. Để vận hành thương mại thật cần hoàn thiện tích hợp thanh toán/vận chuyển, quy trình hoàn tiền–đổi trả, giám sát và kiểm thử tải/bảo mật.", "success")

    add_chapter(doc, "A", "PHỤ LỤC — MA TRẬN TRUY VẾT")
    trace_rows = []
    for code, text, priority, status in FUNCTIONAL_REQUIREMENTS:
        domain = requirement_domain(code)
        uc = {
            "Xác thực": "UC-01/02", "Tài khoản": "UC-02/06", "Catalog": "UC-03/10",
            "Giỏ hàng": "UC-04", "Khuyến mãi": "UC-04/11", "Checkout": "UC-05",
            "Thanh toán": "UC-05/06", "Đơn hàng": "UC-05/06/08", "Tồn kho": "UC-07",
            "Vận chuyển": "UC-08", "Khách hàng": "UC-09", "Quản trị": "UC-10/11",
            "AI": "UC-12", "Kiểm soát": "UC-05/07/08",
        }.get(domain, "—")
        trace_rows.append((code, domain, uc, priority, status))
    add_table(doc, ["FR", "Miền", "Use Case", "Ưu tiên", "Trạng thái"], trace_rows, [1.0, 1.35, 2.25, 0.8, 1.1], 8)
    add_para(doc, "Tài liệu chi tiết theo từng bảng: docs/ba/data-dictionary/. Tài liệu API/vận hành: platform/backend/README.md, platform/ARCHITECTURE.md và platform/ops/runbook.md.", italic=True)

    doc.save(MAIN_DOCX)
    return MAIN_DOCX


def build_requirements_document(diagrams: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc, "NÉT COMMERCE — YÊU CẦU CHỨC NĂNG VÀ PHI CHỨC NĂNG")
    set_properties(doc, "Tài liệu yêu cầu chức năng và phi chức năng NÉT Commerce", "SRS rút gọn, có truy vết và tiêu chí nghiệm thu")
    add_cover(doc, "YÊU CẦU CHỨC NĂNG VÀ PHI CHỨC NĂNG", "SRS đồng bộ với hệ thống Java Spring Boot và frontend hiện tại", "Tài liệu yêu cầu")
    doc.add_heading("1. KIỂM SOÁT TÀI LIỆU", level=1)
    add_table(doc, ["Thuộc tính", "Giá trị"], [
        ("Mục đích", "Đặc tả yêu cầu có thể kiểm thử cho phiên bản NÉT Commerce hiện tại"),
        ("Nguồn", "Mã nguồn, contract API, Flyway V1–V7, test và runtime local"),
        ("Đối tượng", "PM, BA/PO, Frontend, Backend, DevOps, người chấm đồ án"),
        ("Quy ước", "Must/Should; Đã cài đặt/Một phần/Tương lai"),
    ], [1.6, 4.9], 9)
    add_callout(doc, "PHẠM VI", "Tài liệu mô tả chính xác bản local Docker. Payment production, refund aggregate và webhook hãng vận chuyển chưa thuộc phạm vi đã cài đặt. Bốn sai lệch hiện hữu được công khai: guest online-payment truyền token qua query; expiry chưa đồng bộ payments.status; reservation COD không được job expiry quét; và webhook thành công đến muộn có thể làm payment lệch trạng thái order. Vì vậy FR-ORD-02, FR-PAY-02 và FR-INV-05 chỉ đạt một phần.", "warning")
    doc.add_heading("2. BỐI CẢNH VÀ TÁC NHÂN", level=1)
    add_figure(doc, diagrams["context"], "Bối cảnh hệ thống", 1)
    add_table(doc, ["Tác nhân", "Mục tiêu", "Ràng buộc"], [
        ("Guest", "Mua hàng nhanh và theo dõi đơn", "Guest token 07 ngày trong header"),
        ("Customer", "Tài khoản, địa chỉ, wishlist, đơn", "Chỉ dữ liệu của chính mình"),
        ("Staff", "Vận hành catalog/order/kho/giao", "RBAC ở server"),
        ("Manager", "KPI, customer, promotion, AI", "Không vượt policy"),
        ("Admin", "Quản trị toàn hệ thống", "Mọi mutation trọng yếu audit"),
    ], [1.1, 3.1, 2.3], 8.7)
    doc.add_heading("3. YÊU CẦU CHỨC NĂNG", level=1)
    add_table(doc, ["Mã", "Miền", "Mô tả", "Ưu tiên", "Hiện trạng"],
              [(c, requirement_domain(c), t, p, s) for c, t, p, s in FUNCTIONAL_REQUIREMENTS],
              [0.85, 0.9, 3.55, 0.55, 0.65], 7.8)
    doc.add_heading("4. USE CASE TRỌNG YẾU", level=1)
    add_landscape_figures(doc, [
        (diagrams["uc_storefront"], "Use Case storefront", 2),
        (diagrams["uc_admin"], "Use Case admin", 3),
        (diagrams["uc_inventory"], "Use Case inventory", 4),
        (diagrams["uc_fulfillment"], "Use Case fulfillment", 5),
    ])
    for use_case in USE_CASES:
        add_detailed_use_case(doc, use_case)
    doc.add_heading("5. YÊU CẦU PHI CHỨC NĂNG", level=1)
    add_table(doc, ["Mã", "Nhóm", "Yêu cầu định lượng/kiểm chứng", "Bằng chứng"], NON_FUNCTIONAL_REQUIREMENTS, [0.9, 1.0, 3.75, 0.85], 7.8)
    doc.add_heading("6. QUY TẮC CHẤP NHẬN", level=1)
    add_bullets(doc, [
        "Mỗi yêu cầu Must phải có UI/API hoặc bằng chứng backend, error path và test tương ứng.",
        "Mọi total/discount tồn tại trong order phải do server tính lại tại checkout.",
        "Mọi inventory mutation phải có chứng từ/movement, actor, reason và idempotency.",
        "Mọi fulfillment transition phải qua state machine và expectedVersion.",
        "Request guest tracking không được đưa token lên URL; token quá hạn phải bị từ chối.",
        "AI Copilot không được tự ghi dữ liệu hoặc che giấu thời điểm snapshot.",
    ])
    doc.add_heading("7. MA TRẬN TRUY VẾT", level=1)
    add_table(doc, ["Mã", "Use Case", "API/Module", "Dữ liệu", "Kiểm thử"], [
        ("FR-AUT-*", "UC-01/02", "/api/auth; AuthService", "users, roles, refresh_tokens, auth_tokens", "Auth unit/integration"),
        ("FR-CAT/CAR/PRO-*", "UC-03/04", "CatalogService, CartService", "products/variants/carts/promotions", "Catalog/Voucher tests"),
        ("FR-CHK/ORD/PAY-*", "UC-05/06", "OrderService, PaymentService", "orders/items/reservations/payments", "PostgreSQL integration"),
        ("FR-INV-*", "UC-07", "InventoryService", "inventory_items/operations/movements", "Inventory tests"),
        ("FR-FUL-*", "UC-08", "AdminFulfillmentService", "fulfillments/items/events", "State machine + integration"),
        ("FR-CUS-*", "UC-09", "AdminCustomerService", "users/orders/admin_customer_segments", "Segment policy tests"),
        ("FR-ADM/AI/AUD-*", "UC-10/11/12", "Admin modules", "catalog/promotion/audit/outbox", "Admin/AI tests"),
    ], [1.0, 1.0, 1.65, 1.9, 0.95], 7.8)
    doc.save(REQUIREMENTS_DOCX)
    return REQUIREMENTS_DOCX


def build_summary_document(diagrams: dict[str, Path]) -> Path:
    doc = Document()
    configure_document(doc, "NÉT COMMERCE — HỒ SƠ TỔNG HỢP")
    set_properties(doc, "Hồ sơ tổng hợp dự án NÉT Commerce", "Tóm tắt hiện trạng sản phẩm, kỹ thuật, dữ liệu và kế hoạch")
    add_cover(doc, "HỒ SƠ TỔNG HỢP DỰ ÁN NÉT COMMERCE", "Bản đọc nhanh gắn với nội dung đang triển khai", "Báo cáo tổng hợp")
    doc.add_heading("1. TỔNG QUAN", level=1)
    add_para(doc, "NÉT Commerce là đồ án website bán quần áo trực tuyến hoàn chỉnh ở môi trường local Docker. Phạm vi có storefront, tài khoản, giỏ/voucher, checkout, payment mock, theo dõi đơn; cùng backoffice catalog, khách hàng, phân khúc, order, inventory đa kho, fulfillment, báo cáo, tài chính và AI Copilot.")
    add_figure(doc, diagrams["context"], "Bối cảnh hệ thống", 1)
    doc.add_heading("2. HIỆN TRẠNG SẢN PHẨM", level=1)
    add_table(doc, ["Nhóm", "Nội dung đã có", "Giới hạn"], [
        ("Khách hàng", "22 sản phẩm/58 biến thể; auth/account/wishlist/cart/voucher/checkout/tracking", "Chưa có loyalty/recommendation; guest online-payment còn lệch header contract"),
        ("Quản trị", "Catalog/order/customer/segment/inventory/fulfillment/report/AI", "Một số cấu hình ở mức MVP"),
        ("Thanh toán", "COD + online mock; payment events có chữ ký và chống lặp", "Chưa provider/refund; còn hai tình huống lệch aggregate order/payment"),
        ("Tồn kho", "Đa kho, ATP, reservation, receipt/transfer/count/adjust", "Expiry chưa đồng bộ payment và chưa quét reservation COD"),
        ("Vận chuyển", "Tự tạo kiện theo kho; manual carrier/tracking; state machine", "Chưa webhook/label provider"),
    ], [1.1, 3.65, 1.75], 8.5)
    doc.add_heading("3. KIẾN TRÚC", level=1)
    add_figure(doc, diagrams["architecture"], "Kiến trúc triển khai", 2)
    add_table(doc, ["Thành phần", "Công nghệ"], [
        ("Frontend/BFF", "React 19, Vinext, TypeScript, Tailwind"),
        ("Backend", "Java 21, Spring Boot 4.1.1, Security, JDBC, Flyway"),
        ("Data", "PostgreSQL 17, Redis 7.4, upload volume"),
        ("Runtime", "Docker Compose; frontend 8787, API 8080"),
    ], [1.8, 4.7], 9)
    doc.add_heading("4. DỮ LIỆU VÀ NGHIỆP VỤ", level=1)
    add_para(doc, "Schema có 33 bảng nghiệp vụ và view admin_customer_segments. Order, inventory, payment và fulfillment có vòng đời riêng nhưng được nối bằng khóa ngoại, snapshot, state machine, idempotency, audit và outbox.")
    add_figure(doc, diagrams["erd_inventory"], "ERD Inventory đa kho", 3)
    add_figure(doc, diagrams["erd_fulfillment"], "ERD fulfillment", 4)
    add_figure(doc, diagrams["segments"], "Quy tắc phân khúc khách hàng", 5)
    doc.add_heading("5. YÊU CẦU VÀ KIỂM THỬ", level=1)
    add_para(doc, f"Bộ SRS hiện có {len(FUNCTIONAL_REQUIREMENTS)} yêu cầu chức năng và {len(NON_FUNCTIONAL_REQUIREMENTS)} yêu cầu phi chức năng có tiêu chí kiểm chứng. Các yêu cầu được truy vết về use case, API/module, dữ liệu và test.")
    add_table(doc, ["Bằng chứng", "Kết quả"], [
        ("Backend", "69 test pass; 10 PostgreSQL/Testcontainers"),
        ("Frontend", "12 test pass; lint/typecheck/build pass"),
        ("Runtime", "Frontend/API/PostgreSQL/Redis healthy; smoke routes HTTP 200"),
    ], [2.2, 4.3], 9)
    doc.add_heading("6. KẾ HOẠCH VÀ PHÂN CÔNG", level=1)
    add_figure(doc, diagrams["timeline"], "Timeline 08 tuần", 6)
    add_table(doc, ["Vai trò", "Đầu việc"], [
        ("PM", "Scope, backlog, timeline, RAID, UAT, tài liệu"),
        ("Frontend", "Storefront, auth/account, checkout, admin, BFF"),
        ("Backend 1", "Auth, account, catalog, cart, order, security"),
        ("Backend 2", "Inventory, payment, segment, fulfillment"),
        ("DevOps", "Docker, env, health, CI, smoke, backup/runbook"),
    ], [1.4, 5.1], 9)
    doc.add_heading("7. HOW TO RUN", level=1)
    add_numbered(doc, [
        "Tại platform/backend: sao chép .env.example thành .env rồi chạy docker --context default compose up -d --build.",
        "Tại platform: sao chép .env.example thành .env rồi chạy docker --context default compose up -d --build.",
        "Mở http://localhost:8787; API health tại http://localhost:8080/api/health.",
        "Đăng nhập demo owner@example.com / OwnerNet2026; không dùng credential này ngoài local.",
        "Chạy node ops/smoke-test.mjs để kiểm tra nhanh các route trọng yếu.",
    ])
    add_callout(doc, "GIỚI HẠN", "Bản public frontend đơn lẻ không đủ cho các luồng backend. Bản đầy đủ hiện được xác minh ở local Docker, vì Java API/PostgreSQL/Redis phải cùng hoạt động.", "warning")
    # Giữ đoạn kết thúc bắt buộc của Word ở kích thước tối thiểu để không sinh trang trắng.
    trailing_paragraph = doc.paragraphs[-1]
    if not trailing_paragraph.text.strip():
        trailing_paragraph.paragraph_format.space_before = Pt(0)
        trailing_paragraph.paragraph_format.space_after = Pt(0)
        trailing_paragraph.paragraph_format.line_spacing = Pt(1)
        set_run_font(trailing_paragraph.add_run("\u200b"), size=1, color=WHITE)
    doc.save(SUMMARY_DOCX)
    return SUMMARY_DOCX


def main():
    diagrams = generate_diagrams()
    outputs = [
        build_main_document(diagrams),
        build_requirements_document(diagrams),
        build_summary_document(diagrams),
    ]
    for output in outputs:
        print(output.name)


if __name__ == "__main__":
    main()
