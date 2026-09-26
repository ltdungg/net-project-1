from __future__ import annotations

"""Project 1: report from reviewed Markdown, with deterministic diagrams.

Preset: narrative_proposal. Named overrides: A4 Academic Report (6.5 in body),
Academic table (9.5 pt, left-aligned text), editorial cover, quiet header.
No fieldwork is fabricated; the source distinguishes SYS, SYN, PLAN and FIELD.
"""

import math
import re
import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_current_documentation as base

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/project-1/bao-cao-khao-sat-va-mo-hinh-he-thong.md"
DIAGRAMS = ROOT / "artifacts/project1-diagrams"
OUTPUT = ROOT / "deliverables/Project-1-Bao-cao-khao-sat-va-mo-hinh-he-thong-NET.docx"
SYNTHETIC = False
NAVY, BLUE, INK = base.NAVY, base.BLUE, base.INK
GRAY, PALE, GOLD = base.MUTED, base.PALE_BLUE, base.PALE_GOLD


def canvas(title, subtitle, size=(2200, 1600)):
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    d.text((60, 35), title, font=base.img_font(46, True), fill=f"#{NAVY}")
    d.text((60, 100), subtitle, font=base.img_font(29), fill=f"#{GRAY}")
    d.line((60, 150, size[0] - 60, 150), fill="#D0D5DD", width=2)
    return im, d


def text_center(d, xy, text, size=35, bold=False, fill=INK):
    font=base.img_font(size,bold)
    lines=[]
    for part in text.split("\n"):
        lines.extend(base.wrap_text(d,part,font,int(xy[2]-xy[0])))
    line_height=size+8
    total=line_height*len(lines)
    y=xy[1]+(xy[3]-xy[1]-total)/2
    for line in lines:
        width=d.textlength(line,font=font)
        d.text((xy[0]+(xy[2]-xy[0]-width)/2,y),line,font=font,fill=f"#{fill}")
        y+=line_height


def activity(d, x, y, text, w=690, h=145, fill=PALE, size=35):
    r = (x-w/2, y-h/2, x+w/2, y+h/2)
    d.rounded_rectangle(r, radius=22, fill=f"#{fill}", outline=f"#{BLUE}", width=3)
    text_center(d, (r[0]+18, r[1]+10, r[2]-18, r[3]-10), text, size)
    return r


def diamond(d, x, y, text, w=410, h=170):
    d.polygon([(x,y-h/2),(x+w/2,y),(x,y+h/2),(x-w/2,y)], fill=f"#{GOLD}", outline=f"#{BLUE}", width=3)
    text_center(d, (x-w*.30,y-h*.38,x+w*.30,y+h*.38), text, 31, True)


def route(d, pts, label=None, pos=None, arrow=True):
    d.line(pts, fill=f"#{GRAY}", width=4, joint="curve")
    if arrow:
        x1,y1=pts[-2]; x2,y2=pts[-1]; angle=math.atan2(y2-y1,x2-x1)
        d.polygon([(x2,y2),(x2-17*math.cos(angle-.45),y2-17*math.sin(angle-.45)),(x2-17*math.cos(angle+.45),y2-17*math.sin(angle+.45))], fill=f"#{GRAY}")
    if label and pos:
        f=base.img_font(28,True); bb=d.textbbox(pos,label,font=f)
        d.rectangle((bb[0]-7,bb[1]-4,bb[2]+7,bb[3]+4), fill="white")
        d.text(pos,label,font=f,fill=f"#{INK}")


def start(d,x,y):
    d.ellipse((x-17,y-17,x+17,y+17), fill=f"#{NAVY}")


def end(d,x,y):
    d.ellipse((x-24,y-24,x+24,y+24), fill="white", outline=f"#{NAVY}", width=3)
    d.ellipse((x-15,y-15,x+15,y+15), fill=f"#{NAVY}")


def merge(d,x,y):
    d.polygon([(x,y-23),(x+23,y),(x,y+23),(x-23,y)],fill="white",outline=f"#{GRAY}",width=3)


def lanes(title, subtitle, height=2140):
    im,d=canvas(title,subtitle,(1800,height))
    d.rectangle((50,170,750,height-55),outline="#D0D5DD",width=2)
    d.rectangle((750,170,1750,height-55),outline="#D0D5DD",width=2)
    d.rectangle((50,170,750,255),fill="#F4F6F9")
    d.rectangle((750,170,1750,255),fill="#F4F6F9")
    text_center(d,(65,182,735,240),"NGƯỜI DÙNG / NHÂN VIÊN",30,True)
    text_center(d,(765,182,1735,240),"HỆ THỐNG",30,True)
    return im,d


def save(im, name):
    DIAGRAMS.mkdir(parents=True,exist_ok=True)
    im.save(DIAGRAMS/name, optimize=True)


def draw_survey():
    if SYNTHETIC:
        metrics=json.loads((ROOT/"docs/project-1/synthetic/report_metrics.json").read_text(encoding="utf-8"))
        items=metrics["pain_chart"]
        im,d=canvas("Khó khăn khi mua quần áo trực tuyến", "SYN-PH-P1-120-v1 • DỮ LIỆU GIẢ LẬP • Q08 chọn nhiều • không phải khảo sát thật",(2200,1240))
        simple_labels={
            "Không chắc size":"Khó chọn cỡ quần áo",
            "Khác ảnh-mô tả":"Hàng khác ảnh hoặc mô tả",
            "Đổi trả khó-không rõ":"Đổi trả khó hoặc chưa rõ",
            "Phí hoặc tổng tiền xuất hiện muộn":"Phí, tổng tiền hiện muộn",
            "Không biết size-màu còn hàng":"Không rõ cỡ, màu còn hàng",
            "Tracking không rõ":"Khó theo dõi đơn",
            "Lo thanh toán online":"Lo thanh toán trực tuyến",
        }
        for i,item in enumerate(items):
            label,n,total=item["label"],item["count"],item["denominator"]
            label=simple_labels.get(label,label)
            pct=n/total*100; y=240+i*125
            d.text((65,y+15),label,font=base.img_font(35),fill=f"#{INK}")
            d.rounded_rectangle((720,y,720+pct*14,y+70),radius=9,fill=f"#{BLUE}")
            percent=f"{pct:.1f}".replace(".",",")
            d.text((745+pct*14,y+14),f"{n}/{total} · {percent}%",font=base.img_font(35,True),fill=f"#{NAVY}")
        d.text((65,1150),"Nguồn: SYN-PH-P1-120-v1, chỉ tính phiếu hợp lệ. Tổng tỷ lệ có thể vượt 100%.",font=base.img_font(27),fill=f"#{GRAY}")
        save(im,"survey.png")
        return
    im,d=canvas("Các vấn đề mua quần áo trực tuyến", "SYN-01 • DỮ LIỆU MÔ PHỎNG n=200 • chọn nhiều • không phải khảo sát thực địa",(2200,1240))
    items=[("Không chắc kích cỡ",136,68),("Khác ảnh/mô tả",118,59),("Đổi trả khó/không rõ",108,54),("Phí/tổng tiền xuất hiện muộn",94,47),("Không rõ cỡ, màu còn hàng",92,46),("Khó theo dõi đơn",70,35),("Lo thanh toán trực tuyến",52,26)]
    for i,(label,n,pct) in enumerate(items):
        y=240+i*125
        d.text((65,y+15),label,font=base.img_font(35),fill=f"#{INK}")
        d.rounded_rectangle((660,y,660+pct*16,y+70),radius=9,fill=f"#{BLUE}")
        d.text((690+pct*16,y+14),f"{n}/200 · {pct}%",font=base.img_font(35,True),fill=f"#{NAVY}")
    d.text((65,1150),"Nguồn: docs/research/01-bao-cao-khao-sat.md. Mẫu số chung 200; tổng tỷ lệ có thể vượt 100%.",font=base.img_font(26),fill=f"#{GRAY}")
    save(im,"survey.png")


def draw_bfd():
    im,d=canvas("BFD — Các nhóm chức năng", "Mỗi nhánh chia một chức năng lớn thành các phần nhỏ; không chỉ thứ tự làm việc",(2200,1470))
    activity(d,1100,285,"F0 · WEBSITE BÁN QUẦN ÁO NÉT",1020,145,PALE,41)
    xs=[385,1100,1815]
    groups=[("F1 · Phía khách hàng","Sản phẩm, cỡ và màu\nTài khoản và hồ sơ\nSản phẩm yêu thích",385,620),
            ("F2 · Mua và đặt hàng","Giỏ và mã giảm giá\nĐặt hàng, thanh toán\nTheo dõi, xử lý đơn",1100,620),
            ("F3 · Quản lý các kho","Tồn kho và giữ hàng\nNhập / chuyển kho\nKiểm kê / điều chỉnh",1815,620),
            ("F4 · Giao hàng và hoàn","Tạo kiện theo kho\nMã vận đơn, trạng thái\nGiao lỗi / hàng hoàn",385,1120),
            ("F5 · Quản lý bán hàng","Hồ sơ, nhóm khách\nSản phẩm, hình ảnh\nƯu đãi, mã giảm giá",1100,1120),
            ("F6 · Báo cáo và kiểm soát","Báo cáo, tài chính, trợ lý\nQuyền và lịch sử sửa đổi\nVận hành hệ thống",1815,1120)]
    route(d,[(1100,358),(1100,420),(55,420),(55,920)],arrow=False)
    route(d,[(55,420),(1815,420)],arrow=False)
    route(d,[(55,920),(1815,920)],arrow=False)
    for title,body,x,y in groups:
        route(d,[(x,420 if y==620 else 920),(x,y-140)],arrow=False)
        r=(x-285,y-140,x+285,y+165)
        d.rounded_rectangle(r,radius=18,fill="#F4F6F9",outline=f"#{BLUE}",width=3)
        text_center(d,(r[0]+15,r[1]+12,r[2]-15,r[1]+87),title,34,True)
        text_center(d,(r[0]+20,r[1]+103,r[2]-20,r[3]-18),body,33)
    save(im,"bfd-overview.png")


def draw_auth():
    im,d=lanes("A1 — Đăng ký và xác minh", "P1-FR-03 • Bản chạy thử có mã xác minh; chưa nối dịch vụ gửi email")
    start(d,400,300); route(d,[(400,317),(400,357)])
    activity(d,400,430,"Nhập email, mật khẩu\nvà thông tin đăng ký",570)
    route(d,[(685,430),(905,430)])
    activity(d,1250,430,"Kiểm tra thông tin và email")
    route(d,[(1250,503),(1250,565)])
    diamond(d,1250,650,"Hợp lệ?")
    route(d,[(1045,650),(685,650)],"Không",(820,610))
    activity(d,400,650,"Thông báo lỗi\nNgười dùng sửa lại",570,130,GOLD)
    route(d,[(115,650),(85,650),(85,430),(115,430)])
    route(d,[(1250,735),(1250,792)],"Có",(1270,751))
    activity(d,1250,875,"Tạo tài khoản chờ xác minh\nvà mã dùng một lần",690,165)
    route(d,[(1250,958),(1250,1027)])
    activity(d,1250,1110,"Cấp mã cho bản chạy thử\nKhi dùng thật: cần gửi email",690,165)
    route(d,[(905,1110),(685,1110)])
    activity(d,400,1110,"Nhập mã / mở liên kết\nxác minh email",570)
    route(d,[(400,1183),(400,1310),(905,1310)])
    activity(d,1250,1310,"Kiểm tra mã, hạn dùng\nvà việc mã đã được dùng")
    route(d,[(1250,1383),(1250,1455)])
    diamond(d,1250,1540,"Mã đúng?")
    route(d,[(1045,1540),(685,1540)],"Không",(805,1500))
    activity(d,400,1540,"Mã bị từ chối; có thể\nxin mã mới rồi thử lại",570,165,GOLD)
    route(d,[(400,1623),(400,1760)]); end(d,400,1785)
    route(d,[(1250,1625),(1250,1712)],"Có",(1270,1650))
    activity(d,1250,1800,"Xác nhận email hợp lệ\nCho phép dùng tài khoản",690,175)
    route(d,[(1250,1888),(1250,1990)]); end(d,1250,2015)
    save(im,"activity-auth.png")


def draw_checkout():
    im,d=lanes("A2 — Đặt hàng từ nhiều kho", "Tồn khả dụng = tồn vật lý − hàng giữ; mỗi dòng hàng lấy đủ từ một kho",2290)
    start(d,400,285); route(d,[(400,302),(400,320)])
    activity(d,400,405,"Chọn hàng, số lượng, mã giảm\nĐịa chỉ và cách trả tiền",570,170)
    route(d,[(685,405),(1250,405),(1250,515)])
    diamond(d,1250,600,"Trùng mã\nyêu cầu?")
    route(d,[(1250,685),(1250,715)],"Không",(1270,685))
    route(d,[(1045,600),(685,600)],"Có",(825,560))
    activity(d,400,600,"Nhận lại đơn cũ\nKhông tạo thêm giao dịch",570)
    route(d,[(400,673),(400,900)]); end(d,400,925)
    activity(d,1250,800,"Tính lại giá, mã giảm giá\nKiểm tra và khóa lượng tồn",690,170)
    route(d,[(1250,885),(1250,925)])
    diamond(d,1250,1010,"Đủ điều kiện\nvà tồn?")
    route(d,[(1045,1010),(720,1010),(720,1160),(685,1160)],"Không",(780,970))
    activity(d,400,1160,"Nhận lỗi; sửa giỏ/địa chỉ\nKhông có đơn dở dang",570,170,GOLD)
    route(d,[(400,1245),(400,1330)]); end(d,400,1355)
    route(d,[(1250,1095),(1250,1170)],"Có",(1270,1125))
    activity(d,1250,1340,"Ghi cùng một giao dịch:\nĐơn, hàng và giá lúc đặt\nPhần hàng giữ theo kho\nKiện giao và hàng trong kiện\nThanh toán, lịch sử, sự kiện",690,340)
    route(d,[(1250,1510),(1250,1585)])
    diamond(d,1250,1670,"Trả tiền?")
    route(d,[(1045,1670),(400,1670),(400,1795)],"Khi nhận",(700,1630))
    route(d,[(1250,1755),(1250,1795)],"Trực tuyến",(1270,1760))
    activity(d,400,1880,"Đơn đã được xác nhận\nChờ thu tiền khi giao đủ",570,170)
    activity(d,1250,1880,"Đơn đang chờ thanh toán\nChờ xác nhận từ bản mô phỏng",690,170)
    route(d,[(400,1965),(400,2090),(850,2090)],arrow=False)
    route(d,[(1250,1965),(1250,2090),(850,2090)],arrow=False)
    merge(d,850,2090); route(d,[(850,2113),(850,2160)]); end(d,850,2185)
    save(im,"activity-checkout.png")


def draw_inventory():
    im,d=lanes("A3 — Chứng từ nhập / chuyển kho", "Mỗi lệnh ghi một dòng và có hiệu lực ngay; chưa có bước duyệt bản nháp")
    start(d,400,300); route(d,[(400,317),(400,347)])
    activity(d,400,430,"Chọn hàng, kho, số lượng\nGhi lý do và mã chống lặp",570,165)
    route(d,[(685,430),(905,430)])
    activity(d,1250,430,"Kiểm tra quyền và mã chống lặp")
    route(d,[(1250,503),(1250,575)])
    diamond(d,1250,660,"Trùng mã?")
    route(d,[(1045,660),(685,660)],"Có",(825,620))
    activity(d,400,660,"Nhận lại chứng từ đã có",570,130)
    route(d,[(400,725),(400,805)]); end(d,400,830)
    route(d,[(1250,745),(1250,802)],"Không",(1270,760))
    activity(d,1250,890,"Kiểm tra kho, hàng, số lượng\nKhóa dòng tồn cần thay đổi",690,175)
    route(d,[(1250,978),(1250,1045)])
    diamond(d,1250,1130,"Hợp lệ?")
    route(d,[(1045,1130),(685,1130)],"Không",(800,1090))
    activity(d,400,1130,"Từ chối và báo lý do\nKhông đổi tồn",570,165,GOLD)
    route(d,[(400,1213),(400,1310)]); end(d,400,1335)
    route(d,[(1250,1215),(1250,1280)],"Có",(1270,1240))
    activity(d,1250,1400,"Nhập: tăng tồn vật lý\nChuyển: hai kho khác nhau,\nnguồn đủ hàng chưa giữ;\ntrừ nguồn và cộng đích cùng lúc",690,240)
    route(d,[(1250,1520),(1250,1590)])
    activity(d,1250,1700,"Lưu chứng từ có hiệu lực\nBiến động tồn, lịch sử cùng lúc\nXác nhận và trả số tồn mới",690,220)
    route(d,[(1250,1810),(1250,1900)]); end(d,1250,1925)
    save(im,"activity-inventory.png")


def draw_fulfillment():
    im,d=lanes("A4 — Giao hàng và hàng hoàn", "Nhân viên tự cập nhật; chưa nối dịch vụ của hãng giao hàng",2420)
    start(d,400,300); route(d,[(400,317),(400,345)])
    activity(d,400,430,"Mở kiện và đóng gói\nNhập hãng, mã vận đơn",570,170)
    route(d,[(685,430),(905,430)])
    activity(d,1250,430,"Kiểm tra thanh toán, quyền,\ntrạng thái và dữ liệu có bị đổi",690,170)
    route(d,[(1250,515),(1250,575)])
    diamond(d,1250,660,"Đủ điều kiện\ngửi hàng?")
    route(d,[(1045,660),(685,660)],"Không",(800,620))
    activity(d,400,660,"Báo lỗi; sửa dữ liệu\nhoặc tải bản mới nhất",570,165,GOLD)
    route(d,[(115,660),(85,660),(85,430),(115,430)])
    route(d,[(1250,745),(1250,815)],"Có",(1270,767))
    activity(d,1250,910,"Đánh dấu đã gửi hàng\nChỉ trừ hàng thuộc kiện này\nGiảm tồn và hàng giữ, ghi sổ",690,190)
    route(d,[(905,910),(685,910)])
    activity(d,400,910,"Bàn giao hàng cho hãng\nTheo dõi kết quả giao",570,170)
    route(d,[(400,995),(400,1170),(905,1170)])
    activity(d,1250,1170,"Nhận kết quả giao hàng\nKiểm tra trạng thái, dữ liệu mới",690,170)
    route(d,[(1250,1255),(1250,1325)])
    diamond(d,1250,1410,"Đã giao?")
    route(d,[(1045,1410),(400,1410),(400,1530)],"Không",(750,1370))
    route(d,[(1250,1495),(1250,1530)],"Có",(1270,1500))
    activity(d,1250,1640,"Ghi đã giao, cập nhật đơn\nVới trả tiền khi nhận hàng:\nchỉ ghi đã trả khi giao đủ mọi kiện",690,220)
    activity(d,400,1640,"Giao lỗi / hàng hoàn: ghi lý do\nGiao lỗi: sắp xếp giao lại\nHàng hoàn: chưa cộng lại tồn",570,220,GOLD)
    route(d,[(400,1750),(400,1830)])
    activity(d,400,1940,"Nếu hàng hoàn: kiểm tra lại\nChỉ nhập / chỉnh tồn\nkhi hàng đủ điều kiện",570,220,GOLD)
    route(d,[(1250,1750),(1250,2190),(850,2190)],arrow=False)
    route(d,[(400,2050),(400,2190),(850,2190)],arrow=False)
    merge(d,850,2190); route(d,[(850,2213),(850,2280)]); end(d,850,2305)
    save(im,"activity-fulfillment.png")


def entity(d,x,y,name,ref=False,w=475,h=150):
    r=(x-w/2,y-h/2,x+w/2,y+h/2)
    d.rounded_rectangle(r,radius=16,fill=f"#{GOLD if ref else PALE}",outline=f"#{BLUE}",width=3)
    text_center(d,(r[0]+15,r[1]+12,r[2]-15,r[3]-12),name,39,True)
    return r


def rel(d,pts,a,b,apos,bpos,label=None,lpos=None):
    route(d,pts,arrow=False)
    for t,p in [(a,apos),(b,bpos)]:
        f=base.img_font(31,True); bb=d.textbbox(p,t,font=f)
        d.rectangle((bb[0]-6,bb[1]-4,bb[2]+6,bb[3]+4),fill="white")
        d.text(p,t,font=f,fill=f"#{NAVY}")
    if label and lpos:
        d.text(lpos,label,font=base.img_font(26),fill=f"#{GRAY}")


def draw_erd_catalog():
    im,d=canvas("C1 — Tài khoản, sản phẩm và giỏ", "ERD KHÁI NIỆM • các nhóm dữ liệu và số lượng liên kết",(2200,1530))
    for x,y,n in [(350,310,"Tài khoản"),(1050,310,"Giỏ hàng"),(1750,310,"Hàng trong giỏ"),(350,830,"Địa chỉ"),(1050,830,"Sản phẩm"),(1750,830,"Cỡ/màu cụ thể"),(350,1320,"Vai trò"),(1050,1320,"Danh mục")]: entity(d,x,y,n)
    rel(d,[(588,310),(812,310)],"0..1","0..N",(606,260),(730,260))
    rel(d,[(1288,310),(1512,310)],"1","0..N",(1310,260),(1420,260))
    rel(d,[(1750,385),(1750,755)],"0..N","1",(1770,430),(1770,690))
    rel(d,[(1288,830),(1512,830)],"1","1..N",(1310,780),(1420,780))
    rel(d,[(350,385),(350,755)],"1","0..N",(370,425),(370,690))
    rel(d,[(112,310),(55,310),(55,1320),(112,1320)],"0..N","1..N",(65,425),(65,1180))
    rel(d,[(1050,905),(1050,1245)],"0..N","0..N",(1070,940),(1070,1170))
    d.text((1390,1245),"Giỏ không bắt buộc có tài khoản.\nMỗi mẫu đang bán có cỡ/màu cụ thể.\nSản phẩm yêu thích được mô tả riêng,\nkhông trình bày trong hình này.",font=base.img_font(29),fill=f"#{GRAY}",spacing=12)
    save(im,"erd-catalog.png")


def draw_erd_commerce():
    im,d=canvas("C2 — Đơn hàng và thanh toán", "ERD KHÁI NIỆM • các liên kết trong phiên bản hiện tại",(2200,1530))
    for x,y,n,ref in [(350,310,"Tài khoản",True),(1050,310,"Đơn hàng",False),(1750,310,"Thanh toán",False),(350,850,"Cỡ/màu cụ thể",True),(1050,850,"Hàng của đơn",False),(1750,850,"Mã giảm giá",False),(1750,1320,"Chương trình\nkhuyến mãi",False)]: entity(d,x,y,n,ref)
    rel(d,[(588,310),(812,310)],"0..1","0..N",(605,260),(730,260))
    rel(d,[(1288,310),(1512,310)],"1","1..N",(1310,260),(1420,260))
    rel(d,[(1050,385),(1050,775)],"1","1..N",(1070,440),(1070,710))
    rel(d,[(588,850),(812,850)],"0..1","0..N",(608,800),(730,800))
    rel(d,[(1200,385),(1200,605),(1750,605),(1750,775)],"0..N","0..1",(1220,425),(1770,700),"áp dụng",(1450,565))
    rel(d,[(1750,925),(1750,1245)],"0..N","1",(1770,970),(1770,1170))
    d.text((120,1120),"Ô vàng: nhóm dữ liệu đã có ở hình khác.\nĐơn giữ thông tin lúc đặt dù hàng cũ bị xóa.\nĐặt hàng tạo 1 khoản thanh toán; có thể lưu nhiều khoản.\nMỗi đơn dùng tối đa 1 mã giảm, lưu lại giá trị ưu đãi.\nChức năng hoàn tiền đầy đủ chưa được triển khai.",font=base.img_font(29),fill=f"#{GRAY}",spacing=15)
    save(im,"erd-commerce.png")


def draw_erd_fulfillment():
    im,d=canvas("C3 — Tồn kho và giao hàng", "ERD KHÁI NIỆM • một kiện cho mỗi cặp đơn hàng × kho",(2200,1570))
    for x,y,n,ref in [(350,310,"Kho",False),(1050,310,"Tồn kho",False),(1750,310,"Cỡ/màu cụ thể",True),(350,810,"Đơn hàng",True),(1050,810,"Hàng giữ cho đơn",False),(1750,810,"Hàng của đơn",True),(350,1310,"Kiện giao",False),(1050,1310,"Hàng trong kiện",False)]: entity(d,x,y,n,ref)
    rel(d,[(588,310),(812,310)],"1","0..N",(607,260),(730,260))
    rel(d,[(1288,310),(1512,310)],"0..N","1",(1305,260),(1460,260))
    rel(d,[(1050,385),(1050,735)],"1","0..N",(1070,430),(1070,670))
    rel(d,[(588,810),(812,810)],"1","1..N",(607,760),(730,760))
    rel(d,[(350,885),(350,1235)],"1","1..N",(370,930),(370,1170))
    rel(d,[(588,1310),(812,1310)],"1","1..N",(607,1260),(730,1260))
    rel(d,[(1050,885),(1050,1235)],"0..1","0..1",(1070,930),(1070,1170))
    rel(d,[(1750,885),(1750,1310),(1288,1310)],"1","0..N",(1770,930),(1325,1260))
    rel(d,[(112,310),(55,310),(55,1310),(112,1310)],"1","0..N",(65,460),(65,1150))
    d.text((1330,1400),"Mỗi kho có một dòng tồn cho một cỡ/màu.\nMột dòng hàng chưa chia qua nhiều kho.\nPhiếu cũ có thể thiếu liên kết hàng giữ.",font=base.img_font(27),fill=f"#{GRAY}",spacing=10)
    save(im,"erd-fulfillment.png")


def figure(doc,path,caption,number):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(3)
    with Image.open(path) as im:
        width=min(6.35,7.55*im.width/im.height)
    pic=p.add_run().add_picture(str(path),width=Inches(width))
    pic._inline.docPr.set("descr",caption); pic._inline.docPr.set("name",f"Hình {number}")
    cp=doc.add_paragraph(style="Caption"); cp.paragraph_format.keep_with_next=False
    cp.add_run("Hình ")
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"SEQ Hinh \\* ARABIC")
    r=OxmlElement("w:r"); t=OxmlElement("w:t"); t.text=str(number); r.append(t); field.append(r); cp._p.append(field)
    cp.add_run(". "+caption)


def make_numbering(doc,kind):
    n=base.create_numbering_instance(doc)
    root=doc.part.numbering_part.element
    num=root.xpath(f'./w:num[@w:numId="{n}"]')[0]
    aid=num.find(qn("w:abstractNumId")).get(qn("w:val"))
    lvl=root.xpath(f'./w:abstractNum[@w:abstractNumId="{aid}"]/w:lvl')[0]
    lvl.find(qn("w:numFmt")).set(qn("w:val"),"bullet" if kind=="bullet" else "decimal")
    lvl.find(qn("w:lvlText")).set(qn("w:val"),"•" if kind=="bullet" else "%1.")
    ppr=lvl.find(qn("w:pPr")); spacing=OxmlElement("w:spacing")
    for key,val in {"before":"0","after":"80","line":"290","lineRule":"auto"}.items(): spacing.set(qn("w:"+key),val)
    ppr.append(spacing)
    if kind=="decimal":
        ppr.find(qn("w:tabs")).find(qn("w:tab")).set(qn("w:pos"),"650")
        ppr.find(qn("w:ind")).set(qn("w:left"),"650")
        ppr.find(qn("w:ind")).set(qn("w:hanging"),"390")
    return n


def add_list(doc,items,kind):
    nid=make_numbering(doc,kind)
    for text in items:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.208
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Inches(.451 if kind=="decimal" else .375)
        p.paragraph_format.first_line_indent=Inches(-.271 if kind=="decimal" else -.194)
        numpr=OxmlElement("w:numPr"); il=OxmlElement("w:ilvl"); il.set(qn("w:val"),"0"); num=OxmlElement("w:numId"); num.set(qn("w:val"),str(nid)); numpr.extend((il,num)); p._p.get_or_add_pPr().append(numpr)
        run=p.add_run(text)
        if re.match(r"^S\d:",text):
            base.set_run_font(run,size=10)
            p.paragraph_format.line_spacing=1.08


def table(doc,headers,rows,widths):
    t=base.add_table(doc,headers,rows,widths,font_size=9.5)
    # A trailing full-size spacer can create a blank page before a chapter.
    spacer=doc.paragraphs[-1]
    if not spacer.text and not spacer._p.xpath('.//w:drawing'):
        spacer._p.getparent().remove(spacer._p)
    compact=headers[0]=="Thuật ngữ" or (SYNTHETIC and headers[0]=="Nhóm")
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            base.set_cell_margins(cell,top=40 if compact else 80,bottom=40 if compact else 80,start=120,end=120)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing=1.08
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.space_after=Pt(0 if compact else 2)
                p.paragraph_format.keep_with_next=(ri==0)
    return t


def cover(doc):
    p=doc.add_paragraph("HỌC PHẦN PROJECT 1"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(45); p.paragraph_format.space_after=Pt(28)
    base.set_run_font(p.runs[0],size=13,color=BLUE,bold=True)
    for line in ["BÁO CÁO KHẢO SÁT", "VÀ PHÂN TÍCH HỆ THỐNG"]:
        p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(8); base.set_run_font(p.add_run(line),size=26,color=NAVY,bold=True)
    p=doc.add_paragraph("Website bán quần áo trực tuyến NÉT Commerce"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(35)
    base.set_run_font(p.runs[0],size=15,color=BLUE)
    for label,value in [("Trường / Khoa / Lớp","[Người học bổ sung]"),("Giảng viên hướng dẫn","[Người học bổ sung]"),("Nhóm thực hiện","05 người: 01 quản lý, 03 lập trình, 01 triển khai"),("Tên / mã sinh viên","[Người học bổ sung danh sách 05 thành viên]"),("Thời gian khảo sát","02 tuần — bài khảo sát giả lập" if SYNTHETIC else "02 tuần — kế hoạch đề xuất"),("Phiên bản","1.2 · bài khảo sát giả lập · 03/09/2026" if SYNTHETIC else "1.0 · ngày 03/09/2026")]:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9)
        base.set_run_font(p.add_run(label+": "),bold=True); p.add_run(value)
    p=doc.add_paragraph("DỮ LIỆU GIẢ LẬP — 120 phiếu hợp lệ, 12 cuộc phỏng vấn mẫu, 30 hồ sơ mẫu" if SYNTHETIC else "Kế hoạch khảo sát • Tìm hiểu hệ thống • Kết quả minh họa được ghi rõ"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(28); base.set_run_font(p.runs[0],size=10,color=GRAY,italic=True)
    doc.add_page_break()
    p=doc.add_paragraph("MỤC LỤC"); base.set_run_font(p.runs[0],size=16,color=BLUE,bold=True)
    p.paragraph_format.space_after=Pt(12)
    p=doc.add_paragraph(); field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),'TOC \\o "1-2" \\h \\z \\u'); p._p.append(field)


def build(synthetic=False):
    global SOURCE, OUTPUT, DIAGRAMS, SYNTHETIC
    SYNTHETIC=synthetic
    if synthetic:
        SOURCE=ROOT/"docs/project-1/bao-cao-khao-sat-gia-lap.md"
        OUTPUT=ROOT/"deliverables/Project-1-Bao-cao-khao-sat-GIA-LAP-NET.docx"
        DIAGRAMS=ROOT/"artifacts/project1-synthetic-diagrams"
    for fn in [draw_survey,draw_bfd,draw_auth,draw_checkout,draw_inventory,draw_fulfillment,draw_erd_catalog,draw_erd_commerce,draw_erd_fulfillment]: fn()
    doc=Document(); base.configure_document(doc,"PROJECT 1  |  NÉT COMMERCE  |  BÀI KHẢO SÁT GIẢ LẬP — SYN" if synthetic else "PROJECT 1  |  NÉT COMMERCE  |  KHẢO SÁT VÀ PHÂN TÍCH YÊU CẦU")
    if synthetic:
        # Named override: compact academic body. Keep 11 pt text while avoiding
        # isolated continuation pages in this longer survey edition.
        doc.styles["Normal"].paragraph_format.line_spacing=1.15
        doc.styles["Normal"].paragraph_format.space_after=Pt(6)
        for name in ["TOC 1","TOC 2"]:
            st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            st.base_style=doc.styles["Normal"]
            st.font.name="Calibri"; st.font.size=Pt(10)
            st.paragraph_format.line_spacing=1.0
            st.paragraph_format.space_before=Pt(0)
            st.paragraph_format.space_after=Pt(2)
    # Quiet academic header, no rule. Named overrides are applied consistently.
    header=doc.sections[0].header.paragraphs[0]
    for node in header._p.xpath('./w:pPr/w:pBdr'): node.getparent().remove(node)
    for s in ["Heading 1","Heading 2","Heading 3"]: doc.styles[s].paragraph_format.line_spacing=1.1
    for s in ["Title","Subtitle"]:
        st=doc.styles[s]; st.font.name="Calibri"; st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.1
    doc.styles["Title"].font.size=Pt(26); doc.styles["Title"].font.color.rgb=base.rgb(NAVY)
    doc.styles["Subtitle"].font.size=Pt(15); doc.styles["Subtitle"].font.color.rgb=base.rgb(BLUE)
    doc.core_properties.title="Project 1 — Bài khảo sát giả lập và phân tích hệ thống NÉT Commerce" if synthetic else "Project 1 — Khảo sát và mô hình hệ thống NÉT Commerce"
    doc.core_properties.subject="02 tuần khảo sát; phỏng vấn, phiếu hỏi, hồ sơ mẫu; yêu cầu và sơ đồ hệ thống"
    doc.core_properties.author="Nhóm đồ án NÉT Commerce"
    cover(doc)
    lines=SOURCE.read_text(encoding="utf-8").splitlines(); i=0; widths=None; fig=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith("<!-- widths:"):
            widths=[float(x.strip()) for x in line.split(":",1)[1].split("-->")[0].split(",")]; i+=1; continue
        if line=="<!-- pagebreak -->":
            doc.add_page_break(); i+=1; continue
        if line.startswith("#"):
            level=len(line)-len(line.lstrip("#")); text=line[level:].strip()
            p=doc.add_heading(text,level=level)
            if level==1: p.paragraph_format.page_break_before=True
            i+=1; continue
        if line.startswith("!["):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line); fig+=1
            path=(SOURCE.parent/match.group(2)).resolve(); figure(doc,path,match.group(1),fig)
            i+=1; continue
        if line.startswith("|"):
            cells=lambda s:[c.strip() for c in s.strip().strip("|").split("|")]
            headers=cells(line); i+=2; rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"): rows.append(tuple(cells(lines[i]))); i+=1
            table(doc,headers,rows,widths or [6.5/len(headers)]*len(headers)); widths=None; continue
        if line.startswith("- ") or re.match(r'^\d+\. ',line):
            kind="bullet" if line.startswith("- ") else "decimal"; items=[]
            pattern=r'^- ' if kind=="bullet" else r'^\d+\. '
            while i<len(lines) and re.match(pattern,lines[i].strip()):
                items.append(re.sub(pattern,"",lines[i].strip())); i+=1
            add_list(doc,items,kind); continue
        if line.startswith("> "):
            p=doc.add_paragraph(line[2:]); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(10)
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),GOLD); p._p.get_or_add_pPr().append(shade)
            for r in p.runs: base.set_run_font(r,size=10,color=INK)
            i+=1; continue
        p=doc.add_paragraph(line)
        previous=p._p.getprevious()
        if previous is not None and previous.tag==qn("w:tbl"):
            p.paragraph_format.space_before=Pt(6)
        # Question forms and source paths stay left aligned for readability.
        if re.match(r'^Q\d{2}\.',line) or line.startswith("Mã "):
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.keep_together=True
        i+=1
    OUTPUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUTPUT)
    print(OUTPUT)
    print(f"Figures: {fig}; tables: {len(doc.tables)}; paragraphs: {len(doc.paragraphs)}")


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--synthetic",action="store_true",help="Build the explicitly labeled synthetic survey edition; preserve baseline v1.0.")
    build(synthetic=parser.parse_args().synthetic)
